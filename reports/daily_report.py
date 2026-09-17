"""
reports/daily_report.py — Trading System v2

Daily EOD Report Generator - produces 7-sheet xlsx report at 16:05 IST.
Replaces smoke_test_review.xlsx workflow with auto-generated structured report.

Usage (CLI):
    python -m reports.daily_report [--date YYYY-MM-DD] [--db PATH]

Usage (cron):
    5 16 * * 1-5 ubuntu /home/ubuntu/systems/venv/bin/python \\
        /home/ubuntu/systems/trading-system/reports/daily_report.py \\
        >> /home/ubuntu/systems/trading-system/logs/daily_report.log 2>&1

Output: reports/output/daily_report_YYYY-MM-DD.xlsx

Sheets:
    0_EOD_Dashboard     - Master summary (5 sections A-E)
    1_Signals           - Signal funnel with recon check
    2_Orders            - Order lifecycle with deviations
    3_Capital           - Funds tracker with reconciliation
    4_Candles           - Per-candle analysis with tune suggestions
    5_Telegram          - Alert log
    6_Strategy_Analysis - Strategy and time-of-day performance
"""

from __future__ import annotations

import argparse
import json
import logging
import sys
import time
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from reports.style_constants import (
    FONT_BODY, FONT_HEADER, FONT_TITLE, FONT_SECTION_HEADER, FONT_WHITE_BOLD,
    FILL_GREEN, FILL_RED, FILL_AMBER, FILL_GREY, FILL_SEPARATOR, FILL_HEADER,
    FILL_TITLE_BG, FILL_GROUP_HEADER, FILL_SUB_HEADER,
    BORDER_ALL, ALIGN_CENTER, ALIGN_LEFT, ALIGN_RIGHT,
    NUM_FMT_CURRENCY, NUM_FMT_CURRENCY_NEG_PARENS, NUM_FMT_PERCENT, NUM_FMT_RATIO,
    EXIT_REASON_FILLS, ALERT_TYPE_FILLS, DELIVERY_FILLS,
    COLOR_WHITE,
)
from core.time_authority import now_ist
from reports import signal_status as sig_status
from utils.holiday_guard import is_trading_day

# M-R2 (audit 04-Jul): CLOSED_MANUAL trades ARE closed with realized P&L; the report
# filtered status=='CLOSED' only, understating P&L on any EOD-close day. Single source:
_CLOSED_STATUSES = ("CLOSED", "CLOSED_MANUAL")

log = logging.getLogger("daily_report")


# ─────────────────────────────────────────────────────────────────────────────
# Data container
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class ReportData:
    """Container for all data needed by the report."""
    date_iso: str
    mode: str
    account: str
    opening_capital: float

    signals: List[dict]
    trades: List[dict]
    orders: List[dict]
    fm_ledger: List[dict]
    screener_results: List[dict]
    innings: List[dict]
    system_events: List[dict]
    recon_log: List[dict]
    gate_state: List[dict]

    excluded_symbols: List[str]
    candle_map: Dict[Tuple[str, str], dict] = field(default_factory=dict)
    excursion_map: Dict[str, dict] = field(default_factory=dict)
    strategy_min_scores: Dict[str, int] = field(default_factory=dict)
    broker_rates: Dict[str, Any] = field(default_factory=dict)


# ─────────────────────────────────────────────────────────────────────────────
# Holiday check
# ─────────────────────────────────────────────────────────────────────────────

def is_holiday_or_weekend(date_iso: str, config_dir: Path) -> bool:
    """Check if date is a weekend or NSE holiday.

    M-R1 (audit 04-Jul): the old inline check did `date_iso in holidays.get("holidays")`,
    but the production nse_holidays_<year>.yaml lists holidays as DICTS ({date:, name:}),
    so a plain-string membership test NEVER matched — the report generated and Telegrammed
    on NSE holidays (violating the 'no real alerts on non-trading days' rule). Delegate the
    holiday half to utils.holiday_guard (the single source; it parses BOTH the string and
    dict entry formats). Weekend stays a local check so a missing YAML still blocks weekends.
    """
    dt = date.fromisoformat(date_iso)
    if dt.weekday() >= 5:
        return True
    try:
        return not is_trading_day(dt, config_dir)
    except FileNotFoundError:
        # No holiday file for the year -> can't be a listed holiday; same as the old behaviour.
        return False


# ─────────────────────────────────────────────────────────────────────────────
# Data loading
# ─────────────────────────────────────────────────────────────────────────────

def load_report_data(store, date_iso: str, config_dir: Path) -> ReportData:
    """Load all data for the report from state_store (DB-only, v14+)."""
    signals = store.get_signals_for_date(date_iso)
    trades = store.get_trades_for_date(date_iso)
    orders = store.get_orders_for_date(date_iso)
    fm_ledger = store.get_fm_ledger_for_date(date_iso)
    screener_results = store.get_screener_results_for_date(date_iso)
    innings = store.get_innings_for_date(date_iso)
    system_events = store.get_system_events_for_date(date_iso)
    recon_log = store.get_reconciliation_log_for_date(date_iso)

    try:
        gate_state = store.get_all_gate_state()
    except Exception:
        gate_state = []

    session = store.get_session_row()
    mode = session["mode"] if session else "UNKNOWN"
    account = session["account_id"] if session else "UNKNOWN"

    excluded_symbols = []
    global_min = 60
    broker_rates: Dict[str, Any] = {}
    try:
        import yaml
        system_config_path = config_dir / "system_config.yaml"
        if system_config_path.exists():
            with open(system_config_path, "r") as f:
                sys_cfg = yaml.safe_load(f) or {}
            excluded_symbols = sys_cfg.get("excluded_symbols", [])
        scoring_path = config_dir / "scoring_weights.yaml"
        if scoring_path.exists():
            with open(scoring_path, "r") as f:
                scoring_cfg = yaml.safe_load(f) or {}
            global_min = scoring_cfg.get("min_pass_score", 60)
        broker_costs_path = config_dir / "broker_costs.yaml"
        if broker_costs_path.exists():
            with open(broker_costs_path, "r") as f:
                broker_cfg = yaml.safe_load(f) or {}
            broker_rates = broker_cfg.get("zerodha", {})
    except Exception:
        pass

    strategy_min_scores = _build_strategy_min_scores(config_dir, global_min)

    candle_rows = store.get_candles_for_date(date_iso)
    candle_map: Dict[Tuple[str, str], dict] = {}
    for c in candle_rows:
        ts = c.get("ts", "")
        hhmm = ts[11:16] if len(ts) >= 16 else ""
        if hhmm:
            candle_map[(c["symbol"], hhmm)] = {
                "open": c["open"], "high": c["high"],
                "low": c["low"], "close": c["close"],
                "is_synthetic": c.get("is_synthetic", 0),
            }

    if not candle_map:
        candle_csv_dir = Path(config_dir).parent / "data_store" / "candles"
        csv_candles = _load_candle_csv(candle_csv_dir, date_iso)
        candle_map.update(csv_candles)

    excursion_rows = store.get_trade_excursions_for_date(date_iso)
    excursion_map: Dict[str, dict] = {r["trade_id"]: r for r in excursion_rows}

    opening_capital = 0.0

    init_rows = [r for r in fm_ledger if r.get("entry_type") == "INIT"]
    if init_rows:
        opening_capital = init_rows[0].get("balance_after", 0.0)

    return ReportData(
        date_iso=date_iso,
        mode=mode,
        account=account,
        opening_capital=opening_capital,
        signals=signals,
        trades=trades,
        orders=orders,
        fm_ledger=fm_ledger,
        screener_results=screener_results,
        innings=innings,
        system_events=system_events,
        recon_log=recon_log,
        gate_state=gate_state,
        excluded_symbols=excluded_symbols,
        candle_map=candle_map,
        excursion_map=excursion_map,
        strategy_min_scores=strategy_min_scores,
        broker_rates=broker_rates,
    )


# ─────────────────────────────────────────────────────────────────────────────
# Helper functions
# ─────────────────────────────────────────────────────────────────────────────

def _fmt_time(ts_str: Optional[str]) -> str:
    """Extract HH:MM:SS from ISO timestamp."""
    if not ts_str:
        return ""
    try:
        dt = datetime.fromisoformat(ts_str)
        return dt.strftime("%H:%M:%S")
    except (ValueError, TypeError):
        return str(ts_str)[:8] if ts_str else ""


def _fmt_datetime(ts_str: Optional[str]) -> str:
    """Format ISO timestamp to YYYY-MM-DD HH:MM:SS."""
    if not ts_str:
        return ""
    try:
        dt = datetime.fromisoformat(ts_str)
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    except (ValueError, TypeError):
        return str(ts_str)


def _get_order_for_trade_leg(orders: List[dict], trade_id: str, leg: str) -> Optional[dict]:
    """Find order for a specific trade and leg."""
    for o in orders:
        if o.get("trade_id") == trade_id and o.get("leg") == leg:
            return o
    return None


def _calc_slip_pct(sys_price: float, fill_price: float) -> float:
    """Calculate slippage percentage."""
    if sys_price == 0:
        return 0.0
    return abs(fill_price - sys_price) / sys_price * 100


# Routine, non-incident system events — must NOT inflate the CRITICAL tally.
# KILL_AUTO_CLEARED is the daily prior-day kill auto-clear at boot
# (kill_switch._record_cleared_kill); it fires every trading day and is INFO-grade,
# not an incident. It is still surfaced below as a "Kill Switch Events" tally —
# reclassified, not hidden.
_ROUTINE_SYSTEM_EVENT_TYPES = frozenset({"KILL_AUTO_CLEARED"})


def _is_critical_event(event_type: Optional[str]) -> bool:
    """True iff a system_event is incident-grade (event_type contains CRITICAL or KILL)
    AND is not a routine auto-clear. Keyed on the STRUCTURED event_type, never a free-text
    severity: the only KILL/CRITICAL event_type ever written is the routine
    KILL_AUTO_CLEARED, so the previous bare substring match reported a false
    'CRITICAL Count >= 1' every single trading day. A genuine future KILL/CRITICAL
    event_type still counts (anti-vacuity)."""
    et = (event_type or "").upper()
    if et in _ROUTINE_SYSTEM_EVENT_TYPES:
        return False
    return "CRITICAL" in et or "KILL" in et


def _apply_cell_style(cell, font=None, fill=None, border=None, alignment=None, number_format=None):
    """Apply styles to a cell."""
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    if alignment:
        cell.alignment = alignment
    if number_format:
        cell.number_format = number_format


def _set_column_widths(ws: Worksheet, widths: Dict[str, float]):
    """Set column widths by letter."""
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _disable_gridlines(ws: Worksheet):
    """Turn off gridlines for sheet."""
    ws.sheet_view.showGridLines = False


def _add_separator_column(ws: Worksheet, col_idx: int, start_row: int, end_row: int):
    """Add dark blue separator column."""
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = 2
    for row in range(start_row, end_row + 1):
        cell = ws.cell(row=row, column=col_idx)
        cell.fill = FILL_SEPARATOR


# ─────────────────────────────────────────────────────────────────────────────
# Private fallbacks for historical data (eligible_score, cost breakdown, CSV candles)
# ─────────────────────────────────────────────────────────────────────────────

def _build_strategy_min_scores(config_dir: Path, global_min: int) -> Dict[str, int]:
    """Read all strategy YAMLs and return effective min_score per strategy name.

    Private fallback for signals where screener_results.eligible_score is NULL.
    """
    import yaml as _yaml
    strategies_dir = config_dir / "strategies"
    result: Dict[str, int] = {}
    if not strategies_dir.is_dir():
        return result
    for yaml_path in strategies_dir.glob("*.yaml"):
        try:
            with open(yaml_path, "r", encoding="utf-8") as f:
                cfg = _yaml.safe_load(f) or {}
            name = cfg.get("name", "")
            min_score = cfg.get("min_score", 0)
            effective = min_score if min_score > 0 else global_min
            if name:
                result[name] = effective
        except Exception:
            pass
    return result


def _compute_cost_breakdown(trade: dict, z_rates: dict) -> Dict[str, float]:
    """Compute itemised Zerodha cost breakdown for a round-trip intraday trade.

    Private fallback for trades where cost_brokerage is NULL but charges total exists.
    """
    exit_price = trade.get("exit_price")
    if not exit_price or not trade.get("charges"):
        return {}

    qty = trade.get("qty_filled") or trade.get("qty_planned") or 0
    entry_price = trade.get("entry_actual_price") or trade.get("entry_target_price") or 0
    if not qty or not entry_price:
        return {}

    direction = (trade.get("direction") or "LONG").upper()
    entry_side = "BUY" if direction == "LONG" else "SELL"
    exit_side  = "SELL" if direction == "LONG" else "BUY"

    flat  = z_rates.get("brokerage_flat_intraday", 20.0)
    b_pct = z_rates.get("brokerage_pct_intraday", 0.03) / 100
    stt_pct   = z_rates.get("stt_sell_pct", 0.025) / 100
    exch_pct  = z_rates.get("exchange_txn_pct", 0.00297) / 100
    sebi_pct  = z_rates.get("sebi_pct", 0.0001) / 100
    gst_pct   = z_rates.get("gst_pct", 18.0) / 100
    stamp_pct = z_rates.get("stamp_duty_mis_buy_pct", 0.003) / 100

    def _leg(side: str, price: float):
        tv    = qty * price
        brok  = min(flat, b_pct * tv)
        stt   = stt_pct * tv if side == "SELL" else 0.0
        exch  = exch_pct * tv + sebi_pct * tv
        gst   = gst_pct * (brok + exch)
        stamp = stamp_pct * tv if side == "BUY" else 0.0
        return brok, stt, exch, gst, stamp

    e = _leg(entry_side, float(entry_price))
    x = _leg(exit_side, float(exit_price))

    return {
        "brokerage": round(e[0] + x[0], 2),
        "stt":       round(e[1] + x[1], 2),
        "exch":      round(e[2] + x[2], 2),
        "gst":       round(e[3] + x[3], 2),
        "stamp":     round(e[4] + x[4], 2),
    }


def _load_candle_csv(candle_dir: Path, date_iso: str) -> Dict[Tuple[str, str], dict]:
    """Load candle CSV into dict keyed by (symbol, HH:MM).

    Private fallback for when candles table is empty.
    """
    csv_path = candle_dir / f"candle_data_{date_iso}.csv"
    if not csv_path.exists():
        return {}
    import csv as _csv
    candles: Dict[Tuple[str, str], dict] = {}
    with open(csv_path, newline="", encoding="utf-8") as f:
        for row in _csv.DictReader(f):
            hhmm = row["datetime"][11:16]
            candles[(row["symbol"], hhmm)] = {
                "open":  float(row["open"]),
                "high":  float(row["high"]),
                "low":   float(row["low"]),
                "close": float(row["close"]),
                "is_synthetic": 0,
            }
    return candles


# ─────────────────────────────────────────────────────────────────────────────
# Sheet 0: EOD Dashboard
# ─────────────────────────────────────────────────────────────────────────────

def build_sheet_0_dashboard(wb: openpyxl.Workbook, data: ReportData) -> Worksheet:
    """Build EOD Dashboard sheet with 5 sections."""
    ws = wb.create_sheet(title="0_EOD_Dashboard")
    _disable_gridlines(ws)

    row = 1

    ws.cell(row=row, column=1, value="EOD Dashboard")
    ws.cell(row=row, column=1).font = FONT_TITLE
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    row += 2

    def section_header(text: str, r: int) -> int:
        ws.cell(row=r, column=1, value=text)
        ws.cell(row=r, column=1).font = FONT_SECTION_HEADER
        ws.cell(row=r, column=1).fill = FILL_HEADER
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        return r + 1

    def add_row(label: str, value: Any, r: int) -> int:
        ws.cell(row=r, column=1, value=label).font = FONT_BODY
        ws.cell(row=r, column=2, value=value).font = FONT_BODY
        ws.cell(row=r, column=1).border = BORDER_ALL
        ws.cell(row=r, column=2).border = BORDER_ALL
        return r + 1

    if not data.signals and not data.trades and data.mode == "UNKNOWN":
        ws.cell(row=row, column=1, value="WARNING: No data found for this date -- verify DB connection and environment")
        ws.cell(row=row, column=1).font = FONT_WHITE_BOLD
        ws.cell(row=row, column=1).fill = FILL_RED
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        row += 1
        log.warning("daily_report: no data found for %s -- DB may be empty or wrong environment", data.date_iso)

    row = section_header("Section A — Day Overview", row)
    row = add_row("Trading Date", data.date_iso, row)
    row = add_row("Mode", data.mode, row)
    row = add_row("Account", data.account, row)
    row = add_row("Opening Capital", f"₹{data.opening_capital:,.2f}", row)
    # Closing = opening + realized net P&L (closed trades). There is NO per-day broker-funds
    # figure persisted (the boot get_margins().net is never snapshotted), so the old "Closing
    # Capital (Broker)" row was the last fm_ledger balance_after — the RESET_PNL 0.0 on a normal
    # day, a value it never actually had. Report the honest system-side closing instead.
    realized_pnl = sum((t.get("net_pnl") or 0.0) for t in data.trades if t.get("status") in _CLOSED_STATUSES)
    row = add_row("Closing Capital", f"₹{data.opening_capital + realized_pnl:,.2f}", row)
    row = add_row("Net P&L (Realized)", f"₹{realized_pnl:,.2f}", row)

    start_event = next((e for e in data.system_events if e.get("event_type") == "STARTUP"), None)
    end_event = next((e for e in reversed(data.system_events) if e.get("event_type") == "SHUTDOWN"), None)
    market_hours = "09:15 - 15:30"
    if start_event and end_event:
        try:
            s = datetime.fromisoformat(start_event["timestamp"]).replace(tzinfo=None)
            e = datetime.fromisoformat(end_event["timestamp"]).replace(tzinfo=None)
            uptime = (e - s).total_seconds() / 3600
            row = add_row("System Uptime", f"{uptime:.1f} hours", row)
        except (ValueError, KeyError):
            row = add_row("System Uptime", "N/A", row)
    else:
        row = add_row("System Uptime", "N/A", row)

    row = add_row("Market Hours Traded", market_hours, row)
    row = add_row("Excluded Symbols", ", ".join(data.excluded_symbols) or "None", row)
    row += 1

    row = section_header("Section B — Signal Funnel", row)
    total_signals = len(data.signals)
    # Classification is driven by the STRUCTURED status, never by the free-text
    # rejection_reason (see reports/signal_status.py for why, and what it cost).
    dedup_signals = [s for s in data.signals if not sig_status.is_dedup_duplicate(s.get("status"))]
    excluded_count = sum(1 for s in data.signals if s.get("symbol") in data.excluded_symbols)
    after_dedup = len(dedup_signals) - excluded_count
    passed_screen = sum(1 for s in data.signals if s.get("trade_id") is not None or s.get("status") in ("TRADED", "PLACED", "FILLED"))
    converted = sum(1 for s in data.signals if s.get("trade_id") is not None)
    # Sizing/capital rejections, by status. The old test was
    # `"CAPITAL" in rejection_reason.upper()`, which matched every CONCENTRATION rejection
    # because its reason string embeds `capital_qty=` -- reporting 1,189 capital rejections
    # on 2026-07-10 when the true count was 0.
    sizing_rejected = [s for s in data.signals if sig_status.is_sizing_rejection(s.get("status"))]
    constraint_counts: dict = {}
    for s in sizing_rejected:
        c = sig_status.sizing_constraint(s.get("status")) or "UNKNOWN"
        constraint_counts[c] = constraint_counts.get(c, 0) + 1
    top_constraint = max(constraint_counts.items(), key=lambda kv: kv[1]) if constraint_counts else None
    # 'Silent' means NO recorded outcome. REJECTED_* is not the only terminal family --
    # DROPPED_*/SKIPPED_*/QUEUE_FULL/PLACEMENT_FAILED/TIMEOUT are terminal too, and counting
    # only REJECTED_* mislabelled 2,688 explicitly-dispositioned signals as silently dead.
    silent_dead = total_signals - sum(
        1 for s in data.signals
        if s.get("trade_id") is not None or sig_status.has_explicit_disposition(s.get("status"))
    )

    row = add_row("Total Received", f"{total_signals}", row)
    row = add_row("After Dedup/Excluded", f"{after_dedup} ({after_dedup/max(total_signals,1)*100:.1f}%)", row)
    row = add_row("Passed Screening", f"{passed_screen} ({passed_screen/max(total_signals,1)*100:.1f}%)", row)
    row = add_row("Converted to Orders", f"{converted} ({converted/max(total_signals,1)*100:.1f}%)", row)
    row = add_row("Rejected (Sizing/Capital)", f"{len(sizing_rejected)}", row)
    if top_constraint:
        row = add_row("  Binding Constraint", f"{top_constraint[0]} ({top_constraint[1]})", row)
    row = add_row("Silent Dead", f"{max(0, silent_dead)}", row)
    row += 1

    row = section_header("Section C — P&L Summary", row)
    closed_trades = [t for t in data.trades if t.get("status") in _CLOSED_STATUSES]
    gross_pnl = sum(t.get("gross_pnl") or 0.0 for t in closed_trades)
    total_costs = sum(t.get("charges") or 0.0 for t in closed_trades)
    net_pnl = sum(t.get("net_pnl") or 0.0 for t in closed_trades)
    wins = [t for t in closed_trades if (t.get("net_pnl") or 0.0) > 0]
    losses = [t for t in closed_trades if (t.get("net_pnl") or 0.0) < 0]
    breakeven = [t for t in closed_trades if (t.get("net_pnl") or 0.0) == 0]
    win_rate = len(wins) / max(len(closed_trades), 1) * 100

    best_trade = max(closed_trades, key=lambda t: t.get("net_pnl") or 0.0, default=None)
    worst_trade = min(closed_trades, key=lambda t: t.get("net_pnl") or 0.0, default=None)

    row = add_row("Gross P&L", f"₹{gross_pnl:,.2f}", row)
    row = add_row("Total Costs", f"₹{total_costs:,.2f}", row)
    row = add_row("Net P&L", f"₹{net_pnl:,.2f}", row)
    ws.cell(row=row-1, column=2).fill = FILL_GREEN if net_pnl >= 0 else FILL_RED
    row = add_row("W/L/BE", f"{len(wins)}/{len(losses)}/{len(breakeven)}", row)
    row = add_row("Win Rate", f"{win_rate:.1f}%", row)

    if best_trade:
        row = add_row("Best Trade", f"{best_trade['symbol']} Rs{best_trade.get('net_pnl', 0):.2f} ({best_trade.get('exit_reason', '')})", row)
    if worst_trade:
        row = add_row("Worst Trade", f"{worst_trade['symbol']} Rs{worst_trade.get('net_pnl', 0):.2f} ({worst_trade.get('exit_reason', '')})", row)

    rr_values = []
    for t in closed_trades:
        entry_p = t.get("entry_actual_price") or t.get("entry_target_price") or 0.0
        sl_p = t.get("sl_initial") or 0.0
        exit_p = t.get("exit_price") or 0.0
        direction = t.get("direction", "LONG")
        if entry_p > 0 and sl_p > 0 and entry_p != sl_p:
            risk = abs(entry_p - sl_p)
            if direction == "LONG":
                reward = exit_p - entry_p
            else:
                reward = entry_p - exit_p
            rr_values.append(reward / risk)
    avg_rr = sum(rr_values) / len(rr_values) if rr_values else 0.0
    row = add_row("Avg R:R Achieved", f"{avg_rr:.2f}", row)

    if data.opening_capital > 0:
        util = sum(t.get("margin_reserved") or 0.0 for t in data.trades) / data.opening_capital * 100
        row = add_row("Capital Utilization %", f"{util:.1f}%", row)

    if closed_trades and data.opening_capital > 0:
        running_pnl = 0.0
        max_dd = 0.0
        peak = 0.0
        for t in closed_trades:
            running_pnl += t.get("net_pnl") or 0.0
            peak = max(peak, running_pnl)
            dd = peak - running_pnl
            max_dd = max(max_dd, dd)
        dd_pct = max_dd / data.opening_capital * 100
        row = add_row("Max Drawdown %", f"{dd_pct:.2f}%", row)
    else:
        row = add_row("Max Drawdown %", "0.00%", row)
    row += 1

    row = section_header("Section D — System Health", row)
    error_count = sum(1 for e in data.system_events if "ERROR" in (e.get("event_type") or "").upper())
    critical_count = sum(1 for e in data.system_events if _is_critical_event(e.get("event_type")))
    orphan_count = sum(1 for r in data.recon_log if "ORPHAN" in (r.get("check_name") or "").upper())

    kill_count = sum(1 for e in data.system_events if "KILL" in (e.get("event_type") or "").upper())

    row = add_row("ERROR Count", error_count, row)
    row = add_row("CRITICAL Count", critical_count, row)
    row = add_row("Kill Switch Events", kill_count, row)
    row = add_row("Orphan Orders", orphan_count, row)
    row = add_row("Reconcile Status", "OK" if orphan_count == 0 else "REVIEW", row)

    # Group by the STRUCTURED status family, not by the free-text reason. The reason embeds the
    # symbol and three sizing arm values, so grouping by it produced 190 near-unique lines for
    # 7,655 rejections on 2026-07-10 -- every one too small to notice. By family: 9 lines.
    # The per-signal reason text is NOT lost: sheet 1_Signals renders it for every signal.
    rejected_signals = [s for s in data.signals if sig_status.is_rejected(s.get("status"))]
    rejection_reasons: dict = {}
    for s in rejected_signals:
        fam = sig_status.family(s.get("status"))
        rejection_reasons[fam] = rejection_reasons.get(fam, 0) + 1
    if rejection_reasons:
        row = add_row("Rejection Breakdown", "", row)
        for reason, cnt in sorted(rejection_reasons.items(), key=lambda x: -x[1]):
            row = add_row(f"  {reason}", cnt, row)
    row += 1

    row = section_header("Section E — Auto Tuning Signals", row)
    suggestions = _generate_tune_suggestions(data)
    if not suggestions:
        row = add_row("Status", "No tuning suggestions for today", row)
    else:
        for suggestion in suggestions[:10]:
            ws.cell(row=row, column=1, value=suggestion).font = FONT_BODY
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            row += 1

    _set_column_widths(ws, {"A": 25, "B": 50, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 15})

    return ws


def _generate_tune_suggestions(data: ReportData) -> List[str]:
    """Generate auto-tuning suggestions based on trade outcomes."""
    suggestions = []
    closed_trades = [t for t in data.trades if t.get("status") in _CLOSED_STATUSES]

    for t in closed_trades:
        symbol = t.get("symbol", "?")
        exit_reason = t.get("exit_reason", "")
        exit_price = t.get("exit_price") or 0.0
        sl_initial = t.get("sl_initial") or 0.0
        tgt_initial = t.get("tgt_initial") or 0.0
        entry_actual = t.get("entry_actual_price") or t.get("entry_target_price") or 0.0
        entry_target = t.get("entry_target_price") or 0.0
        direction = t.get("direction", "LONG")

        if exit_reason in ("SL_HIT", "SL"):
            if direction == "LONG" and sl_initial > 0:
                adverse_move = (sl_initial - exit_price) / sl_initial * 100 if sl_initial else 0
                if adverse_move > 0.3:
                    suggestions.append(f"⚠️ {symbol}: SL hit at ₹{exit_price:.2f} vs placed ₹{sl_initial:.2f} — consider widening SL by 0.5%")
            elif direction == "SHORT" and sl_initial > 0:
                adverse_move = (exit_price - sl_initial) / sl_initial * 100 if sl_initial else 0
                if adverse_move > 0.3:
                    suggestions.append(f"⚠️ {symbol}: SL hit at ₹{exit_price:.2f} vs placed ₹{sl_initial:.2f} — consider widening SL by 0.5%")

        elif exit_reason == "EOD":
            if direction == "LONG":
                max_favorable = exit_price
                for ing in data.innings:
                    if ing.get("trade_id") == t.get("trade_id"):
                        max_favorable = max(max_favorable, ing.get("exit_price") or exit_price)
                if tgt_initial > 0:
                    gap_pct = (tgt_initial - max_favorable) / tgt_initial * 100
                    if 0 < gap_pct <= 1.0:
                        diff = tgt_initial - max_favorable
                        suggestions.append(f"ℹ️ {symbol}: TGT not reached — max favourable ₹{max_favorable:.2f} (missed by ₹{diff:.2f}). Consider reducing TGT by 1%")

        elif exit_reason in ("TGT_HIT", "TGT"):
            suggestions.append(f"✅ {symbol}: TGT hit perfectly. Strategy {t.get('strategy', '?')} performing as expected")

        if entry_actual > 0 and entry_target > 0:
            slip_pct = abs(entry_actual - entry_target) / entry_target * 100
            if slip_pct > 2.0:
                suggestions.append(f"🚨 CRITICAL: {symbol}: Entry slippage {slip_pct:.1f}% — R:R collapsed. Review entry gate")

    if closed_trades:
        w = len([t for t in closed_trades if (t.get("net_pnl") or 0) > 0])
        l = len(closed_trades) - w
        suggestions.append(f"ℹ️ GENERAL: {w}W {l}L today — review complete.")

    return suggestions


# ─────────────────────────────────────────────────────────────────────────────
# Sheet 1: Signals
# ─────────────────────────────────────────────────────────────────────────────

def build_sheet_1_signals(wb: openpyxl.Workbook, data: ReportData) -> Worksheet:
    """Build Signals sheet with recon columns."""
    ws = wb.create_sheet(title="1_Signals")
    _disable_gridlines(ws)

    headers = [
        "Trading Date", "Received At", "Scanner/Strategy", "Symbol", "Raw Symbol",
        "Total Rcvd", "Queued", "Selected", "Rejected", "Duplicate", "Order Passed", "Excluded", "Delta",
        "Eligible Score\n(Min Tradable)", "Algo Score\n(Stock's Score)", "Score Breakdown",
        "Trigger Price", "Signal Age (s)", "Dedup Status", "Screening Result", "Rejection Reason",
        "Signal ID", "Trade ID"
    ]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_ALL

    ws.freeze_panes = "A2"

    excluded_set = set(data.excluded_symbols)

    screener_map = {r.get("signal_id"): r for r in data.screener_results}
    trade_by_signal = {t.get("signal_id"): t.get("trade_id") for t in data.trades}

    row = 2
    for sig in data.signals:
        signal_id = sig.get("signal_id", "")
        symbol = sig.get("symbol", "")
        status = sig.get("status", "")
        rejection_reason = sig.get("rejection_reason", "")
        trade_id = sig.get("trade_id") or trade_by_signal.get(signal_id) or "—"

        is_dup = "DUPLICATE" in status.upper()
        is_excluded = symbol in excluded_set
        is_queued = not is_dup and not is_excluded
        is_selected = trade_id != "—" or status in ("TRADED", "PLACED", "FILLED", "SIZED", "APPROVED", "RESERVED")
        is_rejected = "REJECTED" in status.upper() and not is_selected
        is_order_passed = trade_id != "—"

        screener_row = screener_map.get(signal_id, {})
        algo_score = screener_row.get("score") if screener_row.get("score") is not None else 0
        step_results_raw = screener_row.get("step_results", "")
        score_breakdown = ""
        if step_results_raw:
            try:
                steps = json.loads(step_results_raw) if isinstance(step_results_raw, str) else step_results_raw
                if isinstance(steps, dict):
                    score_breakdown = " ".join(f"{k}:{v}" for k, v in steps.items() if isinstance(v, (int, float)))
            except (json.JSONDecodeError, TypeError):
                pass
        strategy_name = sig.get("strategy", "")
        eligible_from_db = screener_row.get("eligible_score")
        if eligible_from_db is not None:
            effective_min = eligible_from_db
        else:
            effective_min = data.strategy_min_scores.get(strategy_name, "—")

        try:
            received_dt = datetime.fromisoformat(sig.get("received_at", "")).replace(tzinfo=None)
            triggered_dt = datetime.fromisoformat(sig.get("triggered_at", "")).replace(tzinfo=None)
            signal_age = int((received_dt - triggered_dt).total_seconds())
        except (ValueError, TypeError):
            signal_age = "—"

        row_data = [
            data.date_iso,
            _fmt_time(sig.get("received_at")),
            f"{sig.get('scanner', '')} / {sig.get('strategy', '')}",
            symbol,
            symbol,
            1,
            1 if is_queued else 0,
            1 if is_selected else 0,
            1 if is_rejected else 0,
            1 if is_dup else 0,
            1 if is_order_passed else 0,
            1 if is_excluded else 0,
            f"=F{row}-(G{row}+J{row}+L{row})",
            effective_min,
            algo_score,
            score_breakdown or "---",
            sig.get("trigger_price", ""),
            signal_age,
            "DUP" if is_dup else "OK",
            screener_row.get("status", status),
            rejection_reason or "—",
            signal_id,
            trade_id,
        ]

        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = FONT_BODY
            cell.border = BORDER_ALL

            if col == 15 and isinstance(algo_score, (int, float)) and isinstance(effective_min, (int, float)):
                if algo_score < effective_min:
                    cell.fill = FILL_RED
                else:
                    cell.fill = FILL_GREEN
            elif col == 16:
                cell.alignment = ALIGN_LEFT

        row += 1

    if data.signals:
        ws.cell(row=row, column=1, value="TOTALS").font = FONT_HEADER
        for col in range(6, 13):
            cell = ws.cell(row=row, column=col, value=f"=SUM({get_column_letter(col)}2:{get_column_letter(col)}{row-1})")
            cell.font = FONT_HEADER
            cell.border = BORDER_ALL

        row += 1
        ws.cell(row=row, column=1, value="Recon check: F = G + J + L | G = H + I | K <= H | Delta != 0 -> investigate")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=23)

    return ws


# ─────────────────────────────────────────────────────────────────────────────
# Sheet 2: Orders
# ─────────────────────────────────────────────────────────────────────────────

def build_sheet_2_orders(wb: openpyxl.Workbook, data: ReportData) -> Worksheet:
    """Build Orders sheet matching order_sheet.xlsx design (43 cols, 3 header rows)."""
    ws = wb.create_sheet(title="2_Orders")
    _disable_gridlines(ws)

    # Separator columns (narrow dividers painted with separator blue)
    sep_cols = [8, 13, 20, 23, 29, 37, 40]

    # ── Row 1: Title ─────────────────────────────────────────────────────────
    title_cell = ws.cell(row=1, column=1,
        value=f"AlgoCore Systems v2.0  ·  Orders & Trade Lifecycle  ·  {data.date_iso}  ({data.mode})")
    title_cell.font = FONT_TITLE
    title_cell.fill = FILL_TITLE_BG
    title_cell.alignment = ALIGN_LEFT

    # ── Row 2: Group headers ─────────────────────────────────────────────────
    group_headers = [
        ("IDENTITY", 1, 7),
        ("TIME", 9, 12),
        ("SYSTEM  (Per Qty Price)", 14, 19),
        ("ACTUALLY FILLED  (Per Qty Price)", 21, 22),
        ("ENTRY DEVIATION  &  RESULTS", 24, 28),
        ("REVENUE  &  EXPENSES  (Total Amount)", 30, 36),
        ("NET REVENUE  (Total Amount)", 38, 39),
        ("METADATA IDs", 41, 43),
    ]

    for label, start_col, end_col in group_headers:
        ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=end_col)
        cell = ws.cell(row=2, column=start_col, value=label)
        cell.font = FONT_WHITE_BOLD
        cell.fill = FILL_GROUP_HEADER
        cell.alignment = ALIGN_CENTER

    for col in sep_cols:
        ws.cell(row=2, column=col).fill = FILL_SEPARATOR

    ws.row_dimensions[2].height = 21.75

    # ── Row 3: Sub-headers ───────────────────────────────────────────────────
    sub_headers = [
        # IDENTITY (1-7)
        "Trading\nDate", "Signal\nTime", "Strategy", "Direction", "Symbol",
        "Eligible\nScore\n(Min)", "Algo\nScore",
        "",   # 8 sep
        # TIME (9-12)
        "Order\nPlaced", "Fill\nTime", "Exit\nTime", "Time in\nTrade (min)",
        "",   # 13 sep
        # SYSTEM (14-19)
        "Sys\nQty", "Sys Entry\n(₹/scrip)", "Sys SL\n(₹/scrip)", "Sys TGT\n(₹/scrip)",
        "Sys\nR:R", "Qty\nMatch",
        "",   # 20 sep
        # ACTUALLY FILLED (21-22)
        "Fill Entry\n(₹/scrip)", "Fill SL\n(₹/scrip)",
        "",   # 23 sep
        # ENTRY DEVIATION & RESULTS (24-28)
        "Entry Slip\n(₹/scrip)", "Entry\nSlip %", "Fill R:R", "Exit\nReason", "SL Trail\nCount",
        "",   # 29 sep
        # REVENUE & EXPENSES (30-36)
        "Gross P&L\n(₹)", "Brokerage\n(₹)", "STT\n(₹)", "Exch Charges\n(₹)",
        "Stamp Duty\n(₹)", "GST\n(₹)", "Total Costs\n(₹)",
        "",   # 37 sep
        # NET REVENUE (38-39)
        "Net P&L\n(₹)", "ROI %",
        "",   # 40 sep
        # METADATA (41-43)
        "Trade ID", "Broker\nOrder ID", "Exit Price\n(₹)",
    ]

    for col, header in enumerate(sub_headers, start=1):
        cell = ws.cell(row=3, column=col, value=header if header else None)
        if col in sep_cols:
            cell.fill = FILL_SEPARATOR
        else:
            cell.font = FONT_WHITE_BOLD
            cell.fill = FILL_SUB_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_ALL

    ws.row_dimensions[3].height = 39.75

    # Pre-paint separator columns for all data rows
    max_rows = max(500, len(data.trades) + 20)
    for col in sep_cols:
        for r in range(4, 4 + max_rows):
            ws.cell(row=r, column=col).fill = FILL_SEPARATOR
        ws.column_dimensions[get_column_letter(col)].width = 2

    ws.freeze_panes = "A4"

    # ── Data rows (start at row 4) ───────────────────────────────────────────
    row = 4
    screener_map = {r.get("signal_id"): r for r in data.screener_results}
    signal_map = {s.get("signal_id"): s for s in data.signals}

    for trade in data.trades:
        trade_id = trade.get("trade_id", "")
        signal_id = trade.get("signal_id", "")
        signal = signal_map.get(signal_id, {})
        screener = screener_map.get(signal_id, {})

        strategy = trade.get("strategy", "") or signal.get("strategy", "")
        eligible_from_db = screener.get("eligible_score")
        if eligible_from_db is not None:
            effective_min = eligible_from_db
        else:
            effective_min = data.strategy_min_scores.get(strategy, "—")

        entry_order = _get_order_for_trade_leg(data.orders, trade_id, "ENTRY")
        sl_order = _get_order_for_trade_leg(data.orders, trade_id, "SL")

        sys_qty = trade.get("qty_planned", 0)
        sys_entry = trade.get("entry_target_price", 0)
        sys_sl = trade.get("sl_initial", 0)
        sys_tgt = trade.get("tgt_initial", 0)

        fill_entry = trade.get("entry_actual_price") or sys_entry
        fill_sl = sl_order.get("trigger_price") if sl_order else sys_sl

        if sys_entry and sys_sl and sys_entry != sys_sl:
            risk = abs(sys_entry - sys_sl)
            reward = abs(sys_tgt - sys_entry) if sys_tgt else 0
            sys_rr = reward / risk if risk else 0
        else:
            sys_rr = 0

        fill_rr = 0
        if fill_entry and fill_sl and fill_entry != fill_sl:
            risk = abs(fill_entry - fill_sl)
            reward = abs(sys_tgt - fill_entry) if sys_tgt else 0
            fill_rr = reward / risk if risk else 0

        entry_slip = abs(fill_entry - sys_entry) if fill_entry and sys_entry else 0
        entry_slip_pct = _calc_slip_pct(sys_entry, fill_entry) if sys_entry else 0

        time_in_trade = 0
        if trade.get("entry_time") and trade.get("exit_time"):
            try:
                entry_dt = datetime.fromisoformat(trade["entry_time"]).replace(tzinfo=None)
                exit_dt = datetime.fromisoformat(trade["exit_time"]).replace(tzinfo=None)
                time_in_trade = max(0, int((exit_dt - entry_dt).total_seconds() / 60))
            except ValueError:
                pass

        gross_pnl = trade.get("gross_pnl") or 0
        charges = trade.get("charges") or 0
        net_pnl = trade.get("net_pnl") or 0

        qty_filled = trade.get("qty_filled", 0)
        roi_pct = (net_pnl / (fill_entry * qty_filled) * 100) if fill_entry and qty_filled else 0

        is_closed = trade.get("status") in _CLOSED_STATUSES
        cost_brokerage_db = trade.get("cost_brokerage")
        charges_total = trade.get("charges")
        if not is_closed or charges_total is None:
            brokerage = "—"
            stt       = "—"
            exch_chrg = "—"
            stamp     = "—"
            gst       = "—"
        elif cost_brokerage_db is not None:
            brokerage = cost_brokerage_db
            stt       = trade.get("cost_stt") or 0
            exch_chrg = trade.get("cost_exchange_txn") or 0
            stamp     = trade.get("cost_stamp_duty") or 0
            gst       = trade.get("cost_gst") or 0
        elif charges_total > 0:
            cost_bd = _compute_cost_breakdown(trade, data.broker_rates)
            brokerage = cost_bd.get("brokerage", 0) if cost_bd else 0
            stt       = cost_bd.get("stt", 0) if cost_bd else 0
            exch_chrg = cost_bd.get("exch", 0) if cost_bd else 0
            stamp     = cost_bd.get("stamp", 0) if cost_bd else 0
            gst       = cost_bd.get("gst", 0) if cost_bd else 0
        else:
            brokerage = 0
            stt       = 0
            exch_chrg = 0
            stamp     = 0
            gst       = 0

        trail_count = trade.get("sl_trail_count") or 0

        row_data = [
            data.date_iso,                                                   # 1
            _fmt_time(signal.get("received_at")),                           # 2
            strategy,                                                        # 3
            trade.get("direction", ""),                                      # 4
            trade.get("symbol", ""),                                         # 5
            effective_min,                                                   # 6
            screener.get("score") if screener.get("score") is not None else 0,  # 7
            "",                                                              # 8 sep
            _fmt_time(entry_order.get("placed_at")) if entry_order else "",  # 9
            _fmt_time(trade.get("entry_time")),                             # 10
            _fmt_time(trade.get("exit_time")),                              # 11
            time_in_trade if time_in_trade else ("—" if not trade.get("exit_time") else 0),  # 12
            "",                                                              # 13 sep
            sys_qty,                                                         # 14
            sys_entry,                                                       # 15
            sys_sl,                                                          # 16
            sys_tgt,                                                         # 17
            round(sys_rr, 2),                                                # 18
            "Y" if trade.get("qty_filled") == sys_qty else "N",              # 19
            "",                                                              # 20 sep
            fill_entry,                                                      # 21
            fill_sl,                                                         # 22
            "",                                                              # 23 sep
            round(entry_slip, 2),                                            # 24
            f"{entry_slip_pct:.2f}%",                                       # 25
            round(fill_rr, 2),                                               # 26
            trade.get("exit_reason", "") or "—",                             # 27
            trail_count,                                                     # 28
            "",                                                              # 29 sep
            round(gross_pnl, 2),                                             # 30
            brokerage,                                                       # 31 Brokerage
            stt,                                                             # 32 STT
            exch_chrg,                                                       # 33 Exch Charges
            stamp,                                                           # 34 Stamp Duty
            gst,                                                             # 35 GST
            round(charges, 2) if is_closed and charges_total is not None else "—",  # 36 Total Costs
            "",                                                              # 37 sep
            round(net_pnl, 2),                                               # 38
            f"{roi_pct:.2f}%",                                              # 39
            "",                                                              # 40 sep
            trade_id,                                                        # 41
            entry_order.get("order_id") if entry_order else "",              # 42
            round(trade["exit_price"], 2) if trade.get("exit_price") else "—",  # 43
        ]

        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = FONT_BODY
            if col not in sep_cols:
                cell.border = BORDER_ALL

            if col == 27:
                exit_reason = trade.get("exit_reason", "")
                if exit_reason in EXIT_REASON_FILLS:
                    cell.fill = EXIT_REASON_FILLS[exit_reason]

            if col == 38:
                if net_pnl > 0:
                    cell.fill = FILL_GREEN
                elif net_pnl < 0:
                    cell.fill = FILL_RED

        row += 1

    return ws


# ─────────────────────────────────────────────────────────────────────────────
# Sheet 3: Capital
# ─────────────────────────────────────────────────────────────────────────────

def build_sheet_3_capital(wb: openpyxl.Workbook, data: ReportData) -> Worksheet:
    """Build Capital (Funds Tracker) sheet with reconciliation."""
    ws = wb.create_sheet(title="3_Capital")
    _disable_gridlines(ws)

    headers = [
        "Trading Date", "Time of Order", "Sl.No", "Strategy", "Stock", "Direction",
        "Trade", "Position Value ₹", "Margin Blocked ₹", "SL Risk ₹", "Target Profit ₹",
        "Result", "P&L ₹\n(Net of Costs)", "Additional Funds Added ₹", "Funds After Adjustment ₹"
    ]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_ALL

    ws.freeze_panes = "A2"

    row = 2
    ws.cell(row=row, column=1, value=data.date_iso).border = BORDER_ALL
    ws.cell(row=row, column=2, value="08:10:00").border = BORDER_ALL
    ws.cell(row=row, column=3, value="Opening").font = FONT_HEADER
    ws.cell(row=row, column=3).border = BORDER_ALL
    ws.cell(row=row, column=7, value="Initial Capital").border = BORDER_ALL
    ws.cell(row=row, column=15, value=data.opening_capital).border = BORDER_ALL
    ws.cell(row=row, column=15).number_format = NUM_FMT_CURRENCY
    for col in range(1, 16):
        ws.cell(row=row, column=col).fill = FILL_GREY
    row += 1

    running_balance = data.opening_capital

    signal_map = {s.get("signal_id"): s for s in data.signals}
    order_map = {}
    for o in data.orders:
        order_map.setdefault(o.get("trade_id"), []).append(o)

    sl_no = 1
    for trade in sorted(data.trades, key=lambda t: t.get("created_at", "")):
        trade_id = trade.get("trade_id", "")
        signal_id = trade.get("signal_id", "")
        signal = signal_map.get(signal_id, {})

        strategy = trade.get("strategy", "") or signal.get("strategy") or "UNKNOWN"

        entry_order = _get_order_for_trade_leg(data.orders, trade_id, "ENTRY")

        order_type_display = "LIMIT"
        if entry_order:
            product = entry_order.get("product", "")
            variety = entry_order.get("variety", "")
            order_type = entry_order.get("order_type", "")
            direction = trade.get("direction", "LONG")

            if variety == "co" or product == "CO":
                order_type_display = f"{'BUY' if direction == 'LONG' else 'SELL'} CO"
            elif product == "MIS":
                order_type_display = f"{'BUY' if direction == 'LONG' else 'SELL'} {order_type}"
            else:
                order_type_display = f"{'BUY' if direction == 'LONG' else 'SELL'} {order_type}"

        entry_price = trade.get("entry_actual_price") or trade.get("entry_target_price") or 0
        qty = trade.get("qty_filled") or trade.get("qty_planned") or 0
        position_value = entry_price * qty
        margin = trade.get("margin_reserved", 0)

        sl_price = trade.get("sl_initial", 0)
        tgt_price = trade.get("tgt_initial", 0)
        direction = trade.get("direction", "LONG")

        if direction == "LONG":
            sl_risk = (entry_price - sl_price) * qty if entry_price and sl_price else 0
            tgt_profit = (tgt_price - entry_price) * qty if tgt_price and entry_price else 0
        else:
            sl_risk = (sl_price - entry_price) * qty if entry_price and sl_price else 0
            tgt_profit = (entry_price - tgt_price) * qty if tgt_price and entry_price else 0

        exit_reason = trade.get("exit_reason", "")
        result = exit_reason if exit_reason else trade.get("status", "")

        net_pnl = trade.get("net_pnl") or 0

        sl_display = round(abs(sl_risk), 2) if sl_price else "—"
        tgt_display = round(abs(tgt_profit), 2) if tgt_price else "—"

        running_balance += net_pnl

        row_data = [
            data.date_iso,
            _fmt_time(trade.get("created_at")),
            sl_no,
            strategy,  # Fixed: use strategy variable with fallback
            trade.get("symbol", ""),
            direction,
            order_type_display,
            round(position_value, 2),
            round(margin, 2),
            sl_display,
            tgt_display,
            result,
            round(net_pnl, 2),
            "",
            round(running_balance, 2),
        ]

        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = FONT_BODY
            cell.border = BORDER_ALL

            if col in (8, 9, 10, 11, 13, 15):
                cell.number_format = NUM_FMT_CURRENCY

            if col == 13:
                if net_pnl > 0:
                    cell.fill = FILL_GREEN
                elif net_pnl < 0:
                    cell.fill = FILL_RED

        sl_no += 1
        row += 1

    if data.trades:
        ws.cell(row=row, column=1, value=data.date_iso).border = BORDER_ALL
        ws.cell(row=row, column=2, value="CLOSING (15:17 IST · EOD)").border = BORDER_ALL
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)

        for col in [8, 9, 10, 11]:
            ws.cell(row=row, column=col, value=f"=SUM({get_column_letter(col)}3:{get_column_letter(col)}{row-1})").border = BORDER_ALL
            ws.cell(row=row, column=col).number_format = NUM_FMT_CURRENCY

        ws.cell(row=row, column=12, value="Net Result").border = BORDER_ALL
        ws.cell(row=row, column=13, value=f"=SUM(M3:M{row-1})").border = BORDER_ALL
        ws.cell(row=row, column=13).number_format = NUM_FMT_CURRENCY
        ws.cell(row=row, column=14, value=f"=SUM(N3:N{row-1})").border = BORDER_ALL
        ws.cell(row=row, column=15, value=running_balance).border = BORDER_ALL
        ws.cell(row=row, column=15).number_format = NUM_FMT_CURRENCY

        for col in range(1, 16):
            ws.cell(row=row, column=col).fill = FILL_GREY
            ws.cell(row=row, column=col).font = FONT_HEADER

        row += 4

    ws.cell(row=row, column=1, value="CAPITAL SUMMARY").font = FONT_HEADER
    row += 1

    # System-side capital only. The old "Broker" rows + Reconcile Status derived from the last
    # fm_ledger balance_after (the RESET_PNL 0.0 on a normal day), so the reconcile read REVIEW
    # on ~14 of 15 days against a broker figure that was never captured. The genuine ledger-vs-
    # trades capital reconciliation lives in daily_trade_review's "3 · CAPITAL" block.
    recon_rows = [
        ("Opening Capital ₹", data.opening_capital),
        ("Closing Capital ₹", running_balance),
        ("Net P&L (Realized) ₹", running_balance - data.opening_capital),
    ]

    for label, value in recon_rows:
        ws.cell(row=row, column=1, value=label).font = FONT_BODY
        ws.cell(row=row, column=1).border = BORDER_ALL
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = FONT_BODY
        cell.border = BORDER_ALL
        if isinstance(value, (int, float)):
            cell.number_format = NUM_FMT_CURRENCY
        row += 1

    _set_column_widths(ws, {"A": 12, "B": 15, "C": 8, "D": 15, "E": 15, "F": 10,
                           "G": 15, "H": 15, "I": 15, "J": 12, "K": 12, "L": 12,
                           "M": 15, "N": 18, "O": 20})

    return ws


# ─────────────────────────────────────────────────────────────────────────────
# Sheet 4: Candles
# ─────────────────────────────────────────────────────────────────────────────

def build_sheet_4_candles(
    wb: openpyxl.Workbook,
    data: ReportData,
) -> Worksheet:
    """Build Candles analysis sheet with tune suggestions."""
    ws = wb.create_sheet(title="4_Candles")
    _disable_gridlines(ws)

    group_headers = [
        ("IDENTITY", 1, 6),
        ("CANDLE ANALYSIS", 8, 12),
        ("ACTUAL TRADE System", 14, 17),
        ("DELTA/TUNING", 19, 21),
        ("⚡ TUNE SUGGESTION", 23, 23),
    ]

    for label, start_col, end_col in group_headers:
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        cell = ws.cell(row=1, column=start_col, value=label)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER

    sep_cols = [7, 13, 18, 22]
    for col in sep_cols:
        # Start from row 3 to avoid interfering with merged header rows 1-2
        for r in range(3, 100):
            ws.cell(row=r, column=col).fill = FILL_SEPARATOR
        ws.column_dimensions[get_column_letter(col)].width = 2

    sub_headers = [
        "Trading Date", "Trade ID", "Strategy", "Symbol", "Broker Order ID", "Entry Time",
        "",
        "Open", "High", "Low", "Close", "Synthetic?",
        "",
        "Our Entry", "Our SL", "Our TGT", "Matched?",
        "",
        "Max Favourable", "Max Adverse", "Missed Profit",
        "",
        "Tune Suggestion"
    ]

    for col, header in enumerate(sub_headers, start=1):
        if header:
            cell = ws.cell(row=2, column=col, value=header)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER if col not in sep_cols else FILL_SEPARATOR
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_ALL

    ws.freeze_panes = "A3"

    row = 3
    signal_map = {s.get("signal_id"): s for s in data.signals}

    # Filter to only CLOSED trades - exclude CANCELLED to avoid NaN values
    closed_trades = [t for t in data.trades if t.get("status") in _CLOSED_STATUSES]

    for trade in closed_trades:
        trade_id = trade.get("trade_id", "")
        signal_id = trade.get("signal_id", "")
        signal = signal_map.get(signal_id, {})
        strategy = trade.get("strategy", "") or signal.get("strategy", "")

        entry_order = _get_order_for_trade_leg(data.orders, trade_id, "ENTRY")

        our_entry = trade.get("entry_actual_price") or trade.get("entry_target_price") or 0
        our_sl = trade.get("sl_initial") or 0
        our_tgt = trade.get("tgt_initial") or 0
        exit_price = trade.get("exit_price") or 0
        exit_reason = trade.get("exit_reason", "")
        direction = trade.get("direction", "LONG")

        exc = data.excursion_map.get(trade_id, {})
        max_fav = exc.get("mfe_price") or exit_price
        max_adv = exc.get("mae_price") or exit_price

        if direction == "LONG":
            missed_profit = our_tgt - max_fav if our_tgt > max_fav else 0
        else:
            missed_profit = max_fav - our_tgt if our_tgt < max_fav else 0

        tune_suggestion = ""
        if exit_reason in ("SL_HIT", "SL"):
            if direction == "LONG":
                if our_sl and exit_price and (our_sl - exit_price) / our_sl > 0.003:
                    tune_suggestion = f"⚠️ SL hit at {exit_price:.2f} vs placed {our_sl:.2f} — consider +0.5% SL buffer"
            else:
                if our_sl and exit_price and (exit_price - our_sl) / our_sl > 0.003:
                    tune_suggestion = f"⚠️ SL hit at {exit_price:.2f} vs placed {our_sl:.2f} — consider +0.5% SL buffer"
        elif exit_reason == "EOD" and our_tgt:
            gap_pct = abs(our_tgt - max_fav) / our_tgt * 100 if our_tgt else 0
            if 0 < gap_pct <= 1.0:
                tune_suggestion = f"ℹ️ TGT {our_tgt:.2f} not reached — max favourable {max_fav:.2f} (missed ₹{missed_profit:.2f})"
        elif exit_reason in ("TGT_HIT", "TGT"):
            tune_suggestion = "✅ TGT hit perfectly — no tuning needed"

        entry_hhmm = _fmt_time(trade.get("entry_time"))[:5]
        exc_candle = exc
        candle = data.candle_map.get((trade.get("symbol", ""), entry_hhmm), {})

        row_data = [
            data.date_iso,
            trade_id[:8] + "..." if len(trade_id) > 8 else trade_id,
            strategy,
            trade.get("symbol", ""),
            entry_order.get("order_id", "") if entry_order else "",
            _fmt_time(trade.get("entry_time")),
            "",
            exc_candle.get("entry_candle_open") or candle.get("open") or "—",
            exc_candle.get("entry_candle_high") or candle.get("high") or "—",
            exc_candle.get("entry_candle_low") or candle.get("low") or "—",
            exc_candle.get("entry_candle_close") or candle.get("close") or "—",
            "Yes" if candle.get("is_synthetic") else ("No" if candle else "N/A"),
            "",
            round(our_entry, 2) if our_entry else 0,
            round(our_sl, 2) if our_sl else 0,
            round(our_tgt, 2) if our_tgt else 0,
            "Y" if exit_reason in ("TGT_HIT", "TGT") else "N",
            "",
            round(max_fav, 2) if max_fav else 0,
            round(max_adv, 2) if max_adv else 0,
            round(missed_profit, 2),
            "",
            tune_suggestion,
        ]

        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = FONT_BODY
            if col not in sep_cols:
                cell.border = BORDER_ALL

        row += 1

    ws.column_dimensions["W"].width = 60

    return ws


# ─────────────────────────────────────────────────────────────────────────────
# Sheet 5: Telegram
# ─────────────────────────────────────────────────────────────────────────────

def build_sheet_5_telegram(wb: openpyxl.Workbook, data: ReportData) -> Worksheet:
    """Build Telegram alerts sheet."""
    ws = wb.create_sheet(title="5_Telegram")
    _disable_gridlines(ws)

    group_headers = [
        ("IDENTITY", 1, 5),
        ("SENT Per Qty Prices", 7, 11),
        ("RESULT per Qty", 13, 17),
    ]

    for label, start_col, end_col in group_headers:
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        cell = ws.cell(row=1, column=start_col, value=label)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER

    sep_cols = [6, 12]
    for col in sep_cols:
        for r in range(1, 100):
            ws.cell(row=r, column=col).fill = FILL_SEPARATOR
        ws.column_dimensions[get_column_letter(col)].width = 2

    sub_headers = [
        "Trading Date", "Sent At", "Module", "Alert Type", "Symbol",
        "",
        "Qty", "Entry Price", "SL", "TGT", "Total Amount",
        "",
        "SL Hit ₹", "TGT Hit ₹", "Net P&L ₹", "Trade ID", "Delivery",
        "Retries", "Full Message"
    ]

    for col, header in enumerate(sub_headers, start=1):
        if header:
            cell = ws.cell(row=2, column=col, value=header)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER if col not in sep_cols else FILL_SEPARATOR
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_ALL

    ws.freeze_panes = "A3"

    row = 3
    ws.cell(row=row, column=1, value="Telegram alert history not stored in database.")
    ws.cell(row=row, column=1).font = FONT_BODY
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=19)
    row += 1

    ws.cell(row=row, column=1, value="Reconstructed from trade events:")
    ws.cell(row=row, column=1).font = FONT_HEADER
    row += 1

    signal_map = {s.get("signal_id"): s for s in data.signals}

    for trade in sorted(data.trades, key=lambda t: t.get("created_at", "")):
        trade_id = trade.get("trade_id", "")
        signal_id = trade.get("signal_id", "")
        signal = signal_map.get(signal_id, {})

        entry_price = trade.get("entry_actual_price") or trade.get("entry_target_price") or 0
        qty = trade.get("qty_filled") or trade.get("qty_planned") or 0
        sl = trade.get("sl_initial") or 0
        tgt = trade.get("tgt_initial") or 0
        net_pnl = trade.get("net_pnl") or 0
        exit_reason = trade.get("exit_reason", "")

        alert_type = "ORDER_PLACED"
        if exit_reason in ("SL_HIT", "SL"):
            alert_type = "SL_HIT"
        elif exit_reason in ("TGT_HIT", "TGT"):
            alert_type = "TGT_HIT"
        elif exit_reason == "EOD":
            alert_type = "EOD_EXIT"

        sl_hit_val = net_pnl if alert_type == "SL_HIT" else ""
        tgt_hit_val = net_pnl if alert_type == "TGT_HIT" else ""

        row_data = [
            data.date_iso,
            _fmt_time(trade.get("entry_time") or trade.get("created_at")),
            "order_placer",
            alert_type,
            trade.get("symbol", ""),
            "",
            qty,
            round(entry_price, 2),
            round(sl, 2),
            round(tgt, 2),
            round(entry_price * qty, 2),
            "",
            round(sl_hit_val, 2) if sl_hit_val else "",
            round(tgt_hit_val, 2) if tgt_hit_val else "",
            round(net_pnl, 2),
            trade_id[:12] + "..." if len(trade_id) > 12 else trade_id,
            "SENT",
            0,
            f"{trade.get('direction', '')} {trade.get('symbol', '')} @ {entry_price:.2f}",
        ]

        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = FONT_BODY
            if col not in sep_cols:
                cell.border = BORDER_ALL

            if col == 4:
                if alert_type in ALERT_TYPE_FILLS:
                    cell.fill = ALERT_TYPE_FILLS[alert_type]

            if col == 17:
                cell.fill = FILL_GREEN

        row += 1

    return ws


# ─────────────────────────────────────────────────────────────────────────────
# Sheet 6: Strategy Analysis
# ─────────────────────────────────────────────────────────────────────────────

def build_sheet_6_strategy(
    wb: openpyxl.Workbook,
    data: ReportData,
    config_dir: Path | str = "config",
) -> Worksheet:
    """Build Strategy Analysis sheet with two tables."""
    ws = wb.create_sheet(title="6_Strategy_Analysis")
    _disable_gridlines(ws)

    ws.cell(row=1, column=1, value="STRATEGY-WISE PERFORMANCE").font = FONT_TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=19)

    # V3 side-task A — declared taxonomy (pipeline·horizon), APPENDED as the last two
    # columns so the absolute currency-format indices below (11-16) do not shift.
    # Fail-safe: a name missing from the map (or a bad YAML) simply shows "—".
    # E1 (25-Jul-2026): this called build_taxonomy_map() with NO argument, so the
    # taxonomy columns were always read from ./config relative to CWD no matter
    # what --config-dir said. taxonomy.py:23 was correct all along; the caller was
    # not. Default stays "config" so the no-flag cron invocation is unchanged.
    from strategies.taxonomy import build_taxonomy_map
    _tax_map = build_taxonomy_map(config_dir)

    perf_headers = [
        "Trading Date", "Strategy", "Signals Rcvd", "Processed", "Rejected", "Traded",
        "Wins", "Losses", "Breakeven", "Win Rate %",
        "Gross P&L ₹", "Net P&L ₹", "Avg Net/Trade ₹",
        "Max Win ₹", "Max Loss ₹", "Capital Used ₹", "Drawdown %",
        "Pipeline", "Horizon"
    ]

    for col, header in enumerate(perf_headers, start=1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_ALL

    ws.freeze_panes = "A3"

    # Build signal map for strategy fallback
    signal_map = {s.get("signal_id"): s for s in data.signals}

    strategies = {}
    for trade in data.trades:
        signal_id = trade.get("signal_id", "")
        signal = signal_map.get(signal_id, {})
        strat = trade.get("strategy", "") or signal.get("strategy") or "UNKNOWN"

        if strat not in strategies:
            strategies[strat] = {"trades": [], "signals": 0}
        strategies[strat]["trades"].append(trade)

    for sig in data.signals:
        strat = sig.get("strategy") or "UNKNOWN"
        if strat in strategies:
            strategies[strat]["signals"] += 1
        else:
            strategies[strat] = {"trades": [], "signals": 1}

    _PROCESSED_STATUSES = {"PROCESSED", "TRADED", "FILLED", "CLOSED", "CLOSED_MANUAL", "PLACED", "SIZED", "APPROVED", "RESERVED"}

    row = 3
    for strat, info in strategies.items():
        trades = info["trades"]
        signals = info["signals"]
        closed = [t for t in trades if t.get("status") in _CLOSED_STATUSES]

        wins = [t for t in closed if (t.get("net_pnl") or 0) > 0]
        losses = [t for t in closed if (t.get("net_pnl") or 0) < 0]
        be = [t for t in closed if (t.get("net_pnl") or 0) == 0]

        gross = sum(t.get("gross_pnl") or 0 for t in closed)
        net = sum(t.get("net_pnl") or 0 for t in closed)
        avg_net = net / len(closed) if closed else 0

        pnls = [t.get("net_pnl") or 0 for t in closed]
        max_win = max(pnls) if pnls else 0
        max_loss = min(pnls) if pnls else 0

        capital_used = sum(t.get("margin_reserved") or 0 for t in trades)
        win_rate = len(wins) / len(closed) * 100 if closed else 0
        drawdown_pct = 0.0
        if capital_used:
            drawdown_pct = round(min(max_loss or 0, 0) / capital_used * 100, 1)

        processed_count = sum(1 for s in data.signals if s.get("strategy") == strat and s.get("status", "") in _PROCESSED_STATUSES)

        row_data = [
            data.date_iso,
            strat,
            signals,
            processed_count,
            sum(1 for s in data.signals if s.get("strategy") == strat and sig_status.is_rejected(s.get("status"))),
            len(trades),
            len(wins),
            len(losses),
            len(be),
            round(win_rate, 1),
            round(gross, 2),
            round(net, 2),
            round(avg_net, 2),
            round(max_win, 2),
            round(max_loss, 2),
            round(capital_used, 2),
            drawdown_pct,
            (_tax_map.get(strat, ("—", "—"))[0]),   # Pipeline (declared; "—" if unknown)
            (_tax_map.get(strat, ("—", "—"))[1]),   # Horizon
        ]

        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = FONT_BODY
            cell.border = BORDER_ALL
            if col in (11, 12, 13, 14, 15, 16):
                cell.number_format = NUM_FMT_CURRENCY

        row += 1

    row += 3

    ws.cell(row=row, column=1, value="TIME-OF-DAY ANALYSIS").font = FONT_TITLE
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=17)
    row += 1

    time_headers = [
        "Trading Date", "Time Bucket", "Signals Rcvd", "Processed", "Rejected", "Traded",
        "Wins", "Losses", "Breakeven", "Win Rate %",
        "Gross P&L ₹", "Net P&L ₹", "Avg Net/Trade ₹",
        "Max Win ₹", "Max Loss ₹", "Capital Used ₹", "Drawdown %"
    ]

    for col, header in enumerate(time_headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_ALL

    row += 1

    time_buckets = [
        ("09:15-09:30", "09:15:00", "09:30:00"),
        ("09:30-09:45", "09:30:00", "09:45:00"),
        ("09:45-10:30", "09:45:00", "10:30:00"),
        ("10:30-15:17", "10:30:00", "15:17:00"),
        ("15:17 EOD", "15:17:00", "15:30:00"),
    ]

    for bucket_name, start_time, end_time in time_buckets:
        bucket_trades = []
        bucket_signals = 0

        for trade in data.trades:
            entry_time = _fmt_time(trade.get("entry_time"))
            if entry_time and start_time <= entry_time < end_time:
                bucket_trades.append(trade)

        for sig in data.signals:
            recv_time = _fmt_time(sig.get("received_at"))
            if recv_time and start_time <= recv_time < end_time:
                bucket_signals += 1

        closed = [t for t in bucket_trades if t.get("status") in _CLOSED_STATUSES]
        wins = [t for t in closed if (t.get("net_pnl") or 0) > 0]
        losses = [t for t in closed if (t.get("net_pnl") or 0) < 0]
        be = [t for t in closed if (t.get("net_pnl") or 0) == 0]

        gross = sum(t.get("gross_pnl") or 0 for t in closed)
        net = sum(t.get("net_pnl") or 0 for t in closed)
        avg_net = net / len(closed) if closed else 0

        pnls = [t.get("net_pnl") or 0 for t in closed]
        max_win = max(pnls) if pnls else 0
        max_loss = min(pnls) if pnls else 0

        capital_used = sum(t.get("margin_reserved") or 0 for t in bucket_trades)
        win_rate = len(wins) / len(closed) * 100 if closed else 0
        drawdown_pct = 0.0
        if capital_used:
            drawdown_pct = round(min(max_loss or 0, 0) / capital_used * 100, 1)

        bucket_signals_list = [s for s in data.signals if start_time <= _fmt_time(s.get("received_at")) < end_time]
        processed = sum(1 for s in bucket_signals_list if s.get("status", "") in _PROCESSED_STATUSES)
        traded = len(bucket_trades)
        rejected = max(0, bucket_signals - processed)

        row_data = [
            data.date_iso,
            bucket_name,
            bucket_signals,
            processed,
            rejected,
            traded,
            len(wins),
            len(losses),
            len(be),
            round(win_rate, 1),
            round(gross, 2),
            round(net, 2),
            round(avg_net, 2),
            round(max_win, 2),
            round(max_loss, 2),
            round(capital_used, 2),
            drawdown_pct,
        ]

        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = FONT_BODY
            cell.border = BORDER_ALL
            if col in (11, 12, 13, 14, 15, 16):
                cell.number_format = NUM_FMT_CURRENCY

        row += 1

    row += 3
    ws.cell(row=row, column=1, value="SCORE BREAKDOWN BY STRATEGY").font = FONT_TITLE
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=17)
    row += 1

    strat_step_scores: Dict[str, Dict[str, List[float]]] = {}
    screener_map = {r.get("signal_id"): r for r in data.screener_results}
    for sig in data.signals:
        strat = sig.get("strategy") or "UNKNOWN"
        sr = screener_map.get(sig.get("signal_id"), {})
        raw = sr.get("step_results", "")
        if not raw:
            continue
        try:
            steps = json.loads(raw) if isinstance(raw, str) else raw
        except (json.JSONDecodeError, TypeError):
            continue
        if not isinstance(steps, dict):
            continue
        if strat not in strat_step_scores:
            strat_step_scores[strat] = {}
        for k, v in steps.items():
            if isinstance(v, (int, float)):
                strat_step_scores[strat].setdefault(k, []).append(float(v))

    all_step_names: list = []
    for scores in strat_step_scores.values():
        for k in scores:
            if k not in all_step_names:
                all_step_names.append(k)

    if all_step_names:
        sb_headers = ["Strategy"] + all_step_names
        for col, header in enumerate(sb_headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_ALL
        row += 1

        for strat, scores in strat_step_scores.items():
            ws.cell(row=row, column=1, value=strat).font = FONT_BODY
            ws.cell(row=row, column=1).border = BORDER_ALL
            for ci, step_name in enumerate(all_step_names, start=2):
                vals = scores.get(step_name, [])
                avg = sum(vals) / len(vals) if vals else 0.0
                cell = ws.cell(row=row, column=ci, value=round(avg, 1))
                cell.font = FONT_BODY
                cell.border = BORDER_ALL
            row += 1
    else:
        ws.cell(row=row, column=1, value="No screener data available").font = FONT_BODY
        row += 1

    return ws


# ─────────────────────────────────────────────────────────────────────────────
# Main generator
# ─────────────────────────────────────────────────────────────────────────────

def generate_daily_report(
    store,
    date_iso: str,
    output_dir: Path,
    config_dir: Path,
) -> Path:
    """
    Generate the daily report for the given date.

    Args:
        store: StateStore instance
        date_iso: Date in YYYY-MM-DD format
        output_dir: Directory for output files
        config_dir: Directory containing config files

    Returns:
        Path to generated xlsx file
    """
    log.info("Generating daily report for %s", date_iso)

    data = load_report_data(store, date_iso, config_dir)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    build_sheet_0_dashboard(wb, data)
    build_sheet_1_signals(wb, data)
    build_sheet_2_orders(wb, data)
    build_sheet_3_capital(wb, data)
    build_sheet_4_candles(wb, data)
    build_sheet_5_telegram(wb, data)
    build_sheet_6_strategy(wb, data, config_dir)   # E1: honour --config-dir

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"daily_report_{date_iso}.xlsx"
    wb.save(output_path)

    log.info("Daily report saved to %s", output_path)

    return output_path


# ─────────────────────────────────────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────────────────────────────────────

def _parse_args(argv=None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        prog="reports.daily_report",
        description="Generate 7-sheet daily EOD report.",
    )
    parser.add_argument(
        "--date",
        metavar="YYYY-MM-DD",
        default=None,
        help="Date to generate report for (default: today IST).",
    )
    parser.add_argument(
        "--db",
        metavar="PATH",
        default="data_store/trading_system.db",
        help="Path to SQLite database.",
    )
    parser.add_argument(
        "--output-dir",
        metavar="DIR",
        default="reports/output",
        help="Output directory for reports.",
    )
    parser.add_argument(
        "--config-dir",
        metavar="DIR",
        default="config",
        help="Config directory path.",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Generate even on holidays/weekends.",
    )
    parser.add_argument(
        "--notify",
        action="store_true",
        help="Send Telegram notification on completion.",
    )
    return parser.parse_args(argv)


def main(argv=None) -> int:
    args = _parse_args(argv)

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(name)s %(message)s",
    )

    date_iso = args.date or now_ist().strftime("%Y-%m-%d")

    try:
        date.fromisoformat(date_iso)
    except ValueError:
        print(f"ERROR: invalid date: {date_iso!r} — expected YYYY-MM-DD", file=sys.stderr)
        return 1

    config_dir = Path(args.config_dir)

    if not args.force and is_holiday_or_weekend(date_iso, config_dir):
        log.info("Skipping report — %s is non-trading day", date_iso)
        print(f"Skipping: {date_iso} is a holiday or weekend (use --force to override)")
        return 0

    db_path = Path(args.db)
    if not db_path.exists():
        print(f"ERROR: database not found: {db_path}", file=sys.stderr)
        return 1

    from core.state_store import StateStore
    store = StateStore(db_path)

    started = time.perf_counter()
    try:
        output_path = generate_daily_report(
            store=store,
            date_iso=date_iso,
            output_dir=Path(args.output_dir),
            config_dir=config_dir,
        )
        print(f"Report generated: {output_path}")

        # Observability (Control Tower Phase 1d): record a SUCCESS heartbeat so
        # the cron framework (and the tower's cron freshness check) can verify
        # the xlsx was actually produced. Without this the job ran but emitted
        # no heartbeat/marker -> a recurring false "daily_report missed" finding
        # (the 23-Jun gap). record_heartbeat never raises; job_name MUST match
        # the cron_registry key "daily_report".
        record_daily_report_heartbeat(db_path, "SUCCESS", time.perf_counter() - started)

        if args.notify:
            try:
                from alerts.telegram_notifier import TelegramNotifier
                notifier = TelegramNotifier.from_env()
                if notifier:
                    notifier.send_info(f"Daily report ready: {output_path.name}")
            except Exception as e:
                log.warning("Telegram notification failed: %s", e)

        return 0

    except Exception as e:
        log.exception("Report generation failed")
        print(f"ERROR: {e}", file=sys.stderr)
        record_daily_report_heartbeat(
            db_path, "FAILED", time.perf_counter() - started,
            message=f"{type(e).__name__}: {e}")
        return 2


def record_daily_report_heartbeat(db_path: Path, status: str,
                                  duration_sec: float, message=None) -> None:
    """Record the daily_report cron heartbeat (never raises into the job)."""
    try:
        from utils.cron_heartbeat import record_heartbeat
        record_heartbeat("daily_report", status=status,
                         duration_sec=duration_sec, message=message,
                         db_path=db_path)
    except Exception:  # pragma: no cover — defensive; heartbeat is best-effort
        pass


if __name__ == "__main__":
    sys.exit(main())
