"""
reports/style_constants.py — Trading System v2

Style constants for daily_report.xlsx generation.
All colors, fonts, fills, and number formats defined here.
"""

from openpyxl.styles import (
    Font, Fill, PatternFill, Border, Side, Alignment, NamedStyle
)

# ─────────────────────────────────────────────────────────────────────────────
# Colors (hex without #)
# ─────────────────────────────────────────────────────────────────────────────

COLOR_TITLE_BLUE = "1F4E79"        # Title text
COLOR_SEPARATOR_BLUE = "4472C4"   # Dark blue separator columns
COLOR_BORDER_GREY = "BFBFBF"      # Thin grey borders

COLOR_GREEN_FILL = "E2EFDA"       # Wins, target hits, positive values
COLOR_RED_FILL = "FCE4D6"         # Losses, SL hits, rejections, negative values
COLOR_AMBER_FILL = "FFEB9C"       # Warnings, EOD exits, mismatches
COLOR_GREY_FILL = "D9D9D9"        # Opening/closing rows, summaries

COLOR_WHITE = "FFFFFF"
COLOR_BLACK = "000000"

# ─────────────────────────────────────────────────────────────────────────────
# Fonts
# ─────────────────────────────────────────────────────────────────────────────

FONT_BODY = Font(name="Arial", size=9)
FONT_HEADER = Font(name="Arial", size=9, bold=True)
FONT_TITLE = Font(name="Arial", size=11, bold=True, color=COLOR_TITLE_BLUE)
FONT_SECTION_HEADER = Font(name="Arial", size=10, bold=True, color=COLOR_WHITE)

# ─────────────────────────────────────────────────────────────────────────────
# Fills (PatternFill)
# ─────────────────────────────────────────────────────────────────────────────

FILL_GREEN = PatternFill(start_color=COLOR_GREEN_FILL, end_color=COLOR_GREEN_FILL, fill_type="solid")
FILL_RED = PatternFill(start_color=COLOR_RED_FILL, end_color=COLOR_RED_FILL, fill_type="solid")
FILL_AMBER = PatternFill(start_color=COLOR_AMBER_FILL, end_color=COLOR_AMBER_FILL, fill_type="solid")
FILL_GREY = PatternFill(start_color=COLOR_GREY_FILL, end_color=COLOR_GREY_FILL, fill_type="solid")
FILL_SEPARATOR = PatternFill(start_color=COLOR_SEPARATOR_BLUE, end_color=COLOR_SEPARATOR_BLUE, fill_type="solid")
FILL_HEADER = PatternFill(start_color=COLOR_SEPARATOR_BLUE, end_color=COLOR_SEPARATOR_BLUE, fill_type="solid")

# Sheet 2_Orders specific: 3-row header design
FILL_TITLE_BG     = PatternFill(start_color="DEEAF1", end_color="DEEAF1", fill_type="solid")
FILL_GROUP_HEADER = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
FILL_SUB_HEADER   = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
FONT_WHITE_BOLD   = Font(name="Arial", size=9, bold=True, color="FFFFFF")

# ─────────────────────────────────────────────────────────────────────────────
# Borders
# ─────────────────────────────────────────────────────────────────────────────

THIN_GREY_SIDE = Side(style="thin", color=COLOR_BORDER_GREY)
BORDER_ALL = Border(
    left=THIN_GREY_SIDE,
    right=THIN_GREY_SIDE,
    top=THIN_GREY_SIDE,
    bottom=THIN_GREY_SIDE
)

# ─────────────────────────────────────────────────────────────────────────────
# Alignment
# ─────────────────────────────────────────────────────────────────────────────

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

# ─────────────────────────────────────────────────────────────────────────────
# Number formats
# ─────────────────────────────────────────────────────────────────────────────

NUM_FMT_CURRENCY = '#,##0.00'
NUM_FMT_CURRENCY_NEG_PARENS = '#,##0.00;(#,##0.00);"-"'
NUM_FMT_PERCENT = '0.00%'
NUM_FMT_RATIO = '0.00'
NUM_FMT_INTEGER = '#,##0'

# ─────────────────────────────────────────────────────────────────────────────
# Exit reason -> fill mapping
# ─────────────────────────────────────────────────────────────────────────────

EXIT_REASON_FILLS = {
    "TGT_HIT": FILL_GREEN,
    "TGT": FILL_GREEN,
    "SL_HIT": FILL_RED,
    "SL": FILL_RED,
    "EOD": FILL_AMBER,
    "MANUAL": FILL_AMBER,
    "TIMEOUT": FILL_AMBER,
}

# ─────────────────────────────────────────────────────────────────────────────
# Alert type -> fill mapping
# ─────────────────────────────────────────────────────────────────────────────

ALERT_TYPE_FILLS = {
    "SIGNAL": PatternFill(start_color="DEEBF7", end_color="DEEBF7", fill_type="solid"),  # Light blue
    "ORDER_PLACED": FILL_GREEN,
    "SL_HIT": FILL_RED,
    "TGT_HIT": FILL_GREEN,
    "SOFT_KILL": FILL_AMBER,
    "EOD_SUMMARY": FILL_GREY,
}

# ─────────────────────────────────────────────────────────────────────────────
# Delivery status -> fill mapping
# ─────────────────────────────────────────────────────────────────────────────

DELIVERY_FILLS = {
    "SENT": FILL_GREEN,
    "FAILED": FILL_RED,
}
