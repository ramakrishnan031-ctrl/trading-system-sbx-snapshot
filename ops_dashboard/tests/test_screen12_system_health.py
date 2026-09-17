"""Screen 12 — System Health.

A health screen that shows green for something it did not measure is worse than
no health screen, so most of what follows guards exactly that: UNKNOWN must
never collapse into HEALTHY, an un-instrumented metric must never acquire a
number, and a stale READY must not sit silently beside a dead service.

Plus the approved alert requirement: the alert's own WORDS carry the semantic
colour, and a positive event stays green even inside an Alerts card.

⭐ Assertions are anchored to the fixture's KNOWN values or to a PROPERTY that
must hold. ⛔ No assertion is written by reading back what the code produced.

FIXTURE (conftest), so every number below is traceable:
  preflight run pf_a_1, phase A, overall_status READY, 13 check rows:
    Broker   5 PASS · Database 2 PASS · Engine 3 (2 PASS + capital_deployment WARN)
    State    2 PASS · VM Health 1 PASS
  ⇒ Capital pillar = WARNING (it draws capital_deployment BY NAME)
  autofix log: one SUCCESS, one FAILED, one with NO result ⇒ TRIGGERED
  the trading engine is NOT running under test ⇒ Trading Engine FAILED
  systemctl does not exist on the test host ⇒ every unit UNKNOWN
"""
from __future__ import annotations

import inspect
import io
import os
import re

from backend.services import system_health as sh


def _read(*parts) -> str:
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    with open(os.path.join(here, *parts), encoding="utf-8") as fh:
        return fh.read()


def _tpl() -> str:
    return _read("frontend", "templates", "services.html")


def _css() -> str:
    return _read("frontend", "static", "style.css")


def _shipped(client) -> str:
    html = client.get("/services").data.decode("utf-8")
    return re.sub(r"<!--.*?-->", " ", html, flags=re.S)


def _h(client, extra=""):
    return client.get("/api/system-health" + extra).get_json()


# ── UNKNOWN IS NOT HEALTHY — the screen's core safety property ───────────────
def test_unknown_outranks_healthy_in_every_rollup():
    """⛔ THE failure this guards: an overall HEALTHY computed over components
    that could not be evaluated. Not knowing is not the same as being well."""
    assert sh._worst(["HEALTHY", "UNKNOWN"]) == "UNKNOWN"
    assert sh._worst(["HEALTHY", "HEALTHY"]) == "HEALTHY"
    assert sh._worst(["UNKNOWN", "WARNING"]) == "WARNING"
    assert sh._worst(["WARNING", "FAILED"]) == "FAILED"
    assert sh._worst(["UNKNOWN", "FAILED"]) == "FAILED"
    # ⛔ an EMPTY set is UNKNOWN, never HEALTHY — nothing measured is not "fine"
    assert sh._worst([]) == "UNKNOWN"


def test_a_host_without_systemctl_reports_unknown_not_failed(client):
    """On a non-Linux host `systemctl` does not exist. ⛔ That must not read as
    every service being down (a false incident) nor as up (a false all-clear)."""
    units = [s for s in _h(client)["services"] if s["kind"] == "systemd"]
    assert units, "the fixture configures systemd units"
    for u in units:
        assert u["status"] == "UNKNOWN", u["service"]
        assert u["raw_state"] == "unavailable"


def test_overall_status_is_never_hard_coded_healthy(client):
    """The trading engine is not running under test, so overall MUST reflect it."""
    d = _h(client)
    assert d["overall"]["status"] == "FAILED"
    engine = [s for s in d["services"] if s["service"] == "Trading Engine"][0]
    assert engine["status"] == "FAILED"
    assert "UNKNOWN outranks HEALTHY" in d["overall"]["rollup"]


def test_counts_partition_the_service_population(client):
    """A PROPERTY: the four counts must always add to the total."""
    d = _h(client)
    assert sum(d["counts"].values()) == d["total_services"] == len(d["services"])


# ── NOT INSTRUMENTED never becomes a number ─────────────────────────────────
def test_cpu_ram_and_network_are_live_metrics_that_this_host_cannot_serve(client):
    """⭐ These are read LIVE from /proc — Linux-only. On this (Windows) host the
    source does not exist, so they are NO DATA: real metrics this machine cannot
    serve, ⛔ NOT an instrumentation gap in the system.

    ⛔ Whatever the label, none of the three may show a value or a green status."""
    vm = _h(client)["vm"]
    for key in ("cpu_pct", "ram_pct", "network"):
        assert vm[key]["measured"] is False, key
        assert vm[key]["value"] is None, key
        assert vm[key]["status"] == "UNKNOWN", key      # ⛔ never green
        assert vm[key]["instrumented"] is True, key     # ⛔ not a system gap
        assert vm[key]["reason"], key
    assert "/proc/stat" in vm["cpu_pct"]["reason"]
    assert "/proc/meminfo" in vm["ram_pct"]["reason"]
    assert "/proc/net/dev" in vm["network"]["reason"]


def test_the_live_vm_metrics_populate_from_native_sources_with_no_gui_change(monkeypatch):
    """⭐⭐ THE PRODUCTION GUARANTEE, and the reason this test exists: Rama's
    requirement is that deploying onto the real Linux VM populates CPU / RAM /
    RAM-available / Network / Load *automatically*, with ⛔ no plugin, ⛔ no
    second monitoring system and ⛔ no future GUI change.

    The screen must therefore READ these, ⛔ not hard-code a verdict about them.
    Feeding the reader exactly what a Linux host yields proves the panel renders
    real numbers through the shipped code path."""
    monkeypatch.setattr(sh.host_reader, "vm_stats", lambda cfg: {
        "available": True, "cpu_pct": 37.5, "mem_total_kb": 2_048_000,
        "mem_available_kb": 512_000, "mem_used_pct": 75.0,
        "net_bytes_per_sec": 3_145_728, "loadavg": [1.25, 0.9, 0.7],
        "cpu_count": 2,
        "disk_root": {"used_pct": 41.8}, "disk_data": {"used_pct": 41.8},
    })
    vm = sh._vm_health({"paths": {}})
    assert vm["cpu_pct"]["measured"] is True and vm["cpu_pct"]["value"] == 37.5
    assert vm["ram_pct"]["measured"] is True and vm["ram_pct"]["value"] == 75.0
    assert vm["ram_available_mb"]["measured"] is True
    assert vm["network"]["measured"] is True
    assert vm["network"]["value"] == "3.0 MB/s"          # ⛔ a RATE, not a total
    assert vm["system_load"]["measured"] is True and vm["system_load"]["value"] == 1.25
    assert vm["disk_pct"]["measured"] is True
    # ⛔ and a busy box must not read green just because it answered
    assert vm["cpu_pct"]["status"] == "HEALTHY"          # 37.5% is genuinely fine
    hot = sh._band(97.0, 80, 95)
    assert hot == "FAILED" and sh._band(85.0, 80, 95) == "WARNING"
    assert sh._band(None, 80, 95) == "UNKNOWN"           # ⛔ unread is never green


def test_the_native_proc_parsers_read_the_real_kernel_formats(monkeypatch):
    """/proc/stat and /proc/net/dev publish CUMULATIVE counters, in a specific
    layout. These parsers are fed the genuine kernel formats."""
    import builtins
    from backend.readers import host_reader as hr

    FAKE = {
        # cpu  user nice system idle iowait irq softirq steal …
        "/proc/stat": "cpu  100 5 95 700 100 0 0 0 0 0\ncpu0 1 2 3 4\nintr 9\n",
        # iface: rxbytes rxpkts … (8 rx fields) txbytes …
        "/proc/net/dev": (
            "Inter-|   Receive                          |  Transmit\n"
            " face |bytes    packets errs drop fifo frame compressed multicast|"
            "bytes    packets\n"
            "    lo: 999999  10 0 0 0 0 0 0 999999  10\n"
            "  eth0: 1000     10 0 0 0 0 0 0 2000     20\n"),
    }
    real_open = builtins.open
    monkeypatch.setattr(builtins, "open", lambda p, *a, **k: (
        io.StringIO(FAKE[p]) if p in FAKE else real_open(p, *a, **k)))

    idle, total = hr._read_proc_stat()
    assert idle == 800                      # idle 700 + iowait 100
    assert total == 1000                    # every field summed
    # ⛔ loopback EXCLUDED — counting `lo` reports the box talking to itself
    # as network traffic. eth0 only: 1000 rx + 2000 tx.
    assert hr._read_net_bytes() == 3000


def test_vm_stats_end_to_end_on_a_simulated_linux_host(monkeypatch, tmp_path):
    """⭐⭐ THE REAL DEPLOYMENT REHEARSAL. The tests above stub `vm_stats`; this
    one runs the SHIPPED function with `platform.system()` reporting Linux and
    the genuine kernel file formats behind `open`, so the whole production path
    — two-sample timing, parsing, arithmetic — is exercised, ⛔ not just its
    callers.

    ⚠️ This is as close as a Windows PC can get. It is ⛔ NOT a claim that the
    VM has been measured; that must be confirmed once on the real host."""
    import builtins
    from backend.readers import host_reader as hr

    # /proc/stat advances 100 jiffies between reads, 25 of them idle ⇒ 75% busy.
    stat_reads = iter([
        "cpu  100 5 95 700 100 0 0 0 0 0\n",
        "cpu  160 5 110 725 100 0 0 0 0 0\n",
    ])
    # eth0 gains 200 bytes over the 0.2s window ⇒ 1000 B/s. `lo` is noise.
    net_reads = iter([
        "  lo: 900 1 0 0 0 0 0 0 900 1\n  eth0: 1000 1 0 0 0 0 0 0 2000 1\n",
        "  lo: 999 1 0 0 0 0 0 0 999 1\n  eth0: 1100 1 0 0 0 0 0 0 2100 1\n",
    ])
    MEM = "MemTotal:       2048000 kB\nMemAvailable:    512000 kB\n"
    real_open = builtins.open

    def fake_open(path, *a, **k):
        if path == "/proc/stat":
            return io.StringIO(next(stat_reads))
        if path == "/proc/net/dev":
            return io.StringIO(next(net_reads))
        if path == "/proc/meminfo":
            return io.StringIO(MEM)
        return real_open(path, *a, **k)

    monkeypatch.setattr(builtins, "open", fake_open)
    monkeypatch.setattr(hr.platform, "system", lambda: "Linux")
    monkeypatch.setattr(hr.os, "getloadavg", lambda: (1.25, 0.9, 0.7), raising=False)

    out = hr.vm_stats({"paths": {"data_store": str(tmp_path)}})
    assert out["available"] is True
    assert out["cpu_pct"] == 75.0                    # (100-25)/100
    assert out["mem_total_kb"] == 2048000
    assert out["mem_used_pct"] == 75.0               # (2048000-512000)/2048000
    assert out["net_bytes_per_sec"] == 1000          # 200 bytes / 0.2s, lo excluded
    assert out["loadavg"][0] == 1.25
    assert out["disk_data"] and out["disk_data"]["used_pct"] >= 0

    # ⭐ and the SCREEN renders those numbers through the same shipped path
    monkeypatch.setattr(sh.host_reader, "vm_stats", lambda cfg: out)
    vm = sh._vm_health({"paths": {}})
    # 75% is below both warn bands (CPU 80, RAM 85) ⇒ genuinely HEALTHY
    assert vm["cpu_pct"]["value"] == 75.0 and vm["cpu_pct"]["status"] == "HEALTHY"
    assert vm["ram_pct"]["value"] == 75.0 and vm["ram_pct"]["status"] == "HEALTHY"
    assert vm["network"]["value"] == "1000 B/s"
    assert vm["system_load"]["value"] == 1.25


def test_cpu_percent_is_a_delta_never_a_single_reading():
    """⛔ A single /proc/stat read is a cumulative counter, not a percentage.
    The shipped code differences two samples — pinned at the source so it cannot
    regress into reporting the counter itself."""
    from backend.readers import host_reader as hr
    src = inspect.getsource(hr.vm_stats)
    assert "_RATE_SAMPLE_SEC" in src and "_time.sleep" in src
    assert "d_total > 0" in src             # ⛔ no divide-by-zero, no fake 0%
    # the sample window is short enough not to stall the page
    assert 0 < hr._RATE_SAMPLE_SEC <= 0.5


def test_a_gap_is_never_green_on_any_path(client):
    """⛔ Whatever the filter, an un-instrumented metric stays UNKNOWN."""
    for q in ("", "?status=HEALTHY", "?kind=systemd", "?status=FAILED"):
        vm = _h(client, q)["vm"]
        for k, v in vm.items():
            if isinstance(v, dict) and v.get("measured") is False:
                assert v["status"] == "UNKNOWN", (q, k)
                assert v["value"] is None, (q, k)


def test_disk_is_measured_because_it_genuinely_is(client):
    """⭐ The control: disk uses stdlib shutil and works on every platform, so it
    MUST be measured. Without it, "everything is a gap" would pass vacuously."""
    disk = _h(client)["vm"]["disk_pct"]
    assert disk["measured"] is True
    assert isinstance(disk["value"], (int, float)) and 0 <= disk["value"] <= 100


def test_db_connection_count_is_not_applicable_not_zero(client):
    """SQLite is embedded — a count would describe the dashboard, not the trader."""
    db = _h(client)["database"]
    assert db["connection_count"] is None
    assert "no server" in db["connection_count_note"]


def test_trends_declare_which_series_do_not_exist(client):
    tr = _h(client)["trends"]
    for k in ("cpu", "ram", "response_time"):
        assert tr[k]["measured"] is False, k
        assert tr[k]["instrumented"] is False, k     # ⛔ never merely "no data"
        assert tr[k]["series"] == [], k
        assert tr[k]["reason"], k
    assert "sentinel" in tr["cpu"]["reason"]


def test_the_disk_series_is_real_and_is_not_labelled_not_instrumented(client):
    """⚠️⚠️ REGRESSION PIN. The shipped reader selected `ts` and ordered by `id`,
    but production `system_metrics` (core/analytics_schema.sql:51) has neither —
    it has `timestamp` and NO `id`. Against the real DB the query raised
    `OperationalError: no such column: ts`, which was swallowed into `[]`, so a
    metric the VM genuinely collects rendered as "NOT INSTRUMENTED".

    ⛔ It passed its tests anyway because the FIXTURE invented `id`+`ts`. The
    fixture now mirrors the shipped schema, so this test could not pass against
    the old reader."""
    disk = _h(client)["trends"]["disk"]
    assert disk["instrumented"] is True          # ⛔ NOT an instrumentation gap
    assert disk["measured"] is True
    assert len(disk["series"]) == 6
    assert all(p["ts"] and p["value"] is not None for p in disk["series"])
    # oldest → newest, so the chart's x-axis cannot run backwards
    assert [p["ts"] for p in disk["series"]] == sorted(p["ts"] for p in disk["series"])
    # ⛔ and the -1.0 psutil sentinels seeded alongside are NEVER plotted
    assert all(p["value"] >= 0 for p in disk["series"])


def test_not_instrumented_and_no_data_are_different_states_everywhere(client):
    """⛔ A metric the system DOES measure must never be labelled NOT
    INSTRUMENTED just because this host has recorded nothing. Under-reporting
    the system is as wrong as over-reporting it."""
    d = _h(client)
    vm, bk, tr = d["vm"], d["broker"], d["trends"]
    # ⛔ GENUINELY UNINSTRUMENTED — nothing records these anywhere
    assert bk["last_api_call"]["instrumented"] is False
    assert tr["response_time"]["instrumented"] is False   # never persisted
    assert tr["cpu"]["instrumented"] is False             # collector writes -1.0
    assert tr["ram"]["instrumented"] is False
    # ✅ INSTRUMENTED, simply not available/recorded on THIS host
    for k in ("cpu_pct", "ram_pct", "network", "ram_available_mb", "system_load"):
        assert vm[k]["instrumented"] is True, k
    assert bk["last_successful_order"]["instrumented"] is True
    assert "orders.filled_at" in bk["last_successful_order"]["reason"]
    # ⭐ AND THE DISTINCTION IS THE POINT: the LIVE cpu figure is instrumented
    # (read from /proc) while its persisted HISTORY is not (collector sentinel).
    # Same metric, two different states — collapsing them would mislead either way.
    assert vm["cpu_pct"]["instrumented"] != tr["cpu"]["instrumented"]
    # ⛔ none of them ever produces a number or a green status here
    for m in (vm["cpu_pct"], vm["system_load"], bk["last_successful_order"]):
        assert m["value"] is None and m["status"] == "UNKNOWN"


def test_the_export_spells_out_which_kind_of_empty_each_cell_is(client):
    """⛔ A blank spreadsheet cell reads as zero; a wrong label reads as a gap in
    the system. The XLSX must say which."""
    payload = _h(client)
    sheets = {name: (hdr, rows) for name, hdr, rows in sh.export_sheets(payload)}
    summary = dict(sheets["Summary"][1])
    # live /proc metrics: real, unavailable on this host
    assert summary["CPU %"] == "NO DATA"
    assert summary["Network"] == "NO DATA"
    assert summary["System Load"] == "NO DATA"
    assert summary["Last successful order"] == "NO DATA"
    # genuinely uninstrumented anywhere in the system
    assert summary["Last API call"] == "NOT INSTRUMENTED"
    assert summary["TRADING READINESS (NOW)"] == "NOT READY"
    assert summary["PREFLIGHT VERDICT (HISTORICAL)"] == "READY"
    assert "SUPERSEDED" in summary["Preflight superseded"]


def test_ready_with_warnings_never_collapses_to_a_bare_ready():
    """⛔ The warnings are the whole reason an operator reads that block, so the
    label must carry them even though `ready` is True either way."""
    warned = sh._preflight_verdict({"overall_status": "READY_WITH_WARNINGS",
                                    "phase": "B", "started_at": "2026-08-15T09:14:00",
                                    "checks": [{"check_name": "x", "check_group": "Broker",
                                                "criticality": "LOW", "status": "WARN"}]})
    assert warned["ready"] is True
    assert warned["label"] == "READY (WITH WARNINGS)"
    assert warned["verdict"] == "WARNING"
    clean = sh._preflight_verdict(_PF_READY)
    assert clean["label"] == "READY" and clean["verdict"] == "HEALTHY"


def test_the_frontend_renders_both_emptiness_labels(client):
    body = _shipped(client)
    assert "NOT INSTRUMENTED" in body and "NO DATA" in body
    # the instrumented-but-empty case is amber, not the grey unknown treatment
    assert "d.instrumented" in body
    assert '"num-warn" title=\'' in body or "num-warn" in body


def test_the_disk_reader_queries_the_columns_production_actually_has():
    """⭐ The narrow guard on the exact defect: a source-level assertion that the
    reader can never drift back to the fixture-only column names."""
    import inspect
    from backend.readers import db_reader
    src = inspect.getsource(db_reader.system_metrics_disk_history)
    # ⛔ Slice AFTER the docstring — it quotes the old broken SQL while explaining
    # the fix, so scanning the whole source would match the very thing it forbids.
    body = src[src.index("with _ro(cfg)"):]
    assert "SELECT timestamp, disk_used_pct" in body
    assert "ORDER BY timestamp DESC" in body
    assert "SELECT ts," not in body and "ORDER BY id" not in body


def test_an_instrumented_metric_with_no_rows_says_no_data_not_not_instrumented(tmp_path):
    """⭐ THE THIRD EMPTINESS. A dev host that never ran the 5-minute collector
    has no disk rows — that is ⛔ NOT the same as the metric not existing, and
    conflating the two under-reports the system exactly as a sentinel chart
    would over-report it."""
    empty = tmp_path / "analytics_empty.db"
    import sqlite3
    c = sqlite3.connect(str(empty))
    c.execute("CREATE TABLE system_metrics (timestamp TEXT, disk_used_pct REAL)")
    c.commit()
    c.close()
    out = sh._trends({"paths": {"main_db": str(empty), "analytics_db": str(empty)}})
    assert out["disk"]["instrumented"] is True       # ⛔ not a gap in the system
    assert out["disk"]["measured"] is False
    assert "no snapshots" in out["disk"]["reason"]


# ── SEMANTIC ALERT COLOUR — the approved requirement ────────────────────────
def test_positive_alerts_are_green_even_on_a_critical_feed():
    """⛔ THE defect the approved design names: a positive event painted red
    merely because it arrived in an Alerts card."""
    for title in ("Connection restored", "Recovery Success",
                  "Service healthy again", "Backup completed successfully",
                  "Broker reconnected"):
        assert sh.classify_alert("CRITICAL", title) == "HEALTHY", title
        assert sh.classify_alert("INFO", title) == "HEALTHY", title


def test_negative_alerts_stay_red():
    for title in ("Backtest Engine is down", "Recovery Failed",
                  "Critical failure in order path", "Broker unreachable"):
        assert sh.classify_alert("CRITICAL", title) == "FAILED", title


def test_a_hopeful_title_cannot_downgrade_a_real_failure():
    """⭐ The override is ONE-WAY. A CRITICAL alert is never demoted below its
    declared severity by a word — only a plain RESOLUTION is promoted."""
    assert sh.classify_alert("CRITICAL", "Order rejected") == "FAILED"
    assert sh.classify_alert("WARNING", "Latency above threshold") == "WARNING"
    # a bare, unmarked title keeps its severity
    assert sh.classify_alert("CRITICAL", "Something happened") == "FAILED"
    assert sh.classify_alert("WARNING", "Something happened") == "WARNING"


def test_severity_maps_to_the_four_statuses():
    assert sh.classify_alert("CRITICAL", "x") == "FAILED"
    assert sh.classify_alert("WARNING", "x") == "WARNING"
    assert sh.classify_alert("INFO", "x") == "UNKNOWN"
    assert sh.classify_alert(None, "x") == "UNKNOWN"


def test_alert_TEXT_carries_the_colour_not_only_an_icon():
    """The approved requirement is explicit: colour the words. The alert title
    element must bind the semantic class, ⛔ not just a dot or a badge."""
    t = _tpl()
    block = t.split('class="sysh-alerts"', 1)[1].split("</ul>", 1)[0]
    assert 'class="sysh-al-t"' in block
    assert ':class="txt(a.status)"' in block
    # and the severity label too
    assert 'class="sysh-sev"' in block


def test_the_tall_readiness_card_is_packed_beside_a_stack_not_a_single_short_card():
    """⚠️ REGRESSION PIN FOR THE LAYOUT FIX. Measured panel heights are
    ~220·206·206·173·495px. Laid out 3-then-2, the 495px readiness card sat on
    line 2 beside a 173px Auto-Recovery, so the grid row took the taller and left
    **323px of dead background** — and pushed Health Trends that far down.

    ⭐ The fix is PACKING, ⛔ not shrinking a card: readiness holds column 3 for
    BOTH rows so the other four stack two-by-two beside it. Worst dead space
    measured after: 85px at 1600, 94px at 1920."""
    css = _css()
    assert ".sysh-page .sysh-zone-a .sysh-ready { grid-column: 3; grid-row: 1 / span 2; }" in css
    # ⛔ the 5-across template is gone: it made the readiness column ~390px wide
    # at 1920, wrapping its text and making the WIDEST viewport the worst void.
    assert ".85fr 1fr 1.15fr 1fr 1.05fr" not in css
    assert ".sysh-page .sysh-zone-a { grid-template-columns: 1fr 1fr 1.25fr; }" in css
    # ⛔ and the span must be undone where it cannot fit, or the page overflows
    narrow = css[css.index("@media (max-width: 1100px)"):]
    assert "grid-column: auto; grid-row: auto;" in narrow


#: The header line that opens the NEXT screen's block in `style.css`. ⛔ A CSS
#: window that runs to end-of-file makes every screen added later fail a test
#: that was never about it — a false red, not a finding.
_NEXT_SCREEN_MARK = "\n   SCREEN "


def test_cards_size_to_their_content_rather_than_stretching():
    """⛔ `stretch` puts the slack INSIDE the short card, where it reads as
    missing rows — the N15-10 trap in the opposite direction."""
    css = _css()
    assert ".sysh-page .sysh-zone-row { display: grid; gap: 16px; align-items: start; }" in css
    # ⛔ bounded to THIS screen's own block: a window running to end-of-file
    # would fail on any later screen's CSS, which this test is not about.
    start = css.index(".sysh-page .sysh-zone")
    nxt = css.find(_NEXT_SCREEN_MARK, start)
    block = css[start:nxt] if nxt > 0 else css[start:]
    assert "align-items: stretch" not in block


def test_the_css_does_not_repaint_alert_text_over_the_semantic_class():
    """⛔ A colour on `.sysh-al-t` would override the semantic class and make
    every alert one colour — which is the whole defect."""
    css = _css()
    rule = re.search(r"\.sysh-page \.sysh-al-t \{([^}]*)\}", css)
    assert rule, "the rule must exist"
    assert "color" not in rule.group(1)


def test_semantic_colour_helper_maps_all_four_states():
    t = _tpl()
    body = re.search(r"txt\(status\)\s*\{(.*?)\n    \},", t, re.S).group(1)
    assert "HEALTHY: \"num-pos\"" in body
    assert "WARNING: \"num-warn\"" in body
    assert "FAILED: \"num-neg\"" in body
    assert "UNKNOWN" in body and "sysh-unknown" in body


def test_status_colours_are_the_projects_own(client):
    css = _css()
    assert ".sysh-page .sysh-HEALTHY { background: var(--pos-tint);   color: var(--pos); }" in css
    assert ".sysh-page .sysh-WARNING { background: var(--warn-tint);  color: var(--yellow); }" in css
    assert ".sysh-page .sysh-FAILED  { background: var(--neg-tint);   color: var(--neg); }" in css
    assert ".sysh-page .sysh-UNKNOWN { background: var(--muted-tint); color: var(--dim); }" in css


# ── TRADING READINESS ───────────────────────────────────────────────────────
# ⚠️⚠️ THE BLOCKING DEFECT RAMA REJECTED ON 15-Aug: the screen showed
#     Overall Status = FAILED · Trading Engine = FAILED · Trading Readiness = READY
# because preflight's MORNING verdict was rendered as the CURRENT answer. His
# ruling: "Make the live readiness verdict authoritative … a stale preflight
# READY result must never remain the main current READY state after a required
# live service has failed." These tests hold that line.

def _svc(name, status, kind="systemd"):
    return {"service": name, "kind": kind, "status": status,
            "required": sh._is_required({"service": name, "kind": kind})}


_PF_READY = {"overall_status": "READY", "phase": "A",
             "started_at": "2026-08-15T08:30:00",
             "checks": [{"check_name": "kite_token_fresh_today",
                         "check_group": "Broker", "criticality": "CRITICAL",
                         "status": "PASS"}]}


def test_a_failed_required_service_forces_not_ready_even_when_preflight_said_ready():
    """⛔⛔ THE REGRESSION GUARD FOR THE REJECTED BUILD. A live FAILED trading
    engine must produce NOT READY, ⛔ never a READY carrying a warning."""
    out = sh._live_readiness(sh._preflight_verdict(_PF_READY),
                             [_svc("Trading Engine", "FAILED", "engine"),
                              _svc("Database", "HEALTHY", "database")])
    assert out["ready"] is False
    assert out["label"] == "NOT READY"
    assert out["verdict"] == "FAILED"
    assert "Trading Engine" in out["failed_required"]
    # ⭐ and it SAYS the morning verdict is older than the failure
    assert "older than this failure" in out["reason"]


def test_the_live_screen_no_longer_shows_ready_beside_a_dead_engine(client):
    """The END-TO-END form of the same rule, on the real payload: this fixture
    has preflight READY and a dead engine — exactly the rejected combination."""
    d = _h(client)
    assert d["overall"]["status"] == "FAILED"
    assert d["readiness"]["ready"] is False          # ⛔ was True in the rejected build
    assert d["readiness"]["label"] == "NOT READY"
    assert "Trading Engine" in d["readiness"]["failed_required"]


def test_preflight_is_retained_as_history_and_marked_superseded(client):
    """⛔ Rama: "If historical preflight readiness is retained, show it separately
    as historical context, not as the current READY verdict.\""""
    pf = _h(client)["readiness"]["preflight"]
    assert pf["ready"] is True                       # the morning verdict SURVIVES
    assert pf["overall_status"] == "READY"
    assert pf["label"] == "READY"
    assert pf["evaluated_at"]                        # ...carrying its own timestamp
    assert pf["superseded"]["by"] == "NOT READY"
    assert "SUPERSEDED" in pf["superseded"]["message"]


def test_a_failed_NON_required_service_does_not_force_not_ready():
    """⭐ THE OTHER HALF, and it is what keeps the rule from being a blunt copy of
    Overall Status: a dead ALERT watcher is a real problem but it does ⛔ NOT make
    an open position unsafe. Colour Overall Status, ⛔ do not block trading."""
    out = sh._live_readiness(sh._preflight_verdict(_PF_READY),
                             [_svc("Trading Engine", "HEALTHY", "engine"),
                              _svc("Database", "HEALTHY", "database"),
                              _svc("alert-watcher.service", "FAILED")])
    assert out["ready"] is True
    assert out["label"] == "READY"
    assert out["failed_required"] == []


def test_security_watcher_activating_is_never_a_readiness_blocker():
    """⚠️ `activating/auto-restart` is security-watcher's DESIGNED RestartSec=60
    heartbeat. Requiring it would fire a permanent FALSE NOT READY."""
    assert sh._is_required({"service": "security-watcher.service",
                            "kind": "systemd"}) is False
    assert sh._is_required({"service": "token-watcher.service",
                            "kind": "systemd"}) is False
    assert sh._is_required({"service": "Dashboard API", "kind": "dashboard"}) is False
    assert sh._is_required({"service": "trading-system.service",
                            "kind": "systemd"}) is True


def test_an_unknown_required_service_is_not_confirmed_never_ready():
    """⛔ UNKNOWN is never silently treated as HEALTHY — but it is also not a
    fabricated incident. It is its own third state."""
    out = sh._live_readiness(sh._preflight_verdict(_PF_READY),
                             [_svc("Trading Engine", "UNKNOWN", "engine")])
    assert out["ready"] is None
    assert out["label"] == "NOT CONFIRMED"
    assert out["verdict"] == "UNKNOWN"
    assert "could not be confirmed" in out["reason"]


def test_a_critical_preflight_failure_is_not_cured_by_healthy_services():
    """⭐ THE GATE RUNS BOTH WAYS. A missing holiday calendar stays missing no
    matter how many services are up — live health cannot clear a failed
    critical pre-check."""
    pfv = sh._preflight_verdict({
        "overall_status": "CRITICAL_FAILURE", "phase": "B",
        "started_at": "2026-08-15T09:14:00",
        "checks": [{"check_name": "kite_token_fresh_today", "check_group": "Broker",
                    "criticality": "CRITICAL", "status": "FAIL"}]})
    out = sh._live_readiness(pfv, [_svc("Trading Engine", "HEALTHY", "engine"),
                                   _svc("Database", "HEALTHY", "database")])
    assert out["ready"] is False and out["label"] == "NOT READY"
    assert "kite_token_fresh_today" in " ".join(out["blockers"])


def test_readiness_transitions_across_the_live_lifecycle():
    """Rama's §15 transition list, driven end to end through ONE verdict fn."""
    pfv = sh._preflight_verdict(_PF_READY)

    def verdict(engine_status):
        return sh._live_readiness(pfv, [_svc("Trading Engine", engine_status, "engine"),
                                        _svc("Database", "HEALTHY", "database")])

    assert verdict("HEALTHY")["label"] == "READY"
    warned = verdict("WARNING")
    assert warned["label"] == "READY" and warned["verdict"] == "WARNING"  # degraded ≠ blocked
    assert verdict("FAILED")["label"] == "NOT READY"
    assert verdict("UNKNOWN")["label"] == "NOT CONFIRMED"
    assert verdict("HEALTHY")["label"] == "READY"                        # recovered


def test_the_required_set_is_published_so_the_verdict_is_auditable(client):
    """⛔ A verdict whose basis is buried in code cannot be checked by an operator."""
    rd = _h(client)["readiness"]
    named = {r["service"] for r in rd["required_services"]}
    assert "Trading Engine" in named and "Database" in named
    assert "Dashboard API" not in named      # the GUI never gates trading
    assert rd["basis"] and "BOTH gates" in rd["basis"]


def test_the_five_pillars_are_evaluated_from_real_checks(client):
    """Broker/Database/Engine/State come from `check_group`; Capital has NO group
    of its own and is assembled from named checks — so it proves both paths."""
    rd = _h(client)["readiness"]["preflight"]
    by = {p["pillar"]: p for p in rd["pillars"]}
    assert set(by) == {"Broker", "Database", "Core Services", "Capital", "Risk"}
    assert by["Broker"]["status"] == "HEALTHY" and by["Broker"]["checks"] == 5
    assert by["Database"]["status"] == "HEALTHY" and by["Database"]["checks"] == 2
    assert by["Risk"]["status"] == "HEALTHY" and by["Risk"]["checks"] == 2
    # capital_deployment is seeded WARN ⇒ Capital is WARNING while others are not
    assert by["Capital"]["status"] == "WARNING"
    assert by["Capital"]["checks"] == 3
    assert by["Capital"]["note"] and "no Capital group" in by["Capital"]["note"]
    assert rd["corroborated"] is True and rd["checks_evaluated"] > 0


def test_readiness_without_a_preflight_run_is_not_ready_by_default(client):
    """⛔ Absence of a verdict is not a READY verdict."""
    out = sh._preflight_verdict(None)
    assert out["ready"] is None
    assert out["verdict"] == "UNKNOWN"
    assert "absence of a verdict is not a READY verdict" in out["reason"]
    assert all(p["status"] == "UNKNOWN" for p in out["pillars"])
    # ...and it stays NOT CONFIRMED once the live gate is applied
    live = sh._live_readiness(out, [_svc("Trading Engine", "HEALTHY", "engine")])
    assert live["ready"] is None and live["label"] == "NOT CONFIRMED"


def test_a_ready_verdict_with_no_checks_is_degraded_not_trusted():
    """⭐ A verdict nothing could have contradicted is not evidence. The system's
    own status is still reported, but the badge is degraded and says why."""
    out = sh._preflight_verdict({"overall_status": "READY", "checks": [],
                                 "started_at": "2026-08-15T08:30:00", "phase": "A"})
    assert out["ready"] is True                 # ⛔ not rewritten to False
    assert out["verdict"] == "WARNING"          # but not shown as a clean pass
    assert out["corroborated"] is False
    assert "could not be corroborated" in out["reason"]
    # ⭐ and the uncorroborated warning survives INTO the live verdict
    live = sh._live_readiness(out, [_svc("Trading Engine", "HEALTHY", "engine")])
    assert live["ready"] is True and live["verdict"] == "WARNING"


def test_a_critical_failure_makes_it_not_ready():
    out = sh._preflight_verdict({
        "overall_status": "CRITICAL_FAILURE", "phase": "B",
        "started_at": "2026-08-15T09:14:00",
        "checks": [{"check_name": "kite_token_fresh_today", "check_group": "Broker",
                    "criticality": "CRITICAL", "status": "FAIL"}]})
    assert out["ready"] is False
    assert out["verdict"] == "FAILED"
    assert "kite_token_fresh_today" in " ".join(out["blockers"])
    by = {p["pillar"]: p for p in out["pillars"]}
    assert by["Broker"]["status"] == "FAILED"


# ── AUTO-RECOVERY: triggered is not success ─────────────────────────────────
def test_recovery_distinguishes_triggered_success_and_failed(client):
    """⛔ A TRIGGERED attempt must never be counted as a recovery."""
    rec = _h(client)["recovery"]
    assert rec["success"] == 1 and rec["failed"] == 1 and rec["triggered"] == 1
    by = {e["state"]: e for e in rec["events"]}
    assert by["SUCCESS"]["status"] == "HEALTHY"
    assert by["FAILED"]["status"] == "FAILED"
    # ⭐ in-flight is WARNING, ⛔ not green — the component is not fixed yet
    assert by["TRIGGERED"]["status"] == "WARNING"
    assert "never counted as a recovery" in rec["note"]


def test_a_failed_recovery_carries_its_error(client):
    ev = [e for e in _h(client)["recovery"]["events"] if e["state"] == "FAILED"][0]
    assert ev["error"] == "permission denied"


# ── DEPENDENCIES ────────────────────────────────────────────────────────────
def test_the_five_approved_dependencies_are_present(client):
    names = [d["name"] for d in _h(client)["dependencies"]]
    assert names == ["Chartink", "Broker API", "Database", "Tailscale", "Internet"]


def test_a_dependency_is_never_healthy_merely_because_it_is_configured(client):
    """⛔ Internet has no probe anywhere, so it must be UNKNOWN with a reason —
    ⛔ not quietly inherited from the broker check."""
    by = {d["name"]: d for d in _h(client)["dependencies"]}
    assert by["Internet"]["status"] == "UNKNOWN"
    assert by["Internet"]["measured"] is False
    assert "no internet-reachability probe" in by["Internet"]["source"]
    # the one with a real live probe IS measured — the control
    assert by["Database"]["measured"] is True and by["Database"]["status"] == "HEALTHY"


# ── SERVICES / BROKER / UPTIME ──────────────────────────────────────────────
def test_no_fake_services_are_invented(client):
    """The reference's example list names engines that are THREADS inside one
    unit. Only real, sourced services appear."""
    d = _h(client)
    kinds = {s["kind"] for s in d["services"]}
    assert kinds <= {"systemd", "engine", "engine-component", "database", "dashboard"}
    names = {s["service"] for s in d["services"]}
    assert "Database" in names and "Dashboard API" in names and "Trading Engine" in names
    assert "trading-system.service" in names


def test_response_time_is_absent_where_nothing_is_timed(client):
    """⛔ `systemctl` latency is the dashboard's own subprocess cost, not the
    service's responsiveness — so a systemd row carries no response time."""
    for s in _h(client)["services"]:
        if s["kind"] == "systemd":
            assert s["response_ms"] is None
            assert "N/A" in s["response_note"] or "no request" in s["response_note"]
        if s["kind"] == "database":
            assert isinstance(s["response_ms"], (int, float))


def test_broker_login_and_token_are_separate_facts(client):
    """⛔ Collapsing them would hide a token that EXISTS but is STALE — which
    presents as a silent no-trade morning."""
    bk = _h(client)["broker"]
    for k in ("broker_status", "login_status", "token_status"):
        assert bk[k] in ("HEALTHY", "WARNING", "FAILED", "UNKNOWN")
    assert set(bk["token_detail"]) == {"file_exists", "fresh_today"}
    assert bk["last_api_call"]["measured"] is False


def test_uptime_does_not_come_from_the_dashboard(client):
    """⛔ Uptime must not reset when this page refreshes: it is the trading
    engine's own monotonic clock, read from /health."""
    up = _h(client)["uptime"]
    assert "monotonic" in up["source"]
    assert up["measured"] is False and up["seconds"] is None   # engine is down here
    assert "not answering" in up["reason"]


def test_uptime_formatting_never_fabricates():
    assert sh._fmt_uptime(None) is None
    assert sh._fmt_uptime(0) == "0m"
    assert sh._fmt_uptime(90) == "1m"
    assert sh._fmt_uptime(3700) == "1h 01m"
    assert sh._fmt_uptime(7 * 86400 + 14 * 3600 + 32 * 60) == "7d 14h 32m"


def test_systemd_uptime_refuses_an_unusable_stamp():
    """⛔ None, never 0 — a zero uptime reads as "just restarted", which is a
    materially different and alarming statement from "not known"."""
    from backend.readers import host_reader
    assert host_reader._uptime_from_systemd_stamp("") is None
    assert host_reader._uptime_from_systemd_stamp("n/a") is None
    assert host_reader._uptime_from_systemd_stamp("Fri 2026-13-45 99:99:99 IST") is None


# ── FILTERS / VIEW ALL ──────────────────────────────────────────────────────
def test_the_filter_is_a_view_and_the_counts_stay_whole_population(client):
    """This screen is a LIVE SNAPSHOT: filtering the table must not make the KPI
    deck describe a subset, or an operator would lose the failure they filtered
    away from. `services_all` and the counts stay whole-population."""
    a, b = _h(client), _h(client, "?status=UNKNOWN")
    assert len(b["services"]) < len(a["services"])
    assert all(s["status"] == "UNKNOWN" for s in b["services"])
    assert b["counts"] == a["counts"]
    assert b["services_all"] == a["services_all"] == len(a["services"])


def test_kind_filter_narrows_the_table(client):
    d = _h(client, "?kind=systemd")
    assert d["services"] and all(s["kind"] == "systemd" for s in d["services"])


# ── EXPORT ──────────────────────────────────────────────────────────────────
def test_export_carries_every_panel_and_respects_the_view(client):
    from openpyxl import load_workbook
    r = client.get("/api/export/system-health?status=UNKNOWN")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Services", "Readiness", "Dependencies", "Alerts",
                             "Auto-Recovery", "Service Events", "Throughput",
                             "Summary"]
    ws = wb["Services"]
    header = [c.value for c in ws[1]]
    assert header[:3] == ["Service", "Kind", "Status"]
    body = list(ws.iter_rows(min_row=2, values_only=True))
    assert body and {r[2] for r in body} == {"UNKNOWN"}


def test_export_spells_out_the_gaps(client):
    """⛔ A blank cell would let a reader assume the value was simply zero."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get("/api/export/system-health").data))
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    # ⭐ CPU/RAM/Network are LIVE metrics read from /proc — on this non-Linux
    # host the source is absent, so NO DATA, ⛔ never a blank and ⛔ never zero.
    assert summary["CPU %"] == "NO DATA"
    assert summary["RAM %"] == "NO DATA"
    assert summary["Network"] == "NO DATA"
    assert summary["DB connection count"] == "NOT APPLICABLE"
    assert "/proc/stat" in summary["CPU % why"]
    assert summary["Overall Status"] == "FAILED"


def test_export_readiness_sheet_carries_the_pillars(client):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get("/api/export/system-health").data))
    rows = list(wb["Readiness"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["Broker", "Database", "Core Services",
                                    "Capital", "Risk"]


# ── THE SCREEN ITSELF ───────────────────────────────────────────────────────
def test_route_renders(client):
    r = client.get("/services")
    assert r.status_code == 200
    body = r.data.decode("utf-8")
    assert "sysh-page" in body and "System Health" in body


def test_existing_system_endpoints_are_untouched(client):
    """Screen 12 is ADDITIVE — /api/services and /api/vm keep their contracts."""
    d = client.get("/api/services").get_json()
    assert "units" in d and "cron" in d and "trader" in d
    v = client.get("/api/vm").get_json()
    assert v["cpu_ram_history"]["collected"] is False


def test_real_time_refresh_is_preserved():
    t = _tpl()
    assert '@ops-refresh.window="refresh()"' in t
    body = re.search(r"refresh\(\)\s*\{(.*?)\}", t, re.S).group(1)
    assert "tPage" not in body and "applied" not in body and "cols" not in body


def test_columns_use_the_existing_drag_convention():
    t = _tpl()
    for token in ('draggable="true"', "@dragstart=", "@drop.prevent=", "onDragStart(",
                  "onDrop(", "initCols()", "saveCols()", "resetCols()",
                  "is-drag", "is-over", "screen12.systemhealth.colOrder"):
        assert token in t, token
    body = re.search(r"headClick\(key\)\s*\{(.*?)\n    \},", t, re.S).group(1)
    assert "_dragEndAt" in body and "250" in body


def test_alignment_is_by_column_role_not_blanket():
    css = _css()
    assert ".sysh-page .cap-table th { text-align: center; }" in css
    assert ".sysh-page .cap-table th.lbl { text-align: left; }" in css
    assert ".sysh-page table td { text-align: center" not in css


def test_no_alpine_template_loop_inside_an_svg():
    """⛔ Pinned across the whole tree since Screen 10."""
    import glob
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    offenders = []
    for path in glob.glob(os.path.join(here, "frontend", "templates", "*.html")):
        html = open(path, encoding="utf-8").read()
        for svg in re.findall(r"<svg\b.*?</svg>", html, re.S):
            if re.search(r"<template\b", svg):
                offenders.append(os.path.basename(path))
    assert not offenders, sorted(set(offenders))


def test_no_hard_coded_values_in_the_markup():
    t = _tpl()
    for ghost in ("7d 14h 32m", "18 / 100", "256.4 GB", "120 Mbps", "0.68", "99.82"):
        assert ghost not in t, ghost


def test_css_is_scoped_to_this_screen():
    block = _css().split("SCREEN 12 — SYSTEM HEALTH", 1)[1].split("*/", 1)[1]
    block = re.split(r"SCREEN \d+ [—-]", block, maxsplit=1)[0]
    block = re.sub(r"/\*.*?\*/", " ", block, flags=re.S)
    rules = re.findall(r"([^{}]+)\{([^{}]*)\}", block)
    assert len(rules) > 40
    for selector, _body in rules:
        selector = selector.strip()
        if selector.startswith("@media"):
            continue
        for part in selector.split(","):
            part = part.strip().lstrip("{").strip()
            if part:
                assert part.startswith(".sysh-page"), part


def test_other_screens_are_untouched():
    css = _css()
    for sel in (".cap-page .cap-table th { text-align: center; }",
                ".pnl-page .cap-table th { text-align: center; }",
                ".slp-page .cap-table th { text-align: center; }",
                ".exec-page .cap-table th { text-align: center; }",
                "table td.ctr { text-align: center !important; }"):
        assert sel in css, sel


# ══════════════════════════════════════════════════════════════════════════
# HEALTH-TRENDS X-AXIS — THE LAST-LABEL COLLISION (measured 19-Aug-2026)
#
# ⭐ REACHABILITY FIRST, because a fix to an unreachable path is not worth a
#   test: on the VM `analytics.system_metrics` holds 3,651 rows of which 3,069
#   carry a REAL `disk_used_pct`, and ALL 60 rows in the window the GUI reads
#   are real (cpu_pct / memory_mb are -1.0 sentinels, response_time is never
#   persisted). ⇒ the Disk trend renders a full chart in PRODUCTION TODAY, so
#   this axis is LIVE, ⛔ not latent.
#
# THE DEFECT, MEASURED IN THE BROWSER at 1920x1080 AND 1440x900 with the
# production window (60 snapshots ⇒ step 8 ⇒ stepped marks …48, 56 plus a
# FORCED final 59):
#     '12:27' ended at x=1004.0 and '12:42' began at x=998.6
#     ⇒ a 5.4px OVERLAP, identical at both viewports.
# After the fix: 8 x-labels, minimum gap 37.1px, the END mark retained.
#
# ⭐ This is the SAME defect Screen 11's throughput axis already fixed and
#   pinned (`test_s1011_chart_labels.py`); the rule is carried, ⛔ not
#   re-invented.
# ══════════════════════════════════════════════════════════════════════════
class TestHealthTrendXAxisCollision:

    def test_the_clash_avoidance_rule_is_present_in_trendmarkup(self):
        """⛔ Both halves, or the rule is not the rule: `marks.pop()` drops the
        clashing stepped label and `Math.ceil(step / 2)` is the minimum gap."""
        body = _tpl()
        body = body[body.index("trendMarkup() {"):]
        body = body[:body.index("\n    },")]
        # ⛔ STRIP COMMENTS BEFORE THE NEGATIVE ASSERTION. The block comment that
        # documents this fix QUOTES the defective pattern verbatim, so matching
        # raw source fails on the EXPLANATION rather than on the code — this
        # test caught exactly that on its own first run.
        code = re.sub(r"/\*.*?\*/", "", body, flags=re.S)
        assert "marks.pop()" in code, "the clash-avoidance was removed"
        assert "Math.ceil(step / 2)" in code, "the minimum-gap rule was removed"
        assert "i === s.length - 1" not in code, (
            "the unguarded always-show-last rule is back — that IS the defect")

    def test_the_rule_actually_drops_the_clashing_mark_on_the_production_case(self):
        """Re-implements the rule and runs it on the MEASURED production window.
        ⛔ A rule that never fires would be no rule at all, so this asserts the
        drop happens — not merely that the code contains a branch."""
        n, xmax = 60, 8                      # the exact window the GUI reads
        step = max(1, -(-n // xmax))         # ceil -> 8
        last = n - 1                         # 59
        marks = list(range(0, last + 1, step))
        assert marks[-1] == 56, "precondition changed: the stepped tail is not 56"
        if marks[-1] != last:
            if last - marks[-1] < -(-step // 2):
                marks.pop()
            marks.append(last)
        assert last in marks, "the session-end mark must survive"
        assert 56 not in marks, "the clashing stepped mark should have been dropped"
        gaps = [b - a for a, b in zip(marks, marks[1:])]
        assert min(gaps) >= -(-step // 2), "two marks are still closer than half a step"

    def test_the_rule_leaves_a_well_spaced_tail_alone(self):
        """⛔ The guard must not fire when there is no clash — otherwise it would
        silently delete a legitimate mark. 64 points ⇒ step 8 ⇒ tail 56, last 63,
        gap 7 >= 4, so BOTH survive."""
        n, xmax = 64, 8
        step = max(1, -(-n // xmax))
        last = n - 1
        marks = list(range(0, last + 1, step))
        tail = marks[-1]
        if marks[-1] != last:
            if last - marks[-1] < -(-step // 2):
                marks.pop()
            marks.append(last)
        assert tail in marks, "a well-spaced stepped mark must NOT be dropped"
        assert last in marks

    def test_the_axis_type_was_not_shrunk_to_buy_the_space(self):
        """⛔ Fixing a collision by going under the floor trades one spec
        violation for another. `.sysh-axis` is declared 13px and rendered 13.00px
        at BOTH viewports (viewBox 0 0 560 180 against a 180px-high box ⇒ the
        `xMidYMid meet` scale is pinned at 1.0 by the HEIGHT, so width changes
        cannot shrink it)."""
        css = _css()
        m = re.search(r"\.sysh-page \.sysh-axis \{([^}]*)\}", css)
        assert m, ".sysh-axis rule missing"
        fm = re.search(r"font-size:\s*([0-9.]+)px", m.group(1))
        assert fm and float(fm.group(1)) >= 13, m.group(1)
        assert 'viewBox="0 0 560 180"' in _tpl(), (
            "the trend viewBox changed — the 1.0 scale that keeps 13px "
            "rendered is no longer guaranteed; re-measure before editing this")


# ═════════════════════════════════════════════════════════════════════════════
# THE 30-Aug PASS AGAINST THE ORIGINAL ARTWORK
# ⛔ These pin the APPROVED COMPOSITION, ⛔ not merely "some markup exists".
# ═════════════════════════════════════════════════════════════════════════════
def test_the_top_band_is_five_cards_and_unknown_is_not_one_of_them(client):
    """⭐ Overall · Healthy · Warning · Failed · Uptime — and ⛔ NO Unknown card.

    ⚠️ Unknown is a REAL service status and must not vanish from the screen;
    what it must not be is a headline KPI. So BOTH halves are asserted: gone
    from the band, still present where the distribution is actually read.
    """
    t = _tpl()
    band = t[t.index('class="kpi-row sysh-kpis"'):t.index("sysh-uptime")]
    for label in ("Healthy Services", "Warning Services", "Failed Services"):
        assert label in band
    assert "cnt('UNKNOWN')" not in band          # ⛔ the sixth card, in any form
    assert 'kpi_card("Unknown"' not in t
    # ⭐ BUT UNKNOWN SURVIVES EVERYWHERE IT CARRIES MEANING. ⛔ Dropping the
    # card must not quietly drop the STATE: it stays in the legend, in the
    # colour map, in the counts and in the status filter's own options.
    assert '<i class="dot sysh-dot-UNKNOWN"></i>Unknown' in t
    assert 'UNKNOWN: "sysh-unknown"' in t
    d = _h(client, "")
    assert "UNKNOWN" in d["counts"]
    assert "UNKNOWN" in d["filters"]["options"]["status"]


def test_every_band_card_carries_its_icon_and_its_share():
    t = _tpl()
    assert t.count("sysh-kpi-ico") == 3          # one per count card
    assert t.count("sysh-kpi-pct") == 3
    assert "sysh-ov-ico" in t                    # the artwork's overall shield
    for st in ("HEALTHY", "WARNING", "FAILED"):
        assert "pctOf('%s')" % st in t


def test_a_share_of_an_unknown_fleet_is_a_dash_not_zero_percent():
    """⛔ 0.00% of nothing is a fabricated denominator, ⭐ so the guard is real."""
    js = _tpl()[_tpl().index("pctOf(status)"):]
    assert "if (!total) return" in js.split("},")[0]
    assert "u2014" in js.split("},")[0]          # an em dash, ⛔ not "0.00%"


def test_the_vm_metrics_are_nested_in_the_uptime_card_and_the_rail_is_gone():
    """⭐ The artwork nests them; the separate rail panel no longer exists.

    ⛔ The point is not that a strip exists somewhere — it is that the metrics
    moved INSIDE the uptime card and that NOTHING was lost on the way.
    """
    t = _tpl()
    assert 'class="sysh-rail"' not in t
    assert "VM / System Health" not in t
    card = t[t.index('class="panel sysh-uptime"'):
             t.index("</section>", t.index("sysh-vmstrip"))]
    for field in ("cpu_pct", "ram_pct", "ram_available_mb", "disk_pct",
                  "network", "system_load"):
        assert "vm('%s')" % field in card, field
    # ⛔ still routed through gapCell — a moved metric must not become a bare
    # number, which is exactly how a NOT INSTRUMENTED value would turn green.
    assert card.count("gapCell(vm(") == 6


def test_dependencies_are_five_icon_cards_not_a_list():
    t = _tpl()
    assert 'class="sysh-deps"' not in t          # the old vertical list is gone
    assert "sysh-depcards" in t and "sysh-depcard-ico" in t
    assert "depIcon(d.name)" in t
    # ⭐ the icon map covers exactly the names the BACKEND emits, so a renamed
    # dependency cannot silently render an empty box
    names = {d["name"] for d in sh._dependencies(
        {"units": []}, None, [], {"reachable": True}, "2026-08-30")}
    assert len(names) == 5
    icons = t[t.index("DEP_ICONS:"):t.index("depIcon(name)")]
    for n in names:
        assert '"%s":' % n in icons, n


def test_the_bottom_export_bar_is_wired_to_the_real_route():
    """⛔ Not a gated placeholder — S12 owns a real export route."""
    t = _tpl()
    assert "sysh-exportbar" in t
    assert "Export Current View" in t and "Export XLSX" in t
    assert "/api/export/system-health" in t


# ── THROUGHPUT — in the PNG, absent from the design TXT ─────────────────────
def test_throughput_reads_the_existing_pulse_and_buckets_it(monkeypatch):
    monkeypatch.setattr(sh.db_reader, "activity_pulse", lambda cfg, today: {
        "signals": {"09:15": 3, "09:47": 2, "10:05": 4},
        "orders":  {"09:20": 1, "10:59": 5},
        "trades":  {"10:00": 2},
    })
    tp = sh._throughput({}, "2026-08-30")
    assert tp["measured"] is True and tp["bucket"] == "hour"
    assert tp["series"] == [
        {"label": "09:00", "signals": 5, "orders": 1, "fills": 0},
        {"label": "10:00", "signals": 4, "orders": 5, "fills": 2},
    ]


def test_an_empty_throughput_is_a_real_zero_never_an_instrumentation_gap(monkeypatch):
    """⛔ THE WHOLE POINT OF THE PANEL'S HONESTY. Saying NOT INSTRUMENTED over
    three tables the system demonstrably writes would be a false claim about
    the system itself — the opposite failure to charting a sentinel."""
    monkeypatch.setattr(sh.db_reader, "activity_pulse", lambda cfg, today: {})
    tp = sh._throughput({}, "2026-08-30")
    assert tp["measured"] is False
    assert tp["instrumented"] is True            # ⛔ never False
    assert "REAL ZERO" in tp["reason"]
    assert "NOT INSTRUMENTED" not in tp["reason"]


def test_the_third_series_is_the_fill_not_an_order_acknowledgement():
    """⚠️ `trades.entry_time` — a trade row exists once the ENTRY filled.
    Reading `orders.status` would count a broker ACK as a fill."""
    src = inspect.getsource(sh._throughput)
    assert "entry_time" in src
    assert "orders.status" not in src.split('"""')[2]   # ⛔ not in the code body


def test_the_throughput_panel_renders_no_activity_not_not_instrumented():
    t = _tpl()
    panel = t[t.index('<h3 class="panel-title">Throughput'):t.index("Status Guide")]
    assert "NO ACTIVITY" in panel
    assert "NOT INSTRUMENTED" not in panel
    assert "pulseMarkup()" in panel
    for legend in ("Signals Received", "Orders Created", "Orders Filled"):
        assert legend in panel


def test_throughput_is_in_the_payload_and_in_the_workbook(client):
    d = _h(client, "")
    assert "throughput" in d and d["throughput"]["instrumented"] is True
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get("/api/export/system-health").data))
    assert [c.value for c in wb["Throughput"][1]] == [
        "Bucket", "Signals Received", "Orders Created", "Orders Filled"]
