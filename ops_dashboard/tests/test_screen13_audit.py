"""SCREEN 13 — AUDIT.

The approved design (`gui/13. Audit.png` + `.txt`) is binding for STRUCTURE; the
DATA is this system's. These tests hold both lines: every approved element is
present, and nothing this system fails to record is invented to fill it.

⭐ THE CENTRAL PROPERTY: an accountability screen that shows a plausible-looking
value it did not measure is worse than one that shows a gap. Every "who / when /
from where / what changed" that the schema cannot answer must render as an
explicit gap, and a test here must fail if it ever silently becomes a value.
"""
from __future__ import annotations

import inspect
import io
import json
import os
import re

from backend.services import audit


def _read(*parts) -> str:
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    with open(os.path.join(here, *parts), encoding="utf-8") as fh:
        return fh.read()


def _tpl() -> str:
    return _read("frontend", "templates", "audit.html")


def _css() -> str:
    return _read("frontend", "static", "style.css")


def _shipped(client) -> str:
    """The HTML the browser actually receives, with comments stripped — a Jinja
    `{# #}` never ships and an HTML comment is not UI, so neither may satisfy an
    assertion about what the operator can see."""
    html = client.get("/audit").data.decode("utf-8")
    return re.sub(r"<!--.*?-->", " ", html, flags=re.S)


def _a(client, qs=""):
    return client.get("/api/audit-screen" + qs).get_json()


WIDE = "?start=2000-01-01&end=2099-12-31"


# ── the approved vocabularies ───────────────────────────────────────────────
def test_the_eight_approved_categories_are_the_vocabulary():
    assert audit.CATEGORIES == ("Configuration", "Control", "Strategy", "Risk",
                                "Capital", "Service", "Authentication", "System")
    assert audit.STATUSES == ("Success", "Failed", "Partial")
    assert audit.SOURCES == ("Dashboard", "System", "API", "Scheduler",
                             "Recovery Engine")
    assert audit.TIMELINE_STAGES == ("Requested", "Validated", "Applied", "Confirmed")


def test_the_nine_approved_columns_in_order_and_no_scanner():
    body = _tpl()
    cols = re.search(r"DEFAULT_COLS:\s*\[(.*?)\n    \],", body, re.S).group(1)
    keys = re.findall(r'key:\s*"([a-z_]+)"', cols)
    assert keys == ["date", "time", "category", "module", "action", "user",
                    "status", "source", "ref_id"]
    labels = re.findall(r'label:\s*"([^"]+)"', cols)
    assert labels == ["Date", "Time", "Category", "Module", "Action", "User",
                      "Status", "Source", "Reference ID"]


def _page_content(client) -> str:
    """The AUDIT screen's VISIBLE markup — shared chrome and scripts removed.

    ⛔ Two things must not be scanned, and both would produce a false failure
    that tempts a wrong "fix":
      · the sidebar's `/scanner-attribution` nav link, which belongs to a
        DIFFERENT, untouched screen and must survive;
      · script bodies, whose comments state *"⛔ no Scanner"* — documentation of
        the rule is not a breach of it, and deleting the comment to satisfy a
        grep would remove the reason the rule is kept.
    """
    body = _shipped(client)
    body = re.sub(r"<script\b.*?</script>", " ", body, flags=re.S)
    i = body.find('class="aud-page"')
    assert i > 0, "the audit page root is missing"
    return body[i:]


def test_scanner_appears_nowhere_on_this_screen(client):
    """⛔ Project-wide: KEEP STRATEGY, REMOVE SCANNER. The reader must not even
    SELECT `scanner_name`, so the value cannot leak through a payload later."""
    assert "scanner" not in _page_content(client).lower()
    # ⭐ and the nav link to the separate Scanner Attribution screen is UNTOUCHED
    assert "/scanner-attribution" in _shipped(client)
    from backend.readers import db_reader
    src = inspect.getsource(db_reader.audit_webhook_range)
    sql = src[src.index("SELECT"):src.index("ORDER BY")]
    assert "scanner_name" not in sql
    # ⭐ and Strategy IS retained — it is an approved category, not a casualty
    assert "Strategy" in audit.CATEGORIES
    payload = _a(client, WIDE)
    assert not any("scanner" in json.dumps(r).lower() for r in payload["records"])


# ── the six approved KPI ────────────────────────────────────────────────────
def test_the_six_approved_kpi_exist_with_the_approved_wording(client):
    k = _a(client, WIDE)["kpi"]
    for key in ("total_events", "todays_changes", "configuration_changes",
                "control_actions", "failed_actions", "last_event"):
        assert key in k, key
    body = _tpl()
    for label in ("TOTAL AUDIT EVENTS", "TODAY'S CHANGES", "CONFIGURATION CHANGES",
                  "CONTROL ACTIONS", "FAILED ACTIONS", "LAST AUDIT EVENT"):
        assert label in body, label


def test_kpi_counts_are_computed_over_the_same_population_as_the_table(client):
    """⛔ A KPI computed over a different population than the table it sits above
    is a lie that looks like arithmetic."""
    p = _a(client, WIDE)
    assert p["kpi"]["total_events"] == len(p["records"])
    today = p["today"]
    assert p["kpi"]["todays_changes"] == sum(
        1 for r in p["records"] if r["date"] == today)
    assert p["kpi"]["failed_actions"] == sum(
        1 for r in p["records"] if r["date"] == today and r["status"] == "Failed")


# ── ⭐ OLD → NEW: real, derived from consecutive config snapshots ────────────
def test_old_new_values_are_derived_from_real_consecutive_snapshots(client):
    """⭐⭐ THE APPROVED CENTREPIECE, AND IT IS REAL. The fixture holds two
    snapshots whose only difference is `system.risk.max_daily_trades`. The diff
    must surface exactly that, with both sides, ⛔ nothing invented."""
    recs = _a(client, WIDE)["records"]
    diffs = [r for r in recs if r["src_table"] == "config_snapshots"]
    assert diffs, "the snapshot diff produced no change at all"
    hit = [r for r in diffs if r["field"] == "system.risk.max_daily_trades"]
    assert len(hit) == 1, [r["field"] for r in diffs]
    r = hit[0]
    assert r["old_value"] == "9"          # the earlier snapshot's real value
    assert r["new_value"] is not None and r["new_value"] != r["old_value"]
    assert r["action"] == "max_daily_trades changed"
    # ⭐ classified from the REAL config path, so a risk setting reads as Risk
    assert r["category"] == "Risk" and r["module"] == "Risk Engine"


def test_a_change_is_never_inferred_without_two_snapshots_proving_it():
    """⛔ One snapshot cannot evidence a change. It must yield ZERO, not a guess."""
    one = [{"snapshot_id": 1, "snapshot_ts": "2026-08-01T08:15:00+05:30",
            "snapshot_date": "2026-08-01", "config_hash": "a" * 8,
            "config_json": json.dumps({"system": {"risk": {"x": 1}}})}]
    assert audit._config_changes(one) == []
    # identical consecutive snapshots are not a change either
    same = one + [dict(one[0], snapshot_id=2, snapshot_ts="2026-08-02T08:15:00+05:30")]
    assert audit._config_changes(same) == []

    # ⭐ AND THE CASE THE HASH ALONE WOULD GET WRONG: a DIFFERENT hash whose only
    # difference is `file_hashes` — a bookkeeping field, not a setting anybody
    # changed. The value-level comparison is what protects here, ⛔ not the hash
    # short-circuit, so a reported "change" would be an invented one.
    a = {"system": {"risk": {"x": 1}}, "file_hashes": {"a.yaml": "aaa"}}
    b = {"system": {"risk": {"x": 1}}, "file_hashes": {"a.yaml": "bbb"}}
    pair = [{"snapshot_id": 1, "snapshot_ts": "2026-08-01T08:15:00+05:30",
             "snapshot_date": "2026-08-01", "config_hash": "h1",
             "config_json": json.dumps(a)},
            {"snapshot_id": 2, "snapshot_ts": "2026-08-02T08:15:00+05:30",
             "snapshot_date": "2026-08-02", "config_hash": "h2",
             "config_json": json.dumps(b)}]
    assert audit._config_changes(pair) == []

    # ...while a REAL setting change on the same pair IS reported, so the filter
    # above cannot be hiding everything.
    b2 = {"system": {"risk": {"x": 2}}, "file_hashes": {"a.yaml": "bbb"}}
    pair[1] = dict(pair[1], config_json=json.dumps(b2))
    got = audit._config_changes(pair)
    assert len(got) == 1 and got[0]["field"] == "system.risk.x"
    assert (got[0]["old_value"], got[0]["new_value"]) == ("1", "2")


def test_config_paths_classify_into_the_approved_categories():
    """⭐ Mapping is by REAL config section, ⛔ not by guessing at a name."""
    assert audit._classify_path("system.risk.max_daily_loss_pct")[0] == "Risk"
    assert audit._classify_path("system.capital.starting")[0] == "Capital"
    assert audit._classify_path("system.portfolio_allocator.x")[0] == "Capital"
    assert audit._classify_path("system.entry_gate.min_score")[0] == "Strategy"
    assert audit._classify_path("system.order_monitor.poll")[0] == "Service"
    # ⛔ anything unmapped stays plain Configuration — never forced into a bucket
    assert audit._classify_path("broker_costs.zerodha.x")[0] == "Configuration"
    # longest prefix wins, so system.risk beats a bare system
    assert audit._classify_path("system.risk.a")[1] == "Risk Engine"


# ── ⛔ the gaps: nothing unrecorded may become a value ───────────────────────
def test_user_attribution_is_a_declared_gap_not_an_invented_name(client):
    """⛔ auth.py touches NO database, so login/logout is never persisted and
    there is no human attribution anywhere. A record with no actor must carry
    None — ⛔ never 'admin', never the module name, never a guess."""
    p = _a(client, WIDE)
    assert p["gaps"]["user"]["measured"] is False
    assert "auth.py" in p["gaps"]["user"]["reason"]
    assert p["gaps"]["authentication"]["measured"] is False
    # every actor that DOES exist is a process, and is labelled as one
    for r in p["records"]:
        if r["user"] is not None:
            assert r["user_kind"] == "process", r
            assert r["user"] in ("system", "order_monitor"), r["user"]
    # ⛔ and the config diffs — whose author is genuinely unknown — claim nobody
    for r in p["records"]:
        if r["src_table"] == "config_snapshots":
            assert r["user"] is None, "the editor of a YAML file is not recorded"


def test_config_diff_events_do_not_credit_the_system_as_the_actor(client):
    """⭐ A subtle one: the system DETECTED the config change, it did not MAKE
    it. Crediting 'system' would fabricate accountability for a human edit."""
    recs = _a(client, WIDE)["records"]
    cd = [r for r in recs if r["action"].startswith("Configuration change detected")]
    assert cd, "fixture seeds a CONFIG_DIFF event"
    assert all(r["user"] is None for r in cd)
    # ...while a genuine system action DOES name the system
    st = [r for r in recs if r["action"].startswith("System started")]
    assert st and all(r["user"] == "system" for r in st)


def test_source_is_a_gap_except_where_the_record_proves_its_own_origin(client):
    """⛔ Dashboard/Scheduler/Recovery-Engine origins are not recorded. The ONE
    genuine value is the webhook: that row exists BECAUSE an HTTP POST arrived."""
    p = _a(client, WIDE)
    assert p["gaps"]["source"]["measured"] is False
    for r in p["records"]:
        assert r["source"] in (None, "API"), r["source"]
        if r["source"] == "API":
            assert r["src_table"] == "webhook_audit"


def test_the_timeline_marks_only_applied_as_real(client):
    """⛔ Rama: "do not silently convert missing stages into successful stages."
    Only one instant per record is recorded."""
    p = _a(client, WIDE)
    rec = p["records"][0]
    tl = audit._timeline(rec)
    assert [s["stage"] for s in tl] == list(audit.TIMELINE_STAGES)
    by = {s["stage"]: s for s in tl}
    assert by["Applied"]["measured"] is True and by["Applied"]["ts"] == rec["ts"]
    for stage in ("Requested", "Validated", "Confirmed"):
        assert by[stage]["measured"] is False, stage
        assert by[stage]["ts"] is None, stage
        assert by[stage]["note"], stage


def test_retention_period_is_a_gap_and_the_span_is_not_sold_as_a_policy(client):
    """⛔ No audit retention policy exists. `signal_retention_days` governs
    SIGNAL fingerprints; showing it here would attribute a policy nobody wrote."""
    r = _a(client, WIDE)["retention"]
    assert r["retention_period"]["measured"] is False
    assert "signal_retention_days" in r["retention_period"]["reason"]
    # the observed span IS real, and is labelled a measurement
    assert "MEASUREMENT" in r["span_note"] or "measurement" in r["span_note"]
    assert r["records"] >= 1 and r["oldest"] and r["newest"]


def test_retention_describes_the_store_not_the_current_filter(client):
    """⛔ A retention figure that shrank when the operator narrowed a date range
    would not be a retention figure."""
    wide = _a(client, WIDE)["retention"]
    narrow = _a(client, "?start=2099-01-01&end=2099-01-02")["retention"]
    assert narrow["records"] == wide["records"]
    assert narrow["oldest"] == wide["oldest"]


def test_control_history_gap_is_declared(client):
    """⚠️ kill_switch_state is CHECK(id=1) — one row. It is the CURRENT state,
    ⛔ not a history of control actions, and the screen says so."""
    p = _a(client, WIDE)
    assert p["gaps"]["control_history"]["measured"] is False
    assert "single row" in p["gaps"]["control_history"]["reason"]


# ── reference IDs: stable, deterministic, traceable ─────────────────────────
def test_reference_ids_are_deterministic_stable_and_traceable(client):
    """§24: stable, non-mutating, resolvable. Built from the REAL source table
    plus its REAL primary key — ⛔ no counter, no clock, no random component."""
    a = _a(client, WIDE)["records"]
    b = _a(client, WIDE)["records"]
    assert [r["ref_id"] for r in a] == [r["ref_id"] for r in b], "ids moved on refresh"
    assert len({r["ref_id"] for r in a}) == len(a), "duplicate reference ids"
    for r in a:
        assert r["ref_id"] and r["src_table"] and r["src_pk"] is not None
    # the prefix names the real source, so an id resolves back to a row
    assert any(r["ref_id"].startswith("SYSEVT-") for r in a)


def test_a_reference_id_resolves_to_its_record_with_a_timeline(client):
    p = _a(client, WIDE)
    ref = p["records"][0]["ref_id"]
    got = client.get("/api/audit-detail" + WIDE + "&ref_id=" + ref).get_json()
    assert got["ref_id"] == ref
    assert [s["stage"] for s in got["timeline"]] == list(audit.TIMELINE_STAGES)
    miss = client.get("/api/audit-detail" + WIDE + "&ref_id=NOPE-1")
    assert miss.status_code == 404


def test_ordering_is_deterministic_for_equal_timestamps():
    """⛔ Two records sharing a stamp must not swap between refreshes — that is
    how a polling table appears to duplicate or lose rows."""
    src = inspect.getsource(audit.build_audit)
    assert 'key=lambda r: (str(r.get("ts") or ""), str(r.get("ref_id") or ""))' in src


# ── filters, search, pagination ─────────────────────────────────────────────
def test_filters_actually_filter_and_options_come_from_the_population(client):
    p = _a(client, WIDE)
    opts = p["filters"]["options"]
    assert opts["category"] and set(opts["category"]) <= set(audit.CATEGORIES)
    cat = opts["category"][0]
    f = _a(client, WIDE + "&category=" + cat)
    assert f["records"] and all(r["category"] == cat for r in f["records"])
    assert len(f["records"]) <= len(p["records"])
    # ⛔ an option must never be offered that matches nothing
    for c in opts["category"]:
        assert any(r["category"] == c for r in p["records"]), c


def test_global_search_hits_the_real_fields_not_a_placeholder(client):
    """The approved search must reach Risk / Capital / a strategy or service
    name — i.e. the record's own text, ⛔ not a decorative input."""
    p = _a(client, WIDE)
    hit = _a(client, WIDE + "&q=max_daily_trades")
    assert hit["records"], "search found nothing for a field that exists"
    assert all("max_daily_trades" in json.dumps(r).lower() for r in hit["records"])
    assert len(hit["records"]) < len(p["records"])
    assert _a(client, WIDE + "&q=zzz-no-such-thing")["records"] == []


def test_status_filter_uses_the_structured_value_not_free_text(client):
    p = _a(client, WIDE + "&status=Success")
    assert all(r["status"] == "Success" for r in p["records"])


def test_reset_restores_the_unfiltered_population(client):
    """Reset in the UI clears every filter and re-requests the default range;
    the server contract that makes that work is that no filter is sticky."""
    wide = _a(client, WIDE)
    filtered = _a(client, WIDE + "&category=System")
    again = _a(client, WIDE)
    assert len(again["records"]) == len(wide["records"]) >= len(filtered["records"])


# ── charts reconcile ────────────────────────────────────────────────────────
def test_the_donut_reconciles_to_the_filtered_population(client):
    p = _a(client, WIDE)
    d = p["distribution"]
    assert d["reconciles"] is True
    assert d["total"] == len(p["records"])
    assert sum(s["count"] for s in d["slices"]) == len(p["records"])
    if d["total"]:
        assert abs(sum(s["pct"] for s in d["slices"]) - 100.0) < 0.6
    for s in d["slices"]:
        assert s["category"] in audit.CATEGORIES or s["category"] == "Unknown"


def test_events_over_time_conserves_counts_and_fills_every_day(client):
    """⛔ Plotting only the days that had activity would join two non-adjacent
    days as if they were consecutive — a quiet day must be a REAL zero."""
    p = _a(client, "?start=2026-08-01&end=2026-08-07")
    ot = p["over_time"]
    assert [x["day"] for x in ot["points"]] == [
        "2026-08-0%d" % d for d in range(1, 8)]
    assert ot["conserved"] is True
    assert ot["total"] + ot["outside_range"] == len(p["records"])


def test_top_users_reconciles_and_never_invents_a_person(client):
    p = _a(client, WIDE)
    tu = p["top_users"]
    assert tu["attributed"] == sum(u["changes"] for u in tu["users"])
    assert tu["attributed"] + tu["unattributed"] == len(p["records"])
    assert all(u["kind"] == "process" for u in tu["users"])
    assert "PROCESSES" in tu["note"] or "processes" in tu["note"]
    if tu["attributed"]:
        assert abs(sum(u["pct"] for u in tu["users"]) - 100.0) < 0.6
    # ⛔ unattributed records are excluded from the ranking, not given to anyone
    assert "excluded" in tu["unattributed_note"]


def test_critical_changes_are_derived_from_real_records(client):
    """⛔ Not the approved example list hard-coded: a change is critical when it
    touched Risk / Capital / Control, or when it FAILED."""
    p = _a(client, WIDE)
    refs = {r["ref_id"] for r in p["records"]}
    for c in p["critical"]:
        assert c["ref_id"] in refs, "a critical row that is not a real record"
        assert c["category"] in ("Risk", "Capital", "Control") or c["status"] == "Failed"


# ── export ──────────────────────────────────────────────────────────────────
def test_export_carries_the_filtered_rows_only(client):
    from openpyxl import load_workbook
    p = _a(client, WIDE + "&category=System")
    r = client.get("/api/export/audit-screen" + WIDE + "&category=System")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Audit Events", "Category Distribution",
                             "Events Over Time", "Top Actors", "Critical Changes",
                             "Retention", "Summary"]
    ws = wb["Audit Events"]
    assert [c.value for c in ws[1]] == [h for h, _k in audit.EXPORT_COLS]
    assert ws.max_row - 1 == len(p["records"])
    # ⛔ Scanner is not an export column
    assert not any("scanner" in str(c.value).lower() for c in ws[1])


def test_export_spells_out_gaps_rather_than_leaving_blanks(client):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(
        client.get("/api/export/audit-screen" + WIDE).data))
    ws = wb["Audit Events"]
    hdr = [c.value for c in ws[1]]
    ucol = hdr.index("User")
    vals = {ws.cell(row=i, column=ucol + 1).value for i in range(2, ws.max_row + 1)}
    assert "NOT INSTRUMENTED" in vals, "an unrecorded actor must say so"
    assert None not in vals, "a blank cell reads as zero or as 'nobody'"
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["Retention Period"] == "NOT INSTRUMENTED"


# ── the approved layout is present ──────────────────────────────────────────
def test_every_approved_panel_is_on_the_page(client):
    body = _shipped(client)
    for heading in ("FILTERS", "AUDIT EVENTS", "AUDIT DETAILS", "TIMELINE",
                    "CRITICAL CHANGES", "AUDIT CATEGORY DISTRIBUTION",
                    "AUDIT EVENTS OVER TIME", "TOP USERS", "RETENTION OVERVIEW"):
        assert heading in body, heading
    for control in ("Reset Filters", "Apply Filters", "Global Search",
                    "Export XLSX", "Export Current View", "Rows per page"):
        assert control in body, control
    for label in ("OLD VALUE", "NEW VALUE", "Audit Records", "Oldest Record",
                  "Newest Record", "Retention Period"):
        assert label in body, label


def test_the_donut_and_the_line_chart_are_both_present_and_not_substituted():
    """⛔ The approved design names a pie/donut AND a line graph. Neither may be
    replaced by another chart type."""
    body = _tpl()
    assert "aud-donut" in body and 'r="15.9155"' in body     # donut geometry
    assert "donutMarkup()" in body
    assert "aud-line" in body and "lineMarkup()" in body
    assert "polyline" in body


def test_no_template_x_for_inside_svg_anywhere_in_the_tree():
    """⛔ THE PINNED RENDER TRAP: `<template x-for>` inside `<svg>` is parsed as
    an SVG element named "template", so Alpine repeats NOTHING and the panel is
    a silent empty box. Every repeated SVG child must be injected via x-html."""
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    tdir = os.path.join(here, "frontend", "templates")
    offenders = []
    for name in sorted(os.listdir(tdir)):
        if not name.endswith(".html"):
            continue
        html = _read("frontend", "templates", name)
        # ⛔ STRIP COMMENTS FIRST. A Jinja `{# #}` never reaches the browser, and
        # the audit template's own header comment EXPLAINS this trap by quoting
        # `<template x-for>` and `<svg>` verbatim — scanning raw source would
        # flag the very documentation that prevents the bug.
        html = re.sub(r"\{#.*?#\}", " ", html, flags=re.S)
        html = re.sub(r"<!--.*?-->", " ", html, flags=re.S)
        for svg in re.findall(r"<svg\b.*?</svg>", html, re.S):
            if re.search(r"<template[^>]*x-for", svg):
                offenders.append(name)
    assert not offenders, offenders


def test_column_drag_reuses_the_established_convention():
    body = _tpl()
    for token in ('draggable="true"', "onDragStart", "onDrop", "onDragEnd",
                  "is-drag", "is-over", "localStorage", "screen13.audit.colOrder"):
        assert token in body, token
    # the 250ms guard so the click that ENDS a drag is not read as a sort
    assert "_dragEndAt" in body and "250" in body
    # a merge that can neither lose nor duplicate a column
    assert "next.length === this.DEFAULT_COLS.length" in body


def test_every_table_mixin_member_the_page_calls_actually_exists():
    """⚠️ THIS CLASS OF BUG BIT TWICE ON THIS SCREEN, and both times silently:
    `tSize` (really `tPageSize`) rendered "Showing NaN to NaN", and `tPages`
    (really `tPageCount`) made the NUMBERED page buttons vanish — the approved
    design shows `1 2 3 4 5 … 855`, so a pager with only ‹ › is a missing
    element, not a cosmetic nit.

    Alpine swallows an undefined-method binding, so nothing throws visibly and
    every contract test still passes. Checking the names against the mixin's
    real source is the only cheap way to catch it."""
    tpl = _tpl()
    mixin = _read("frontend", "static", "components.js")
    defined = set(re.findall(r"^\s{4}(t[A-Za-z]+)\s*:", mixin, re.M))
    assert {"tPageCount", "tPaged", "tPageSize", "tPage", "tSort"} <= defined, defined
    used = set(re.findall(r"\b(t[A-Z][A-Za-z]*)\b", tpl))
    unknown = sorted(u for u in used if u not in defined)
    assert not unknown, "template calls table-mixin members that do not exist: %s" % unknown


def test_the_pager_renders_numbered_pages_not_just_arrows():
    """The approved design's pager is `‹ 1 2 3 4 5 … 855 ›`."""
    tpl = _tpl()
    assert "pageList()" in tpl
    body = re.search(r"pageList\(\)\s*\{(.*?)\n    \},", tpl, re.S).group(1)
    assert "tPageCount" in body and "…" in body


def test_header_and_body_share_one_column_list_so_they_cannot_desync():
    """⛔ The classic drag defect: headers reorder, cells do not. Both loops must
    iterate the SAME `cols` array."""
    body = _tpl()
    assert body.count('x-for="c in cols"') >= 2


def test_status_text_carries_the_semantic_colour(client):
    """The project-wide rule confirmed on Screen 12: the TEXT is coloured, not
    merely an icon."""
    css = _css()
    assert ".aud-page .aud-st-Success { background: var(--pos-tint);  color: var(--pos); }" in css
    assert ".aud-page .aud-st-Failed  { background: var(--neg-tint);  color: var(--neg); }" in css
    assert ".aud-page .aud-st-Partial { background: var(--warn-tint); color: var(--yellow); }" in css


def test_all_eight_categories_have_a_distinct_badge_and_arc():
    css = _css()
    for c in audit.CATEGORIES:
        assert ".aud-cat-%s" % c in css, c
        assert ".aud-dot-%s" % c in css, c
        assert ".aud-arc-%s" % c in css, c


def test_css_is_scoped_to_this_screen_only():
    """⛔ A bare selector here would repaint every other approved screen."""
    css = _css()
    start = css.index("SCREEN 13 — AUDIT")
    nxt = re.search(r"SCREEN \d+ [—-]", css[start + 20:])
    block = css[start: start + 20 + nxt.start()] if nxt else css[start:]
    for line in block.splitlines():
        line = line.strip()
        if line.startswith(".") and "{" in line:
            assert ".aud-page" in line, line


def test_table_text_meets_the_thirteen_pixel_floor():
    """The brief sets a 13px minimum for readable text."""
    css = _css()
    start = css.index("SCREEN 13 — AUDIT")
    nxt = re.search(r"SCREEN \d+ [—-]", css[start + 20:])
    block = css[start: start + 20 + nxt.start()] if nxt else css[start:]
    sizes = [float(m) for m in re.findall(r"font-size:\s*(\d+(?:\.\d+)?)px", block)]
    small = [s for s in sizes if s < 13]
    # the only sub-13px value is the drag grip, expressed in rem as on Screens
    # 11/12 — a decorative handle glyph, not text.
    assert not small, small
    assert "--dt-cell" in block or "13px" in block
    # ⭐ SHARED component classes come in BELOW the floor (measured: th 10.88px,
    # .mono 12.16px, .rt-btn/.panel-sub 11.84px). They are lifted here and ONLY
    # here — a global change would repaint every other approved screen.
    for rule in (".aud-page .rt-btn { font-size: 13px; }",
                 ".aud-page .panel-sub { font-size: 13px; }",
                 ".aud-page th { font-size: 13px; }",
                 ".aud-page .mono { font-size: 13px; }"):
        assert rule in block, rule


def test_the_line_chart_is_not_scaled_by_a_viewbox():
    """⛔ THE SCREEN-10 TRAP: a fixed-width viewBox in a narrower panel scales
    the axis TEXT down with it. The width is measured instead, so one user unit
    is one CSS pixel. ⛔ And `:viewBox` would bind the lowercase `viewbox`, which
    SVG ignores — it would look right in source and do nothing."""
    body = _tpl()
    svg = re.search(r'<svg class="aud-line".*?>', body, re.S).group(0)
    assert "viewBox" not in svg and "viewbox" not in svg
    assert 'preserveAspectRatio="none"' not in svg
    assert "measure()" in body and "clientWidth" in body


def test_real_time_uses_the_existing_refresh_path(client):
    """⛔ No second monitoring architecture: the page listens to the dashboard's
    own `ops-refresh` broadcast."""
    body = _shipped(client)
    assert "ops-refresh.window" in body
    assert "pageBase(\"/api/audit-screen\")" in body


def test_repeated_loads_neither_duplicate_nor_lose_rows(client):
    """⛔ A polling table must not gain or drop rows just because it polled."""
    first = _a(client, WIDE)["records"]
    for _ in range(3):
        again = _a(client, WIDE)["records"]
        assert len(again) == len(first)
        assert [r["ref_id"] for r in again] == [r["ref_id"] for r in first]


def test_the_ist_convention_is_stated_on_the_page(client):
    body = _shipped(client)
    assert "IST" in body and "Asia/Kolkata" in body


def test_the_page_declares_its_instrumentation_gaps_in_words(client):
    """An accountability screen must say what it cannot account for."""
    p = _a(client, WIDE)
    for key in ("user", "authentication", "source", "timeline",
                "retention_period", "control_history"):
        assert p["gaps"][key]["measured"] is False, key
        assert len(p["gaps"][key]["reason"]) > 30, key
    assert "OBSERVED" in p["note"] or "observed" in p["note"]


# ═════════════════════════════════════════════════════════════════════════════
# THE 30-Aug PASS AGAINST THE ORIGINAL ARTWORK
# ⛔ These pin the three corrections. Each names a string absent before the pass.
# ═════════════════════════════════════════════════════════════════════════════
def test_critical_changes_sits_with_the_lower_analytics_not_in_the_table_rail():
    """⭐ The design TXT is explicit: *"the original PNG places this as a
    right-side panel aligned with the lower analytics area. Preserve that
    relationship."*

    ⚠️ It is also a DEAD-SPACE fix. 🔬 With a record selected the rail stood
    1474px against a 528px table — 946px of void beside it, visible only AFTER
    a row is clicked, which is why an unselected screenshot never showed it.
    """
    t = _tpl()
    rail = t[t.index('<aside class="aud-rail">'):t.index("</aside>")]
    assert "aud-critical" not in rail, "CRITICAL CHANGES is back in the table rail"
    lower = t[t.index('<div class="aud-lower">'):t.index('<div class="aud-footer">')]
    assert "aud-critical" in lower and "aud-zone" in lower
    # ⭐ the lower band must MIRROR .aud-main, or the panel stops lining up
    # underneath AUDIT DETAILS and the whole point of the move is lost.
    css = _css()
    main = css[css.index(".aud-page .aud-main {"):]
    main = main[:main.index("}")]
    low = css[css.index(".aud-page .aud-lower {"):]
    low = low[:low.index("}")]
    assert "minmax(0, 1fr) 330px" in main and "minmax(0, 1fr) 330px" in low


def test_the_bottom_export_bar_keeps_the_note_left_and_the_buttons_right():
    """⚠️ The rule was ALREADY `space-between`. `flex-wrap: wrap` plus the long
    honesty note wrapped the buttons onto their own line at the LEFT, so the
    bar read nothing like the artwork. ⛔ `nowrap` is the load-bearing part."""
    t, css = _tpl(), _css()
    foot = t[t.index('<div class="aud-footer">'):]
    assert foot.count('class="btn-export"') == 2, "not the shared S09/S11/S12 control"
    assert "Export Current View" in foot and "Export XLSX" in foot
    assert "/api/export/audit-screen" in t or "exportUrl()" in foot
    rule = css[css.index(".aud-page .aud-footer { flex-wrap: nowrap"):]
    assert "nowrap" in rule[:80]


def test_every_filter_select_is_one_width():
    """⛔ Action rendered ~2x its siblings purely because its option strings are
    longer — a control's width must not encode its data."""
    css = _css()
    assert ".aud-page .aud-frow .aud-f > .sel { width: 168px; }" in css


def test_the_pass_touched_no_other_screen():
    """🔬 Every rule the S13 pass added is `.aud-page`-scoped.

    ⚠️ THE WINDOW MUST END AT THE NEXT SCREEN'S BANNER. The first version read
    to END-OF-FILE, so the very next screen to append its own correctly-scoped
    block failed this test — 🔬 S14's `.tlg-page .tlg-row4` did exactly that.
    ⛔ That was the test's window being wrong, ⛔ not the other screen; the
    assertion itself is right and stays.
    """
    css = _css()
    block = css[css.index("SCREEN 13 — the pass against the original artwork"):]
    nxt = re.search(r"SCREEN \d+ . the pass against the original artwork", block[1:])
    if nxt:
        block = block[:nxt.start() + 1]
    for line in block.splitlines():
        line = line.strip()
        if "{" in line and not line.startswith(("/*", "*", "@")):
            assert line.startswith(".aud-page"), "unscoped selector: %s" % line
