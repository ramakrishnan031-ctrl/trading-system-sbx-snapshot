"""SCREEN 22 — HOLDINGS.  The broker-first reconciliation centre.

The approved design (`gui/22. Holdings.png` + `.txt`) is binding for STRUCTURE;
the DATA is this system's.

⭐ THE CENTRAL PROPERTY: the BROKER side is real and comes from
`position_reconciliation` — the per-symbol comparison the 15:45
`reconcile_positions` cron writes — and the SYSTEM side is the currently-open
position book. Delta is broker minus system at the SOURCE'S OWN GRAIN.

⛔⛔ AND THE THREE THINGS THIS SCREEN MUST NEVER DO, each pinned below:
  1. count broker quantities at ROW grain. The shared fixture puts FOUR open
     trades on symbol AAA, so a row-grain sum reports 195 where the truth is 75.
  2. render a market value. There is NO live price source in this dashboard, so
     Current Price / Market Value / Unrealized P&L are gaps — ⛔ and the entry
     price is never substituted, which would make every unrealised P&L exactly
     zero and read as a measurement.
  3. read "no reconciliation record" as "Matched". Absence of a comparison is
     not agreement.

⚠️ The shared fixture seeds ONE ROW OF EVERY STATUS the writer can produce, and
TWO RUNS whose verdicts DISAGREE for the same symbol — so a reader that keyed on
MAX(date) instead of the run stamp returns two verdicts for AAA and goes red.
"""
from __future__ import annotations

import io
import os
import re
import sqlite3

import _js_syntax

from backend.readers import db_reader
from backend.services import holdings, scanner_attribution

from conftest import TODAY, YDAY, _ts


def _read(*parts) -> str:
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    with open(os.path.join(here, *parts), encoding="utf-8") as fh:
        return fh.read()


def _tpl() -> str:
    return _read("frontend", "templates", "holdings.html")


def _css() -> str:
    return _read("frontend", "static", "style.css")


def _page(client) -> str:
    html = client.get("/holdings").data.decode("utf-8")
    return re.sub(r"<!--.*?-->", " ", html, flags=re.S)


def _s(client, qs=""):
    return client.get("/api/holdings/screen" + qs).get_json()


def _css_block() -> str:
    """Screen 22's CSS block ONLY — bounded at the next screen's header.

    🔴 CORRECTED 17-Aug-2026, ⛔ property unchanged. This window used to run to
    END OF FILE, which was invisible while Screen 22 was the last block in
    style.css. When Screen 17's block was appended after it, this window silently
    adopted Screen 17's rules and reported them as Screen 22's own unscoped
    selectors.

    📌 This is the SAME defect the Screen-12, Screen-18 and Screen-20 windows were
    corrected for on 16-Aug — *"a CSS window that ends at EOF silently adopts
    whatever is appended after it."* Screen 22's was never bounded only because
    nothing had followed it yet. ⛔ Screen 22's CSS and rendering are untouched.
    """
    css = _css()
    start = css.index("SCREEN 22 — HOLDINGS")
    nxt = css.find("\n   SCREEN ", start + 10)
    return css[start:] if nxt == -1 else css[start:nxt]


def _conn(gui_config):
    return sqlite3.connect(gui_config["paths"]["main_db"])


def _seed_delivery(gui_config, *, symbol="LT", strategy="range_breakout_long",
                   qty=10, price=3150.0, days_ago=4, product="CNC",
                   broker_qty=10, status="OK"):
    """A CNC position carried from an earlier day, plus its reconciliation row.

    ⭐ SEEDED HERE RATHER THAN IN THE SHARED FIXTURE, the precedent Screen 18
    set with `_seed_same_day_chain`: adding an open CNC trade to conftest would
    move `open_positions_count`, the capacity rollup and every per-strategy
    order count for screens that have nothing to do with delivery.
    ⚠️ WITHOUT IT the delivery panel is all zeros and every delivery assertion
    is vacuous — a fixture with no CNC row cannot tell a working product filter
    from one that returns nothing.
    """
    import datetime as _dt
    day = (_dt.date(*(int(x) for x in TODAY.split("-")))
           - _dt.timedelta(days=days_ago)).isoformat()
    ts = day + "T10:05:00+05:30"
    tid = "trd_dlv_" + symbol
    c = _conn(gui_config)
    c.execute(
        "INSERT INTO trades(trade_id,signal_id,symbol,direction,strategy,sector,"
        "qty_planned,qty_filled,entry_target_price,entry_actual_price,sl_initial,"
        "tgt_initial,margin_reserved,risk_amount,created_at,entry_time,status,"
        "actual_position_value_rs) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (tid, "sig_" + tid, symbol, "LONG", strategy, "IT", qty, qty, price,
         price, price * 0.97, price * 1.05, price * qty, 100.0, ts, ts, "OPEN",
         price * qty))
    c.execute(
        "INSERT INTO orders(order_id,trade_id,leg,transaction_type,order_type,"
        "product,variety,qty_requested,qty_filled,status,placed_at,filled_at,"
        "updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
        ("ord_" + tid, tid, "ENTRY", "BUY", "LIMIT", product, "regular", qty,
         qty, "COMPLETE", ts, ts, ts))
    c.execute(
        "INSERT INTO position_reconciliation(date,symbol,broker_qty,system_qty,"
        "status,resolved_at,created_at) VALUES(?,?,?,?,?,?,?)",
        (TODAY, symbol, broker_qty, qty, status, None, _ts("15:45:07")))
    c.commit()
    c.close()
    return {"trade_id": tid, "symbol": symbol, "qty": qty, "price": price,
            "days": days_ago, "product": product}


# ═════════════════════════════════════════════════════════════════════════════
# THE BROKER SIDE IS REAL — and it is read at the source's own grain
# ═════════════════════════════════════════════════════════════════════════════
def test_the_broker_side_comes_from_position_reconciliation(gui_config):
    """⭐ The per-symbol table `scripts/reconcile_positions.py` writes at 15:45.
    ⚠️ This supersedes the old screen's "Pending Broker Source (G4/P1)" claim:
    P1 is a DATE-LEVEL verdict table with no symbols and was never this source.
    """
    p = holdings.build_holdings_screen(gui_config)
    rc = p["reconciliation"]
    assert rc["available"] is True
    assert "position_reconciliation" in rc["source"]
    assert rc["grain"] == "symbol"
    assert rc["last_run"]
    assert any(r["broker_qty"] is not None for r in p["rows"])


def test_only_the_latest_run_is_read_not_the_latest_date(gui_config):
    """⛔⛔ `reconcile_positions` INSERTS — it never upserts — so a re-run
    appends a SECOND full set of rows for the same date. Keying on MAX(date)
    returns two verdicts per symbol.

    ⭐ The fixture seeds an OLDER run that calls AAA QTY_MISMATCH and a CURRENT
    one that calls it OK, so this check has a real failing input.
    """
    rows = db_reader.position_reconciliation_latest(gui_config)
    stamps = {r["created_at"] for r in rows}
    assert len(stamps) == 1, stamps
    symbols = [r["symbol"] for r in rows]
    assert len(symbols) == len(set(symbols)), symbols
    aaa = next(r for r in rows if r["symbol"] == "AAA")
    assert aaa["status"] == "OK", "the older run's verdict leaked through"


def test_the_fixture_really_holds_two_disagreeing_runs(gui_config):
    """⭐ Without this the latest-run check above could pass on a single run."""
    c = _conn(gui_config)
    rows = c.execute("SELECT symbol, status, created_at FROM position_reconciliation "
                     "WHERE symbol='AAA' ORDER BY created_at").fetchall()
    c.close()
    assert len(rows) == 2, rows
    assert rows[0][1] != rows[1][1], rows


def test_broker_quantities_are_symbol_grain_not_row_grain(gui_config):
    """⛔⛔ GETTING THE GRAIN WRONG MULTIPLIES MONEY. The comparison is written
    per SYMBOL while the table is per (symbol, product); the fixture puts FOUR
    open trades on AAA, so a row-grain sum reports 195 where the truth is 75."""
    p = holdings.build_holdings_screen(gui_config)
    recon = db_reader.position_reconciliation_latest(gui_config)
    inst = [r for r in recon if r["symbol"] != "ALL"]
    assert p["delta"]["broker_qty"] == sum(int(r["broker_qty"]) for r in inst)
    assert p["delta"]["system_qty"] == sum(int(r["system_qty"]) for r in inst)
    assert p["delta"]["delta_qty"] == (p["delta"]["broker_qty"]
                                       - p["delta"]["system_qty"])
    # the row-grain figure would be strictly larger — prove the fixture can tell
    row_grain = sum(int(r["broker_qty"]) for r in p["rows"]
                    if r["broker_qty"] is not None)
    assert row_grain > p["delta"]["broker_qty"], (
        "the fixture has no repeated symbol, so the grain distinction is untested")


def test_the_kpi_broker_count_is_symbol_grain(gui_config):
    p = holdings.build_holdings_screen(gui_config)
    recon = db_reader.position_reconciliation_latest(gui_config)
    want = sum(1 for r in recon if r["symbol"] != "ALL" and r["broker_qty"])
    assert p["kpi"]["broker_holdings"] == want
    assert p["broker_snapshot"]["holdings_count"] == want
    assert p["side_by_side"]["broker"]["positions"] == want


def test_the_two_kpi_sides_publish_their_different_bases(gui_config):
    """⭐ "28 broker vs 27 system" is meaningful only if a reader can tell what
    each side counted."""
    p = holdings.build_holdings_screen(gui_config)
    assert p["kpi"]["broker_base"] and p["kpi"]["system_base"]
    assert p["kpi"]["broker_base"] != p["kpi"]["system_base"]


def test_a_run_level_broker_failure_is_not_drawn_as_an_instrument(gui_config):
    """⛔ `reconcile_positions` writes ONE row with symbol='ALL' when the broker
    fetch itself failed. That is a run-level error, ⛔ not a holding."""
    c = _conn(gui_config)
    c.execute("DELETE FROM position_reconciliation")
    c.execute("INSERT INTO position_reconciliation(date,symbol,broker_qty,"
              "system_qty,status,created_at) VALUES(?,?,?,?,?,?)",
              (TODAY, "ALL", None, 4, "ERROR", _ts("15:45:07")))
    c.commit()
    c.close()
    p = holdings.build_holdings_screen(gui_config)
    assert not any(r["symbol"] == "ALL" for r in p["rows"])
    assert p["reconciliation"]["broker_reachable"] is False
    assert p["kpi"]["broker_holdings"] == 0


# ═════════════════════════════════════════════════════════════════════════════
# THE THREE-WAY STATUS — from the structured source, never a numeric band
# ═════════════════════════════════════════════════════════════════════════════
def test_the_three_approved_statuses_come_from_the_structured_source(gui_config):
    p = holdings.build_holdings_screen(gui_config)
    assert list(p["statuses"]) == ["Matched", "Minor Difference", "Mismatch"]
    want = {"OK": "Matched", "QTY_MISMATCH": "Minor Difference",
            "ORPHAN_AT_BROKER": "Mismatch", "MISSING_AT_BROKER": "Mismatch"}
    seen = set()
    for r in p["rows"]:
        src = r["recon_source_status"]
        if src in want:
            assert r["recon_status"] == want[src], r
            seen.add(src)
    assert seen == set(want), "the fixture does not exercise every status: %s" % seen


def test_no_numeric_threshold_splits_minor_from_mismatch():
    """⛔ Splitting on a quantity band would be an invented classifier, and the
    artwork's own numbers contradict any single one (Δ5-of-40 minor,
    Δ10-of-100 a mismatch). The distinction is structural."""
    src = _read("backend", "services", "holdings.py")
    body = src[src.index("_STATUS_MAP = {"):src.index("HEALTH_TILES")]
    assert not re.search(r"[<>]=?\s*\d", body), body
    assert "abs(" not in body


def test_a_symbol_with_no_reconciliation_record_is_not_matched(gui_config):
    """⛔ ABSENCE OF A COMPARISON IS NOT AGREEMENT. A screen that renders it as
    agreement is worse than one that renders nothing."""
    c = _conn(gui_config)
    c.execute("DELETE FROM position_reconciliation")
    c.commit()
    c.close()
    p = holdings.build_holdings_screen(gui_config)
    assert p["rows"], "no system rows — the check is vacuous"
    for r in p["rows"]:
        assert r["recon_status"] is None, r
        assert r["broker_qty"] is None and r["delta_qty"] is None
        assert r["recon_note"], r
    assert p["kpi"]["matched"] == 0
    assert p["kpi"]["matched_pct"] is None
    assert p["kpi"]["not_reconciled"] == len(p["rows"])
    assert p["reconciliation"]["available"] is False


def test_a_symbol_level_verdict_shown_on_many_rows_says_so(gui_config):
    """⭐ The fixture puts four open trades on AAA, so its one verdict appears
    four times. Every such row is flagged rather than passing a symbol-level
    call off as a product-level one."""
    p = holdings.build_holdings_screen(gui_config)
    aaa = [r for r in p["rows"] if r["symbol"] == "AAA"]
    assert len(aaa) > 1, "the fixture has no repeated symbol"
    assert all(r["recon_shared"] for r in aaa)
    singles = [r for r in p["rows"] if r["symbol"] != "AAA"]
    assert singles and not any(r["recon_shared"] for r in singles)


def test_delta_is_broker_minus_system_on_every_row(gui_config):
    p = holdings.build_holdings_screen(gui_config)
    checked = 0
    for r in p["rows"]:
        if r["broker_qty"] is None or r["system_qty"] is None:
            assert r["delta_qty"] is None, r
            continue
        assert r["delta_qty"] == int(r["broker_qty"]) - int(r["system_qty"]), r
        checked += 1
    assert checked, "no reconciled row — the check is vacuous"


# ═════════════════════════════════════════════════════════════════════════════
# THE ORPHAN DETECTOR — two-sided
# ═════════════════════════════════════════════════════════════════════════════
def test_the_orphan_detector_is_two_sided(gui_config):
    """⛔ A one-sided detector finds only the orphans it was pointed at; the
    position missing from the side you trusted is the one that costs money."""
    p = holdings.build_holdings_screen(gui_config)
    o = p["orphans"]
    assert o["broker_without_system"]["n"] >= 1
    assert o["system_without_broker"]["n"] >= 1
    recon = db_reader.position_reconciliation_latest(gui_config)
    assert o["broker_without_system"]["n"] == sum(
        1 for r in recon if r["status"] == "ORPHAN_AT_BROKER")
    assert o["system_without_broker"]["n"] == sum(
        1 for r in recon if r["status"] == "MISSING_AT_BROKER")


def test_a_broker_only_symbol_still_appears_as_a_row(gui_config):
    """⛔ A symbol the broker holds that the system never opened produces no
    system row — without the second pass it would be INVISIBLE, which is the
    one row an operator most needs to see."""
    p = holdings.build_holdings_screen(gui_config)
    zzz = next((r for r in p["rows"] if r["symbol"] == "ZZZ"), None)
    assert zzz is not None, [r["symbol"] for r in p["rows"]]
    assert zzz["source"] == "Broker Only"
    assert zzz["origin"] == "broker"
    assert zzz["strategy"] is None and zzz["quantity"] is None


def test_a_recon_only_row_never_inflates_the_system_count(gui_config):
    """⭐ CCC was MISSING_AT_BROKER at 15:45 and has since left the open set. It
    is shown, but it is not a currently-open system position."""
    p = holdings.build_holdings_screen(gui_config)
    ccc = next((r for r in p["rows"] if r["symbol"] == "CCC"), None)
    assert ccc is not None
    assert ccc["origin"] == "recon"
    assert ccc["source"] == "System Only"
    system = [r for r in p["rows"] if r["origin"] in ("system", "both")]
    assert ccc not in system
    assert p["kpi"]["system_holdings"] == len(system)


def test_the_orphan_note_carries_the_t1_blindness(gui_config):
    p = holdings.build_holdings_screen(gui_config)
    assert "positions()" in p["orphans"]["note"]


# ═════════════════════════════════════════════════════════════════════════════
# NO LIVE PRICE — and the entry price is never substituted
# ═════════════════════════════════════════════════════════════════════════════
def test_the_three_price_fields_are_unavailable_on_every_row(gui_config):
    """⛔ `ops_dashboard` has ZERO live-price call sites."""
    p = holdings.build_holdings_screen(gui_config)
    assert p["rows"]
    for r in p["rows"]:
        for k in holdings.PRICE_GAP:
            assert r[k] is None, (r["symbol"], k, r[k])


def test_the_entry_price_is_never_substituted_for_the_current_price(gui_config):
    """⛔ Substituting it would make every unrealised P&L exactly zero, which a
    reader takes as a measurement. The avg price IS shown, under its own name."""
    p = holdings.build_holdings_screen(gui_config)
    priced = [r for r in p["rows"] if r["avg_price"] is not None]
    assert priced, "no priced row — the check is vacuous"
    for r in priced:
        assert r["current_price"] is None
        assert r["unrealized_pnl"] is None


def test_the_two_market_kpis_are_unavailable_with_a_reason(gui_config):
    p = holdings.build_holdings_screen(gui_config)
    assert p["kpi"]["total_value"] is None
    assert p["kpi"]["total_pnl"] is None
    gap = p["gaps"]["price"]
    assert gap["measured"] is False and gap["reason"] and gap["short"]
    assert "quote" in gap["reason"] or "price" in gap["reason"]


def test_the_measured_cost_value_is_shown_under_its_own_name(gui_config):
    """⭐ Filled qty × filled entry price IS measured, and the artwork's own
    SIDE-BY-SIDE panel already distinguishes "Market Value" from "Expected
    Value". ⛔ It is never labelled market value."""
    p = holdings.build_holdings_screen(gui_config)
    system = [r for r in p["rows"] if r["origin"] in ("system", "both")]
    want = sum(r["value_at_cost"] for r in system if r["value_at_cost"] is not None)
    assert p["kpi"]["value_at_cost"] == round(want, 2)
    assert p["side_by_side"]["system"]["expected_value"] == round(want, 2)
    assert p["side_by_side"]["broker"]["market_value"] is None


def test_the_cost_value_uses_the_readers_own_definition(gui_config):
    """⛔ ONE definition of "value" — Screen 06 and Screen 22 must not differ."""
    rows = db_reader.holdings_system_rows(gui_config)
    assert rows
    p = holdings.build_holdings_screen(gui_config)
    by_trade = {r["trade_id"]: r for r in p["rows"] if r["trade_id"]}
    for s in rows:
        assert by_trade[s["trade_id"]]["value_at_cost"] == \
            db_reader._position_value_of(s)


def test_price_mismatch_reports_unavailable_not_zero(gui_config):
    """⛔ A 0 there would claim the comparison was made and found nothing. No
    broker price is recorded anywhere, so it cannot be made at all."""
    p = holdings.build_holdings_screen(gui_config)
    tile = next(t for t in p["health"]["tiles"] if t["label"] == "Price Mismatch")
    assert tile["measured"] is False
    assert tile["n"] is None and tile["pct"] is None
    assert p["gaps"]["price_mismatch"]["reason"]


# ═════════════════════════════════════════════════════════════════════════════
# THE T+1 BLINDNESS — load-bearing, and it must be stated
# ═════════════════════════════════════════════════════════════════════════════
def test_the_reconciler_really_reads_positions_only():
    """⭐ MEASURED against the production script, ⛔ not assumed: if it ever
    starts calling holdings(), this screen's caveat must be revisited."""
    here = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    path = os.path.join(here, "scripts", "reconcile_positions.py")
    if not os.path.isfile(path):
        return
    with io.open(path, encoding="utf-8") as fh:
        src = fh.read()
    assert "kite.positions()" in src
    assert "kite.holdings()" not in src, (
        "reconcile_positions now reads holdings() — the T+1 blindness caveat on "
        "Screen 22 must be revisited")


def test_the_screen_states_the_t1_blindness(gui_config, client):
    p = holdings.build_holdings_screen(gui_config)
    g = p["gaps"]["delivery_blind"]
    assert g["measured"] is False
    assert "kite.holdings()" in g["reason"]
    assert p["delivery"]["blind_to_t1"] is True
    assert "positions() only" in p["note"]
    assert "positions() only" in _page(client)


# ═════════════════════════════════════════════════════════════════════════════
# PRODUCT — from orders.product, via a LEFT JOIN
# ═════════════════════════════════════════════════════════════════════════════
def test_product_comes_from_the_entry_order_not_from_trades(gui_config):
    """⛔⛔ THERE IS NO `trades.product` COLUMN — it lives on `orders` and is
    reached by `LEFT JOIN … leg='ENTRY'`."""
    src = _read("backend", "readers", "db_reader.py")
    q = src[src.index("def holdings_system_rows"):]
    assert "LEFT JOIN orders o" in q[:2000]
    assert "o.leg = 'ENTRY'" in q[:2000]
    assert "t.product" not in q[:2000]


def test_a_trade_with_no_entry_order_keeps_its_row_and_shows_no_product(gui_config):
    """⛔ The join is LEFT on purpose: an INNER join would silently hide exactly
    the positions most worth seeing. NULL is UNKNOWN, ⛔ never a product."""
    c = _conn(gui_config)
    c.execute("DELETE FROM orders WHERE trade_id = 'trd_o1' AND leg = 'ENTRY'")
    c.commit()
    c.close()
    p = holdings.build_holdings_screen(gui_config)
    row = next(r for r in p["rows"] if r["trade_id"] == "trd_o1")
    assert row["product"] == "OTHER"
    assert row["product_raw"] is None
    assert row["is_delivery"] is False


def test_a_recon_only_row_carries_no_product_at_all(gui_config):
    """⛔ It has no ENTRY order, so it has no recorded product — and it must not
    be bucketed under "Other", which would inflate a product breakdown with
    rows that have no product."""
    p = holdings.build_holdings_screen(gui_config)
    recon_only = [r for r in p["rows"] if r["origin"] in ("broker", "recon")]
    assert recon_only
    assert all(r["product"] is None for r in recon_only)
    assert p["by_product"]["unknown_product"] == len(recon_only)
    assert p["by_product"]["total"] == sum(s["n"] for s in p["by_product"]["slices"])


def test_the_product_donut_and_the_categories_strip_share_one_computation(gui_config):
    """⭐ The artwork draws both; ⛔ they must not be able to disagree."""
    p = holdings.build_holdings_screen(gui_config)
    assert p["by_product"] == p["categories"]


def test_other_appears_only_when_a_row_actually_has_that_product(gui_config):
    p = holdings.build_holdings_screen(gui_config)
    keys = [s["key"] for s in p["by_product"]["slices"]]
    system = [r for r in p["rows"] if r["origin"] in ("system", "both")]
    if not any(r["product"] == "OTHER" for r in system):
        assert "OTHER" not in keys, keys


# ═════════════════════════════════════════════════════════════════════════════
# DELIVERY PORTFOLIO — seeded locally so it is not vacuous
# ═════════════════════════════════════════════════════════════════════════════
def test_a_delivery_position_is_categorised_valued_and_dated(gui_config):
    seeded = _seed_delivery(gui_config)
    p = holdings.build_holdings_screen(gui_config)
    row = next(r for r in p["rows"] if r["symbol"] == seeded["symbol"])
    assert row["product"] == "CNC"
    assert row["is_delivery"] is True
    assert row["holding_days"] == seeded["days"]
    assert row["value_at_cost"] == round(seeded["qty"] * seeded["price"], 2)

    d = p["delivery"]
    assert d["positions"] == 1
    assert d["value_at_cost"] == round(seeded["qty"] * seeded["price"], 2)
    assert d["avg_holding_days"] == float(seeded["days"])
    assert d["dated"] == 1
    # ⛔ and the market figures stay unavailable
    assert d["current_value"] is None and d["unrealized_pnl"] is None

    slices = {s["key"]: s["n"] for s in p["by_product"]["slices"]}
    assert slices.get("CNC") == 1


def test_nrml_counts_as_delivery_too(gui_config):
    _seed_delivery(gui_config, symbol="MARUTI", product="NRML", qty=5,
                   price=11200.0, days_ago=2, broker_qty=5)
    p = holdings.build_holdings_screen(gui_config)
    row = next(r for r in p["rows"] if r["symbol"] == "MARUTI")
    assert row["product"] == "NRML" and row["is_delivery"] is True
    assert p["delivery"]["positions"] == 1


def test_holding_days_is_none_not_zero_when_undated(gui_config):
    """⛔ 0 would read as "opened today", a different and misleading claim."""
    assert holdings._holding_days(None, TODAY) is None
    assert holdings._holding_days("", TODAY) is None
    assert holdings._holding_days(TODAY, TODAY) == 0
    assert holdings._holding_days(YDAY, TODAY) == 1


def test_the_delivery_panel_is_empty_without_a_delivery_position(gui_config):
    """⭐ The control: the shared fixture is all MIS, so an implementation that
    counted every position as delivery would fail here."""
    p = holdings.build_holdings_screen(gui_config)
    assert p["delivery"]["positions"] == 0
    assert p["delivery"]["value_at_cost"] is None
    assert p["delivery"]["avg_holding_days"] is None


# ═════════════════════════════════════════════════════════════════════════════
# BROKER SNAPSHOT · TIMELINE · SYNC
# ═════════════════════════════════════════════════════════════════════════════
def test_the_broker_account_is_read_from_the_session_row(gui_config):
    """⛔ Never hard-coded. The artwork prints "Zerodha (ZB1234)" because that is
    what its account is."""
    p = holdings.build_holdings_screen(gui_config)
    session = db_reader.get_session_info(gui_config)
    snap = p["broker_snapshot"]
    assert snap["account_id"] == session["account_id"]
    assert snap["broker"] == session["broker"].title()
    assert session["account_id"] in snap["account_label"]


def test_last_broker_sync_is_the_reconciliation_run_not_the_cash_sync(gui_config):
    """⛔ NOT `capital_snapshot.last_broker_sync`, which is the CASH sync and
    answers a different question."""
    p = holdings.build_holdings_screen(gui_config)
    tl = db_reader.position_reconciliation_timeline(gui_config)
    assert p["broker_snapshot"]["last_sync"] == tl["last_run"]
    src = _read("backend", "services", "holdings.py")
    body = src[src.index("def _broker_snapshot"):src.index("def _timeline")]
    assert "capital_snapshot" not in body.replace("`capital_snapshot.\n", "")


def test_the_sync_control_is_disabled_and_says_why(gui_config, client):
    """⛔ This dashboard is read-only and has no broker path, so a working Sync
    Now cannot exist. The footprint is kept; the capability is not faked."""
    p = holdings.build_holdings_screen(gui_config)
    assert p["broker_snapshot"]["sync_control"] is False
    assert p["gaps"]["broker_sync_control"]["reason"]
    page = _page(client)
    btn = page[page.index("hld-sync-btn"):]
    btn = btn[:btn.index("</button>")]
    assert "disabled" in btn, btn


def test_the_cadence_is_the_real_cron_not_the_artworks_60_seconds(gui_config, client):
    """⛔ The artwork says "Auto-sync every 60 sec". The real cadence is the
    15:45 weekday cron, and the panel prints that."""
    p = holdings.build_holdings_screen(gui_config)
    assert "15:45" in p["broker_snapshot"]["cadence"]
    assert "60 sec" not in _page(client)


def test_the_timeline_reads_each_stamp_from_its_own_column(gui_config):
    p = holdings.build_holdings_screen(gui_config)
    tl = p["timeline"]
    c = _conn(gui_config)
    last_run = c.execute("SELECT MAX(created_at) FROM position_reconciliation").fetchone()[0]
    last_bad = c.execute("SELECT MAX(created_at) FROM position_reconciliation "
                         "WHERE status <> 'OK'").fetchone()[0]
    last_fix = c.execute("SELECT MAX(resolved_at) FROM position_reconciliation").fetchone()[0]
    c.close()
    assert tl["last_reconciliation"] == last_run
    assert tl["last_mismatch"] == last_bad
    assert tl["last_correction"] == last_fix
    assert tl["correction_note"]


def test_last_correction_is_none_when_nothing_was_ever_resolved(gui_config):
    """⛔ `resolved_at` is null until a mismatch is manually resolved — the
    schema says so — so None here is a real "never", not a missing feature."""
    c = _conn(gui_config)
    c.execute("UPDATE position_reconciliation SET resolved_at = NULL")
    c.commit()
    c.close()
    p = holdings.build_holdings_screen(gui_config)
    assert p["timeline"]["last_correction"] is None


def test_the_fixture_exercises_a_real_correction(gui_config):
    """⭐ The control for the test above: without a resolved row, "None" would
    pass whether the reader worked or not."""
    p = holdings.build_holdings_screen(gui_config)
    assert p["timeline"]["last_correction"] is not None


# ═════════════════════════════════════════════════════════════════════════════
# HEALTH SUMMARY · FILTERS
# ═════════════════════════════════════════════════════════════════════════════
def test_the_six_approved_health_tiles_are_present_in_order(gui_config):
    p = holdings.build_holdings_screen(gui_config)
    assert [t["label"] for t in p["health"]["tiles"]] == list(holdings.HEALTH_TILES)


def test_orphan_position_is_the_two_sided_total(gui_config):
    p = holdings.build_holdings_screen(gui_config)
    t = {x["label"]: x["n"] for x in p["health"]["tiles"]}
    assert t["Orphan Position"] == t["Broker Only"] + t["System Only"]
    assert t["Broker Only"] >= 1 and t["System Only"] >= 1


def test_every_measured_tile_publishes_the_base_its_percentage_uses(gui_config):
    """⛔ EVERY PERCENTAGE CARRIES ITS BASE OR IT IS NOT A NUMBER. Two bases are
    in play here and mixing them silently is the failure mode."""
    p = holdings.build_holdings_screen(gui_config)
    for t in p["health"]["tiles"]:
        if t["measured"]:
            assert t["base"], t
        else:
            assert t["base"] is None, t
    bases = {t["base"] for t in p["health"]["tiles"] if t["measured"]}
    assert len(bases) == 2, bases


def test_the_filters_offer_only_values_present_in_the_data(gui_config):
    """⭐ A dropdown offering a choice that matches nothing reads as broken."""
    p = holdings.build_holdings_screen(gui_config)
    f = p["filters"]
    present_products = {r["product"] for r in p["rows"] if r["product"]}
    assert {o["key"] for o in f["product"]} == present_products
    assert set(f["status"]) <= set(p["statuses"])
    assert set(f["source"]) == {r["source"] for r in p["rows"] if r["source"]}
    assert None not in f["symbol"]


def test_each_filter_actually_narrows_the_set(client):
    full = _s(client)
    assert full["count"] == full["total"]
    for arg, value in (("product", "MIS"), ("status", "Mismatch"),
                       ("source", "Broker Only")):
        got = _s(client, "?%s=%s" % (arg, value))
        assert got["count"] < full["total"], (arg, value, got["count"])
        assert got["total"] == full["total"], "total must stay the unfiltered base"
        assert got["count"] > 0, (arg, value)


def test_the_symbol_search_matches_symbol_and_strategy(client):
    by_symbol = _s(client, "?symbol=AAA")
    assert by_symbol["count"] and all(r["symbol"] == "AAA" for r in by_symbol["rows"])
    by_strategy = _s(client, "?symbol=gap_fade")
    assert by_strategy["count"]
    assert all("gap_fade" in (r["strategy"] or "") for r in by_strategy["rows"])


def test_the_trade_type_filter_uses_the_strategy_yaml_intent(client):
    """⭐ The SAME strategy-level value Screens 19 and 20 use."""
    got = _s(client, "?trade_type=Intraday")
    assert all(r["trade_type"] == "Intraday" for r in got["rows"])
    assert got["active"]["trade_type"] == "Intraday"


# ═════════════════════════════════════════════════════════════════════════════
# THE RELATED REJECTION PANEL
# ═════════════════════════════════════════════════════════════════════════════
def test_the_rejection_panel_uses_screen_21s_own_classifier(gui_config):
    """⛔ Two screens must never disagree about how many signals were rejected
    today or which reason led."""
    h = holdings.build_holdings_screen(gui_config)
    s = scanner_attribution.build_scanner_attribution_screen(gui_config)
    assert h["rejection_insights"]["rejected"] == s["rejection_analysis"]["total"]
    assert h["rejection_insights"]["top"] == s["rejection_analysis"]["top"]
    assert h["rejection_insights"]["href"] == "/scanner-attribution"
    src = _read("backend", "services", "holdings.py")
    assert "scanner_attribution._top_reason" in src


# ═════════════════════════════════════════════════════════════════════════════
# EXPORT
# ═════════════════════════════════════════════════════════════════════════════
def test_the_export_header_is_the_approved_table_order():
    tpl = _tpl()
    block = tpl[tpl.index("DEFAULT_COLS:"):tpl.index("/* the three the artwork")]
    labels = re.findall(r'label:\s*"([^"]+)"', block)
    assert labels == list(holdings.EXPORT_HEADER), (labels, holdings.EXPORT_HEADER)
    assert labels == ["Trade Date", "Trade Time", "Strategy", "Symbol", "Product",
                      "Quantity", "Avg Price", "Current Price", "Market Value",
                      "Unrealized P&L", "Broker Qty", "System Qty", "Delta Qty",
                      "Reconciliation Status", "Source"]


def test_the_price_columns_export_as_not_instrumented_never_blank(gui_config):
    """⛔ A blank cell in a spreadsheet reads as zero, and a zero market value is
    a materially different claim from an unmeasured one."""
    p = holdings.build_holdings_screen(gui_config)
    rows = holdings.export_rows(p)
    hdr = rows[0]
    idx = [hdr.index(c) for c in ("Current Price", "Market Value", "Unrealized P&L")]
    assert rows[1:], "no exported rows — the check is vacuous"
    for row in rows[1:]:
        for i in idx:
            assert row[i] == holdings.NA, row


def test_no_export_cell_is_ever_blank(gui_config):
    p = holdings.build_holdings_screen(gui_config)
    for title, header, rows in holdings.export_sheets(p):
        assert all(h for h in header), title
        for row in rows:
            for cell in row:
                assert cell is not None and cell != "", (title, row)


def test_the_export_carries_the_orphans_and_the_gaps(gui_config):
    """⭐ A workbook that omitted them would let a reader total the measured
    columns and believe they account for the whole book."""
    p = holdings.build_holdings_screen(gui_config)
    sheets = {t: (h, r) for t, h, r in holdings.export_sheets(p)}
    assert set(sheets) == {"Holdings", "Orphans", "Not Instrumented"}
    o = sheets["Orphans"][1]
    assert len(o) == (p["orphans"]["broker_without_system"]["n"]
                      + p["orphans"]["system_without_broker"]["n"])
    gaps = sheets["Not Instrumented"][1]
    assert len(gaps) == 4
    assert all(len(g[2]) > 20 for g in gaps), "a gap exported without its reason"


def test_the_export_is_a_real_parseable_workbook(client):
    from openpyxl import load_workbook
    r = client.get("/api/export/holdings")
    assert r.status_code == 200
    assert r.data[:4] == b"PK\x03\x04"
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Holdings", "Orphans", "Not Instrumented"]
    assert [c.value for c in wb["Holdings"][1]] == list(holdings.EXPORT_HEADER)


def test_the_export_writes_only_the_filtered_rows(client):
    from openpyxl import load_workbook
    full = _s(client)
    filtered = _s(client, "?product=MIS")
    assert 0 < filtered["count"] < full["total"]
    r = client.get("/api/export/holdings?product=MIS")
    wb = load_workbook(io.BytesIO(r.data))
    assert wb["Holdings"].max_row == filtered["count"] + 1


# ═════════════════════════════════════════════════════════════════════════════
# THE PAGE
# ═════════════════════════════════════════════════════════════════════════════
def test_the_page_and_both_endpoints_are_reachable(client):
    assert client.get("/holdings").status_code == 200
    assert client.get("/api/holdings/screen").status_code == 200
    assert client.get("/api/export/holdings").status_code == 200


def test_the_legacy_holdings_endpoint_still_answers(client):
    """⭐ ADDITIVE — the G5d gtt_state-mirror endpoint keeps its contract."""
    d = client.get("/api/holdings").get_json()
    assert "rows" in d and "count" in d and "banner" in d


def test_the_approved_panels_are_all_on_the_page(client):
    page = _page(client)
    for title in ("HOLDINGS", "Broker-First Reconciliation Center",
                  "Broker = Final Source of Truth",
                  "HOLDINGS RECONCILIATION TABLE", "DELTA ANALYSIS",
                  "BROKER SNAPSHOT", "RECONCILIATION TIMELINE",
                  "POSITIONS BY PRODUCT", "HOLDINGS HEALTH SUMMARY",
                  "DELIVERY PORTFOLIO", "ORPHAN DETECTOR",
                  "SIDE-BY-SIDE COMPARISON", "HOLDINGS CATEGORIES",
                  "REJECTION / QUALITY INSIGHTS", "EXPORT"):
        assert title in page, title


def test_the_six_approved_kpi_cards_are_present_in_order():
    tpl = _tpl()
    block = tpl[tpl.index("KPIS: ["):tpl.index("DEFAULT_COLS:")]
    labels = re.findall(r'label:\s*"([^"]+)"', block)
    assert labels == ["BROKER HOLDINGS", "SYSTEM HOLDINGS", "MATCHED HOLDINGS",
                      "MISMATCHES", "TOTAL VALUE", "TOTAL UNREALIZED P&L"], labels


def test_the_column_widths_are_bound_to_the_column_key():
    tpl = _tpl()
    block = tpl[tpl.index("DEFAULT_COLS:"):tpl.index("/* the three the artwork")]
    for line in block.splitlines():
        if "key:" in line and "label:" in line:
            assert "w:" in line, line.strip()


def test_the_screen_has_no_write_path(client):
    """L4 — this dashboard never writes to the trading system, and the disabled
    Sync Now button must not be a form either."""
    page = _page(client)
    root = page[page.index('<div class="hld-page"'):]
    assert "<form" not in root.lower()
    assert 'method="post"' not in root.lower()
    assert 'method: "post"' not in root.lower()
    assert '"POST"' not in root and "'POST'" not in root


def test_the_screen_is_not_period_scoped(client):
    """⚠️ Holdings are a LIVE POSITION STATE."""
    d = _s(client)
    assert "period" not in d and "from" not in d
    assert d["today"]


def test_no_alpine_x_if_branch_ships_two_root_elements():
    tpl = _tpl()
    for block in re.findall(r'<template x-if="[^"]+">(.*?)</template>', tpl, re.S):
        depth, top = 0, 0
        for close, tag, attrs in re.findall(r"<(/?)(\w+)([^>]*)>", block):
            if close:
                depth -= 1
                continue
            if depth == 0:
                top += 1
            if not attrs.rstrip().endswith("/"):
                depth += 1
        assert top == 1, "an x-if branch has %d root elements: %s" % (top, block[:90])


def test_no_alpine_x_for_template_sits_inside_an_svg():
    tpl = _tpl()
    for svg in re.findall(r"<svg\b.*?</svg>", tpl, re.S):
        assert "<template" not in svg, svg[:140]


def test_the_screen_css_is_scoped_to_this_page():
    for line in _css_block().splitlines():
        line = line.strip()
        if not line or line.startswith(("/*", "*", "@", "}")) or "{" not in line:
            continue
        selector = line.split("{")[0].strip()
        assert ".hld-page" in selector or "main.content" in selector, selector


def test_the_readability_floor_holds_for_text():
    """⚠️ The donut labels are SVG USER UNITS — the floor is checked against the
    EFFECTIVE size, ⛔ not the raw declaration."""
    for rule in re.finditer(r"([^{}]+)\{([^{}]*)\}", _css_block()):
        selector, body = rule.group(1), rule.group(2)
        m = re.search(r"font-size:\s*([\d.]+)px", body)
        if not m:
            continue
        size = float(m.group(1))
        if "donut-t" in selector:
            size *= 112.0 / 42.0          # .hld-donut is painted at 112px
        assert size >= 12.9, (selector.strip(), m.group(0))


#: The ONE control size this screen may restate, and only because Screens 14 and
#: 15 already ship it verbatim for the SAME control. ⚠️ Compared as an exact
#: declaration, ⛔ not as a class-name exemption: a fourth height smuggled in
#: under the same class still fails.
_RPP_DECL = "height: 30px; padding: 0 8px;"


def test_no_one_off_control_sizing_is_introduced():
    """⛔ Height, padding, radius and background come from the SHARED rules.

    ⚠️ NARROWED 16-Aug for the artwork's rows-per-page select. It is genuinely
    smaller than a filter control in the artwork, and the codebase already has
    that control at ONE size — `.tlg-rpp .sel` and `.slg-rpp .sel`, both
    `height: 30px; padding: 0 8px`. Screen 22 reuses that declaration rather
    than inventing a fourth, and this guard now proves the reuse instead of
    forbidding the control.
    """
    for line in _css_block().splitlines():
        if ".sel" in line and "height" in line:
            if ".hld-rpp .sel" in line and _RPP_DECL in line:
                continue
            raise AssertionError("one-off control sizing: " + line.strip())


def test_the_rows_per_page_control_matches_screens_14_and_15_exactly():
    """⭐ The reuse is asserted against the OTHER screens' own lines, so if any
    of the three ever changes size the three stop matching and this goes red."""
    css = _css()
    decls = re.findall(r"\.(?:tlg|slg|hld)-(?:page )?\.?\w*rpp \.sel \{([^}]*)\}", css)
    assert len(decls) == 3, decls
    norm = {" ".join(d.split()) for d in decls}
    assert len(norm) == 1, norm
    assert _RPP_DECL.replace(";", "") in list(norm)[0].replace(";", "")


def test_the_page_joins_the_shared_filter_and_token_blocks():
    """🔴 MEASURED DEFECT on Screens 19/20: a page outside these blocks renders
    its filter controls with no background and no border, and every
    `rgba(var(--t-*), …)` tint fully transparent."""
    css = _css()
    head = css[:css.index("--t-blue: 56,139,253")]
    assert ".hld-page" in head[head.rindex("\n.dash-page"):]
    for anchor in (".flt-row {", ".flt-field {", ".flt-k {", ".flt-actions {"):
        i = css.index(anchor)
        selector = css[css.rindex("\n", 0, css.rindex("\n", 0, i) if "\n" in css[:i] else 0):i + len(anchor)]
        assert ".hld-page" in css[max(0, i - 400):i + len(anchor)], anchor


def test_no_javascript_string_literal_spans_a_newline():
    bad = _js_syntax.string_literals_spanning_a_newline(
        _js_syntax.script_of(_tpl()))
    assert not bad, bad


def test_the_page_script_has_balanced_braces():
    assert _js_syntax.unbalanced_braces(_js_syntax.script_of(_tpl())) == 0


# ═════════════════════════════════════════════════════════════════════════════
# THE EMPTY BOOK
# ═════════════════════════════════════════════════════════════════════════════
def test_an_empty_book_produces_no_invented_numbers(gui_config):
    c = _conn(gui_config)
    for tbl in ("trades", "orders", "position_reconciliation", "signals"):
        c.execute("DELETE FROM " + tbl)
    c.commit()
    c.close()
    p = holdings.build_holdings_screen(gui_config)
    assert p["rows"] == [] and p["count"] == 0 and p["total"] == 0
    k = p["kpi"]
    assert k["broker_holdings"] == 0 and k["system_holdings"] == 0
    assert k["matched"] == 0 and k["mismatches"] == 0
    assert k["matched_pct"] is None and k["mismatch_pct"] is None
    assert k["value_at_cost"] is None
    assert k["total_value"] is None and k["total_pnl"] is None
    assert p["delta"] == {"broker_qty": 0, "system_qty": 0, "delta_qty": 0,
                          "symbols": 0, "excluded": 0,
                          "base": p["delta"]["base"]}
    assert p["delivery"]["positions"] == 0
    assert p["orphans"]["broker_without_system"]["n"] == 0
    assert p["timeline"]["last_reconciliation"] is None
    assert p["reconciliation"]["available"] is False
    # the export still produces a real workbook with its headers
    sheets = holdings.export_sheets(p)
    assert [t for t, _h, _r in sheets] == ["Holdings", "Orphans",
                                           "Not Instrumented"]
    assert sheets[0][1] == list(holdings.EXPORT_HEADER)


# ═════════════════════════════════════════════════════════════════════════════
# THE 16-AUG OLD-DESIGN COMPLIANCE GATE — D5
# ═════════════════════════════════════════════════════════════════════════════
def test_d5_the_pagination_row_has_the_three_approved_controls():
    """⛔ THE DEFECT THIS PINS: only prev/next existed. The artwork draws
    `‹ [1] 2 3 ›` with the current page boxed AND a `10 / page` selector, and
    the rows-per-page control was missing entirely."""
    tpl = _tpl()
    foot = tpl[tpl.index('<div class="hld-tbl-foot">'):]
    foot = foot[:foot.index("</div>\n        <p")] if "</div>\n        <p" in foot else foot[:2600]
    assert "pageList()" in foot, "no numbered page buttons"
    assert "tPage = Math.max(1, tPage - 1)" in foot, "no previous control"
    assert "tPage = Math.min(tPageCount(rows()), tPage + 1)" in foot, "no next control"
    assert "hld-rpp" in foot and "tPageSizes" in foot, "no rows-per-page selector"
    assert "n + ' / page'" in foot, "the selector does not read 'N / page'"


def test_d5_pagination_reuses_the_existing_mixin_state():
    """⭐ `tPage` / `tPageSize` / `tPageSizes` are tableMixin's own — ⛔ no second
    pagination model was written, and `pageSizes` was already [10, 25, 50]."""
    tpl = _tpl()
    assert "pageSizes: [10, 25, 50]" in tpl
    js = _js_syntax.script_of(tpl)
    assert "tPage:" not in js and "tPageSize:" not in js, \
        "the screen redeclares pagination state instead of using the mixin"


def test_d5_the_page_list_is_the_artworks_numbers_and_cannot_overflow():
    """⭐ Every page is listed while they fit; beyond that it windows with an
    ellipsis so the row cannot outgrow its panel. ⛔ Not a different design."""
    tpl = _tpl()
    body = tpl[tpl.index("pageList() {"):]
    body = body[:body.index("\n    },")]
    assert "PAGE_WINDOW" in tpl
    assert '"…"' in body, body


def test_d5_at_28_rows_the_default_page_size_yields_three_pages(gui_config):
    """⭐ The artwork's own worked example: 28 holdings at 10/page = pages 1-3.
    Seeded here so the assertion has a real 28-row population, ⛔ not assumed."""
    c = _conn(gui_config)
    for i in range(28 - 7):                     # the fixture already carries 7
        sym = "PAG%02d" % i
        c.execute("INSERT INTO position_reconciliation(date,symbol,broker_qty,"
                  "system_qty,status,resolved_at,created_at) VALUES(?,?,?,?,?,?,?)",
                  (TODAY, sym, 10, 10, "OK", None, _ts("15:45:07")))
    c.commit()
    c.close()
    p = holdings.build_holdings_screen(gui_config)
    assert p["count"] == 28, p["count"]
    import math
    assert math.ceil(p["count"] / 10) == 3


def test_d5_the_rows_per_page_control_is_the_established_one():
    """⛔ Not a new size — Screens 14 and 15 already ship this exact control."""
    block = _css_block()
    assert ".hld-page .hld-rpp .sel" in block
    line = [l for l in block.splitlines() if ".hld-rpp .sel" in l][0]
    assert "height: 30px" in line and "padding: 0 8px" in line, line


def test_every_filter_resets_the_page_index():
    """⭐ A filter applied from page 3 must not strand the reader past the end
    of the smaller filtered set.

    ⛔ Without this, filtering 28 rows down to 21 while `tPage` is 3 renders
    "Showing 21 to 21 of 21 holdings" and ONE row — which a reader takes to
    mean the filter matched almost nothing, not that the page is stale.
    🔬 Measured in the browser on 01-Sep-2026 before the fix.
    """
    tpl = _tpl()
    handlers = re.findall(r'@(?:change|keyup\.enter)="([^"]*\bload\(\)[^"]*)"', tpl)
    assert handlers, "the filter controls no longer call load() — retarget this test"
    for h in handlers:
        assert "tPage = 1" in h, (
            "a filter handler calls load() without resetting the page: %r" % h)


def test_the_shared_refresh_does_not_reset_the_page_index():
    """⛔ The counterpart guard. `@ops-refresh.window` fires on every poll, so
    resetting `tPage` there would yank a reader back to page 1 mid-read — the
    reason the reset lives at the filter call sites and ⛔ NOT inside load()."""
    tpl = _tpl()
    m = re.search(r'@ops-refresh\.window="([^"]*)"', tpl)
    assert m, "the shared refresh binding is gone — retarget this test"
    assert "tPage" not in m.group(1), m.group(1)
    body = re.search(r"\n    load\(\)\s*\{(.*?)\n    \},", tpl, re.S)
    if body:
        assert "tPage = 1" not in body.group(1), "load() must not reset the page"


def test_a_stale_timestamp_is_rendered_with_its_date():
    """⛔ The broker side can be DAYS old — `reconcile_positions` writes a row
    only when there is a position to compare, so a flat book leaves the last
    row standing indefinitely.

    🔬 Measured 01-Sep-2026: `last_sync` was 2026-08-18T15:45:02, and the old
    `slice(11,19)` formatter drew it as a bare "15:45:02" beside a green ● CRON
    dot — indistinguishable from today at 15:45. ⭐ The formatter must compare
    the stamp's DAY against the payload's `today` and show the date when they
    differ; ⛔ a bare time-only slice is the defect.
    """
    tpl = _tpl()
    m = re.search(r"hhmmss\(ts\)\s*\{(.*?)\n    \},", tpl, re.S)
    assert m, "hhmmss() is gone or reshaped — retarget this test"
    body = m.group(1)
    assert "today" in body, "hhmmss() no longer compares the stamp against today"
    assert re.search(r"slice\(\s*0\s*,\s*10\s*\)", body), \
        "hhmmss() no longer extracts the stamp's date"
    one_line = re.match(r"\s*return ts \?", body)
    assert not one_line, "hhmmss() reverted to the bare time-only slice"


def test_the_screen_names_the_base_of_each_holdings_count(gui_config):
    """⭐ `system_holdings` counts CURRENTLY-OPEN positions while the table's
    System Qty comes from the last reconciliation RECORD. The two legitimately
    differ, so each must publish the base it counted — otherwise "0 system
    holdings" beside a row reading "System Qty 1" looks like a defect.
    """
    p = holdings.build_holdings_screen(gui_config)
    k = p["kpi"]
    assert k["system_base"], "the system count must name its base"
    assert k["broker_base"], "the broker count must name its base"
    assert k["system_base"] != k["broker_base"], "the two bases are not the same"
    for tile in p["health"]["tiles"]:
        if tile["measured"]:
            assert tile["base"], ("an unmeasured base: %s" % tile["label"])
