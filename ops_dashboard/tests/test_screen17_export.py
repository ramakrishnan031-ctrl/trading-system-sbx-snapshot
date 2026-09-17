"""SCREEN 17 — EXPORT (artwork panel 14).

Added 17-Aug-2026 with the panel itself. Every sibling screen export is tested
(06, 07, 10, 12, 21) and this one was not; a new route with no coverage is how
an auth regression or a silent shape drift ships unnoticed.

🔑 WHAT THESE PIN, beyond "it returns 200":
  (a) the route is login_required — an export is a data-egress path;
  (b) the workbook is built from the SAME payload the screen is served, so an
      exported cell cannot disagree with the panel above it;
  (c) ⛔ NO SCANNER ROW. scanner = strategy 1:1 (Rama, 17-Aug), so a scanner
      line would print one identity twice. This test is what stops it being
      reintroduced later "to match the artwork";
  (d) an unavailable value exports as an explicit marker, ⛔ never as a blank —
      a blank spreadsheet cell reads as zero, or as "nothing happened".
"""
from __future__ import annotations

import io

from backend.services import controls

_XLSX_MIME = ("application/vnd.openxmlformats-officedocument"
              ".spreadsheetml.sheet")


def test_export_requires_login(app):
    """⛔ An unauthenticated caller must not be able to pull the control data."""
    assert app.test_client().get("/api/export/controls-screen").status_code == 401


def test_export_returns_a_real_xlsx(client):
    r = client.get("/api/export/controls-screen")
    assert r.status_code == 200
    assert r.headers["Content-Type"].startswith(_XLSX_MIME)
    assert "attachment" in r.headers["Content-Disposition"]
    assert ".xlsx" in r.headers["Content-Disposition"]
    # A real workbook, not an empty or truncated body.
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames, "workbook has no sheets"


def test_sheets_agree_with_the_screen(client, gui_config):
    """The export and the screen are built from ONE payload, so they cannot drift."""
    payload = controls.build_controls_screen(gui_config)
    sheets = controls.export_sheets(payload)
    titles = [t for t, _h, _r in sheets]
    assert "Active Controls Summary" in titles
    assert "Strategy Controls" in titles

    by_title = {t: rows for t, _h, rows in sheets}
    # Every configured strategy appears exactly once — the control inventory.
    assert len(by_title["Strategy Controls"]) == payload["strategies"]["total"]

    summary = {str(k): v for k, v in by_title["Active Controls Summary"] if k}
    acs = payload["active_controls_summary"]
    assert summary["Active Strategies"] == acs["active_strategies"]
    assert summary["Trading Status"] == acs["trading_status"]


def test_no_scanner_identity_anywhere_in_the_export(client, gui_config):
    """⛔ scanner = strategy 1:1 — the export must not restate it under a second name."""
    sheets = controls.export_sheets(controls.build_controls_screen(gui_config))
    for title, header, rows in sheets:
        assert "scanner" not in title.lower(), f"scanner sheet reintroduced: {title}"
        blob = " ".join(str(c) for c in header).lower()
        assert "scanner" not in blob, f"scanner column reintroduced in {title}"
        for row in rows:
            first = str(row[0]).lower() if row and row[0] is not None else ""
            assert "scanner" not in first, f"scanner row reintroduced in {title}: {row[0]}"


def test_unavailable_exports_as_an_explicit_marker_never_blank(client, gui_config):
    """A missing value must SAY it is missing. ⛔ A blank cell reads as zero."""
    payload = controls.build_controls_screen(gui_config)
    # Force the gap the fixture cannot produce: no recorded control action.
    payload["active_controls_summary"]["last_control_change"] = None
    sheets = dict((t, rows) for t, _h, rows in controls.export_sheets(payload))
    summary = {str(k): v for k, v in sheets["Active Controls Summary"] if k}
    assert summary["Last Control Change"] == "UNAVAILABLE"
    assert summary["Last Control Change"] not in ("", None)
