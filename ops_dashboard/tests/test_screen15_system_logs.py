"""SCREEN 15 — SYSTEM LOGS.

The approved design (`gui/15. System_Logs.png` + `.txt`) is binding for
STRUCTURE; the DATA is this system's.

⭐ THE CENTRAL PROPERTY: three sources answer DIFFERENT parts of the approved
table, and each field must come from the one that actually records it. A field
no source records must render NOT INSTRUMENTED — and a test here must fail if it
ever silently becomes a plausible value.
"""
from __future__ import annotations

import io
import json
import os
import re

from backend.services import system_logs


def _read(*parts) -> str:
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    with open(os.path.join(here, *parts), encoding="utf-8") as fh:
        return fh.read()


def _tpl() -> str:
    return _read("frontend", "templates", "logs.html")


def _css() -> str:
    return _read("frontend", "static", "style.css")


def _shipped(client) -> str:
    html = client.get("/logs").data.decode("utf-8")
    return re.sub(r"<!--.*?-->", " ", html, flags=re.S)


def _page_content(client) -> str:
    """This screen's VISIBLE markup — shared chrome and scripts removed.

    🔴 MATCH THE ROOT BY ITS CLASS, ⛔ never by one exact attribute string.
    `0ca38e2` moved S15 onto the wide-page layout, so the root became
    `class="dash-page slg-page"` — and the literal `'<div class="slg-page"'`
    stopped matching. It did not fail loudly: it returned -1 and took EIGHT
    tests down with it, every one of them reporting the page root as missing
    rather than the real change. A class-aware match cannot be blinded by a
    class being added alongside.
    """
    body = _shipped(client)
    m = re.search(r'<div\s[^>]*class="[^"]*\bslg-page\b[^"]*"', body)
    assert m, "the system-logs page root is missing"
    body = body[m.start():]
    j = body.find("<script>")
    return body[:j] if j > 0 else body


WIDE = "?start=2000-01-01&end=2099-12-31"


def _s(client, qs=WIDE):
    return client.get("/api/system-logs/screen" + qs).get_json()


# ── the approved vocabularies ───────────────────────────────────────────────
def test_the_approved_vocabularies_are_exact():
    assert system_logs.SERVICES == (
        "Signal Engine", "Risk Engine", "Capital Engine", "Order Engine",
        "Broker Connector", "Database", "Dashboard API", "Scheduler",
        "Recovery Engine", "System Monitor")
    assert system_logs.SEVERITIES == ("Info", "Warning", "Error", "Critical")
    assert system_logs.EVENT_TYPES == (
        "Service Started", "Service Stopped", "Service Restarted",
        "Connection Established", "Connection Lost", "Database Error",
        "API Error", "Timeout", "Authentication Failure",
        "Recovery Triggered", "Recovery Completed")
    assert system_logs.TIMELINE_STAGES == (
        "Event Detected", "Logged", "Action Taken", "Resolved")
    assert system_logs.TRADING_IMPACTS == (
        "No Impact", "Minor Impact", "Major Impact", "Trading Halt Risk")


def test_the_services_panel_lists_the_ten_approved_names(client):
    panel = _s(client)["services"]
    assert [s["service"] for s in panel] == list(system_logs.SERVICES)


# ── ⛔ the gaps, and they must never become values ──────────────────────────
def test_trading_impact_is_never_inferred(client):
    """⭐ THE MOST IMPORTANT TEST ON THIS SCREEN. Nothing classifies an event's
    trading impact. Deriving it from severity would be a guess dressed as a
    measurement — an ERROR in a report generator and an ERROR in the order path
    are not the same risk, and nothing stored tells them apart."""
    p = _s(client)
    assert p["gaps"]["trading_impact"]["measured"] is False
    for r in p["records"]:
        assert r["trading_impact"] is None
    body = _page_content(client)
    assert "TRADING IMPACT" in body
    assert "NOT INSTRUMENTED" in body
    # ⭐ The four approved classes are still SHOWN — they are data-driven, so the
    # payload is where they live; the panel iterates them. ⛔ Asserting on the
    # server HTML would test the wrong artefact, because Alpine renders the list.
    assert p["trading_impacts"] == list(system_logs.TRADING_IMPACTS)
    assert 'x-for="t in data.trading_impacts' in _tpl()


def test_error_code_and_resolution_time_are_never_invented(client):
    p = _s(client)
    for r in p["records"]:
        assert r["error_code"] is None
        assert r["resolution_time"] is None
    assert p["gaps"]["error_code"]["measured"] is False
    assert p["gaps"]["resolution"]["measured"] is False


def test_uninstrumented_event_types_show_no_count(client):
    """⛔ A type with no source is NOT zero — it is unmeasurable. A count beside
    the words NOT INSTRUMENTED would be a contradiction on screen."""
    p = _s(client)
    uninstrumented = [t for t in p["event_types"]
                      if t not in p["instrumented_event_types"]]
    assert uninstrumented, "every type instrumented? then this screen changed"
    for t in uninstrumented:
        assert p["by_event_type"][t] == 0
        for r in p["records"]:
            assert r["event_type"] != t


def test_the_instrumented_list_matches_what_the_builders_can_emit():
    """⚠️ The list and the emitters must agree, or a type marked uninstrumented
    would carry rows. Every value `_from_*` can produce must be declared."""
    src = _read("backend", "services", "system_logs.py")
    emitted = set(re.findall(r'event_type="([^"]+)"', src))
    emitted |= set(re.findall(r'"[A-Z_]+": "([A-Za-z ]+)",', src))
    emitted = {e for e in emitted if e in system_logs.EVENT_TYPES}
    assert emitted <= set(system_logs.INSTRUMENTED_EVENT_TYPES), (
        emitted - set(system_logs.INSTRUMENTED_EVENT_TYPES))


def test_every_declared_gap_carries_a_real_reason(client):
    gaps = _s(client)["gaps"]
    for key in ("trading_impact", "event_type", "error_code", "resolution",
                "timeline", "recovery_duration", "debug_log"):
        assert gaps[key]["measured"] is False
        assert gaps[key]["value"] is None
        assert len(gaps[key]["reason"]) > 30, key


# ── ⛔ no Scanner FIELD (but a real message is never censored) ──────────────
def test_there_is_no_scanner_field_anywhere(client):
    """⛔ The requirement is the absence of a Scanner FIELD — column, filter,
    detail row, export column. ⚠️ It is NOT a grep of the message text:
    production genuinely logs `warnings=['scanner_unreachable']`, and redacting
    a real log message to satisfy a string search would falsify the record."""
    tpl = _tpl()
    cols = re.search(r"DEFAULT_COLS:\s*\[(.*?)\n    \],", tpl, re.S).group(1)
    sels = re.search(r"SELECTS:\s*\[(.*?)\n    \],", tpl, re.S).group(1)
    assert "scanner" not in cols.lower()
    assert "scanner" not in sels.lower()
    assert "Scanner" not in system_logs.EXPORT_HEADER
    for r in _s(client)["records"]:
        assert "scanner" not in {k.lower() for k in r}


def test_the_approved_eight_columns_in_the_approved_order():
    tpl = _tpl()
    cols = re.search(r"DEFAULT_COLS:\s*\[(.*?)\n    \],", tpl, re.S).group(1)
    assert re.findall(r'key:\s*"([a-z_]+)"', cols) == [
        "date", "time", "service", "module", "severity", "event_type",
        "status", "message"]
    assert re.findall(r'label:\s*"([^"]+)"', cols) == [
        "Date", "Time", "Service", "Module", "Severity", "Event Type",
        "Status", "Message"]


def test_the_filters_are_the_approved_set():
    tpl = _tpl()
    sels = re.search(r"SELECTS:\s*\[(.*?)\n    \],", tpl, re.S).group(1)
    assert re.findall(r'key:\s*"([a-z_]+)"', sels) == [
        "service", "module", "severity", "event_type", "status"]
    assert 'x-model="f.start"' in tpl and 'x-model="f.end"' in tpl


# ── the approved SEARCH section (binding TXT §SEARCH) ──────────────────────
def test_the_search_panel_has_the_five_approved_fields(client):
    """The binding TXT requires a SEARCH section with these five. ⛔ The small
    table-toolbar search does NOT satisfy it — that is a global box, not the
    approved per-field panel, and both are present."""
    body = _page_content(client)
    i = body.find("SEARCH")
    assert i > 0, "the approved SEARCH panel is missing"
    panel = body[i:i + 2600]
    for field in ("Service Name", "Module", "Error Code", "Message",
                  "Reference ID"):
        assert field in panel, field
    # the toolbar search is retained alongside, not substituted for it
    assert "Search logs" in body


def test_search_by_service_name_narrows_the_same_population(client):
    p = _s(client)
    svc = next((r["service"] for r in p["records"] if r["service"]), None)
    assert svc, "fixture produced no service-attributed event"
    narrow = _s(client, WIDE + "&service_q=" + svc.split()[0])
    assert 0 < narrow["count"] <= p["count"]
    assert all(svc.split()[0].lower() in (r["service"] or "").lower()
               for r in narrow["records"])
    assert narrow["kpi"]["total_events"] == narrow["count"]


def test_search_by_reference_id_finds_that_exact_record(client):
    """⭐ Reference ID IS instrumented — every event carries a deterministic id
    built from its real source, so the approved field genuinely works."""
    p = _s(client)
    ref = p["records"][0]["ref_id"]
    narrow = _s(client, WIDE + "&ref_q=" + ref.replace("#", "%23"))
    assert narrow["count"] >= 1
    assert any(r["ref_id"] == ref for r in narrow["records"])


def test_search_by_module_and_message_are_contains_matches(client):
    """⛔ CONTAINS, not exact: the FILTERS dropdowns are exact because they list
    values that exist, but a typed fragment must not silently match nothing."""
    p = _s(client)
    mod = next((r["module"] for r in p["records"] if r["module"]), None)
    assert mod
    frag = str(mod)[: max(3, len(str(mod)) // 2)]
    narrow = _s(client, WIDE + "&module_q=" + frag)
    assert narrow["count"] >= 1
    assert all(frag.lower() in (r["module"] or "").lower()
               for r in narrow["records"])


def test_the_error_code_search_field_is_shown_but_not_queryable(client):
    """⛔ The approved field is SHOWN because the design requires it, and it is
    DISABLED because no error-code scheme exists — a box that accepted one could
    only ever return nothing. It says NOT INSTRUMENTED instead of pretending."""
    body = _page_content(client)
    i = body.find("Error Code")
    assert i > 0
    assert "disabled" in body[i:i + 700]
    assert "NOT INSTRUMENTED" in body[max(0, i - 200):i + 400]
    # ⛔ and the builder accepts no error-code parameter at all
    src = _read("backend", "services", "system_logs.py")
    assert 'f.get("error_code_q")' not in src
    assert _s(client)["gaps"]["error_code"]["measured"] is False


def test_search_and_filters_are_separate_parameters():
    """⛔ Binding a typed fragment to a dropdown's EXACT match would make the
    search return nothing for every partial word."""
    src = _read("backend", "services", "system_logs.py")
    body = re.search(r"def _matches\(.*?\n\n\n", src, re.S).group(0)
    assert 'has("service", f.get("service_q"))' in body
    assert 'eq("service", f.get("service"))' in body


# ── the reader: three formats, and only one of them is this screen's ────────
def test_only_the_dated_json_system_logs_are_read():
    """⚠️ MEASURED: the logs directory holds THREE formats — JSON system logs,
    PLAIN-TEXT debug logs, and a different JSON shape in the alert-failure logs
    (`severity`/`title`/`source_module`, no `level`). Reading them all as system
    logs produced 405 rows with a NULL severity and a NULL module."""
    from backend.readers import log_reader
    assert log_reader._SYSTEM_STEMS == ("system", "reconciler", "trades")
    rx = log_reader._syslog_re()
    assert rx.match("system_2026-08-03.log")
    assert rx.match("reconciler_2026-08-03.log")
    assert rx.match("trades_2026-08-03.log")
    assert not rx.match("debug_2026-08-03.log")      # plain text
    assert not rx.match("failed_alerts.log")         # different shape
    assert not rx.match("test_failed.log")


def test_a_foreign_shaped_record_is_skipped_and_counted(tmp_path):
    """⛔ Never rendered as an event with a null severity — skipped, and the
    count RETURNED so the screen can say what it did not read."""
    from backend.readers import log_reader
    d = tmp_path / "logs"
    d.mkdir()
    (d / "system_2026-08-03.log").write_text(
        '{"ts":"2026-08-03T10:00:00+05:30","level":"INFO","logger":"main","msg":"ok"}\n'
        '{"ts":"2026-08-03T10:00:01+05:30","severity":"CRITICAL","title":"x",'
        '"source_module":"fund_manager"}\n'
        'not json at all\n', encoding="utf-8")
    out = log_reader.read_system_logs({"paths": {"logs_dir": str(d)}},
                                      "2026-08-01", "2026-08-31")
    assert len(out["rows"]) == 1
    assert out["rows"][0]["logger"] == "main"
    assert out["skipped_foreign_shape"] == 1
    assert out["skipped_not_json"] == 1


def test_a_log_reference_id_is_deterministic(tmp_path):
    """⛔ No clock and no counter — a re-read addresses the same record, so a
    poll cannot renumber the feed."""
    from backend.readers import log_reader
    d = tmp_path / "logs"
    d.mkdir()
    (d / "system_2026-08-03.log").write_text(
        '{"ts":"2026-08-03T10:00:00+05:30","level":"INFO","logger":"main","msg":"a"}\n',
        encoding="utf-8")
    cfg = {"paths": {"logs_dir": str(d)}}
    a = log_reader.read_system_logs(cfg, "2026-08-01", "2026-08-31")["rows"]
    b = log_reader.read_system_logs(cfg, "2026-08-01", "2026-08-31")["rows"]
    assert [r["ref_id"] for r in a] == [r["ref_id"] for r in b]
    assert a[0]["ref_id"].startswith("LOG-system_2026-08-03.log#")


# ── determinism / one population ───────────────────────────────────────────
def test_repeated_builds_neither_duplicate_nor_lose_events(client):
    first = _s(client)["records"]
    for _ in range(3):
        again = _s(client)["records"]
        assert len(again) == len(first)
        assert [r["ref_id"] for r in again] == [r["ref_id"] for r in first]


def test_reference_ids_are_unique_and_name_their_source(client):
    rows = _s(client)["records"]
    ids = [r["ref_id"] for r in rows]
    assert len(set(ids)) == len(ids)
    assert {i.split("-")[0] for i in ids} <= {"LOG", "SYSEVT", "CRON", "RECOV"}


def test_equal_timestamps_order_deterministically(client):
    rows = _s(client)["records"]
    keyed = [(r["ts"], r["ref_id"]) for r in rows]
    assert keyed == sorted(keyed, reverse=True)


def test_kpi_and_panels_reconcile_to_the_table(client):
    p = _s(client)
    rows = p["records"]
    k = p["kpi"]
    assert k["total_events"] == len(rows) == p["count"]
    for sev, key in (("Info", "info"), ("Warning", "warnings"),
                     ("Error", "errors"), ("Critical", "critical")):
        assert k[key] == sum(1 for r in rows if r["severity"] == sev)
    assert sum(p["by_severity"].values()) == p["count"]


def test_the_donut_closes_at_one_hundred_percent(client):
    p = _s(client)
    if not p["count"]:
        return
    assert abs(sum(d["pct"] for d in p["severity_donut"]) - 100.0) < 0.05
    for d in p["severity_donut"]:
        assert d["count"] == p["by_severity"][d["severity"]]


def test_a_filter_narrows_every_panel_together(client):
    wide = _s(client)
    narrow = _s(client, WIDE + "&severity=Info")
    assert narrow["count"] <= wide["count"]
    assert all(r["severity"] == "Info" for r in narrow["records"])
    assert narrow["kpi"]["total_events"] == narrow["count"]
    assert sum(narrow["by_severity"].values()) == narrow["count"]


def test_retention_describes_the_store_not_the_filter(client):
    wide = _s(client)["retention"]
    narrow = _s(client, "?start=2099-01-01&end=2099-01-02")["retention"]
    assert wide["records"] == narrow["records"]


# ── detail / timeline / replay ─────────────────────────────────────────────
def _a_ref(client):
    rows = _s(client)["records"]
    assert rows, "fixture produced no system log events"
    return rows[0]["ref_id"]


def test_the_timeline_has_the_four_approved_stages_in_order(client):
    d = client.get("/api/system-logs/detail" + WIDE
                   + "&ref_id=" + _a_ref(client)).get_json()
    assert [s["stage"] for s in d["timeline"]] == list(system_logs.TIMELINE_STAGES)


def test_only_logged_is_measured_for_a_plain_log_event(client):
    """⛔ A log line has ONE timestamp. Detection, action and resolution are not
    separately stamped, so they come back unavailable WITH their reason."""
    d = client.get("/api/system-logs/detail" + WIDE
                   + "&ref_id=" + _a_ref(client)).get_json()
    by = {s["stage"]: s for s in d["timeline"]}
    assert by["Logged"]["measured"] is True
    assert by["Event Detected"]["measured"] is False
    assert by["Event Detected"]["gap"]


def test_a_job_runtime_is_never_relabelled_as_a_recovery_duration(client):
    """🔴 CAUGHT IN THE BROWSER: a scheduled job's `duration_sec` was rendering
    under RECOVERY TRACKING as "Recovery Duration 4.2 sec" on a row that was not
    a recovery at all. Real number, wrong name — the quantity is a job runtime."""
    p = _s(client)
    ref = next((r["ref_id"] for r in p["records"]
                if r["ref_id"].startswith("CRON-")), None)
    assert ref, "fixture has no scheduler row"
    d = system_logs.event_detail(p, ref)
    assert d["recovery"]["duration_sec"] is None
    assert d["recovery"]["triggered"] is None
    assert d["recovery"]["result"] is None
    # ⭐ the real number is kept, under its own name
    assert "job_duration_sec" in d


def test_a_non_error_event_has_no_resolution_status(client):
    """⛔ 'Resolution Status: Success' beside 'Error Message: —' would imply an
    error was investigated and closed when none was raised."""
    p = _s(client)
    ref = next((r["ref_id"] for r in p["records"]
                if r["severity"] == "Info"
                and not r["ref_id"].startswith("RECOV-")), None)
    assert ref, "fixture has no non-error, non-recovery row"
    d = system_logs.event_detail(p, ref)
    assert d["error"]["error_message"] is None
    assert d["error"]["resolution_status"] is None


def test_replay_returns_the_recorded_sequence_not_a_simulation(client):
    ref = _a_ref(client)
    d = client.get("/api/system-logs/detail" + WIDE + "&ref_id=" + ref).get_json()
    replay = d["replay"]
    assert replay, "replay must contain at least the focused event"
    assert sum(1 for r in replay if r["is_focus"]) == 1
    refs = {r["ref_id"] for r in replay}
    known = {r["ref_id"] for r in _s(client)["records"]}
    assert refs <= known          # ⛔ nothing invented into the sequence
    assert [r["ts"] for r in replay] == sorted(r["ts"] for r in replay)


def test_an_unknown_reference_id_is_a_404(client):
    assert client.get("/api/system-logs/detail" + WIDE
                      + "&ref_id=nope").status_code == 404


# ── XLSX ───────────────────────────────────────────────────────────────────
def test_export_is_the_filtered_set_in_the_approved_order(client):
    from openpyxl import load_workbook
    r = client.get("/api/export/system-logs" + WIDE + "&severity=Info")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.data))["System Log Events"]
    header = [c.value for c in ws[1]]
    assert header[:8] == ["Date", "Time", "Service", "Module", "Severity",
                          "Event Type", "Status", "Message"]
    assert "Scanner" not in header
    body = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(body) == _s(client, WIDE + "&severity=Info")["count"]
    assert {row[4] for row in body} <= {"Info"}
    # ⛔ an unrecorded field exports as NOT INSTRUMENTED, never as a blank cell
    assert all(row[6] for row in body)


# ── layout / reuse / conventions ───────────────────────────────────────────
def test_every_approved_panel_is_on_the_page(client):
    body = _page_content(client)
    for heading in ("SYSTEM LOGS", "FILTERS", "SYSTEM LOG EVENTS", "SERVICES",
                    "SEVERITY LEVELS", "TRADING IMPACT", "EVENT TYPES",
                    "EVENT DETAILS", "COMPONENT TIMELINE", "ERROR INVESTIGATION",
                    "RECOVERY TRACKING", "DEPENDENCY STATUS", "SYSTEM REPLAY",
                    "RETENTION OVERVIEW", "EXPORT"):
        assert heading in body, heading
    for control in ("Reset Filters", "Apply Filters", "Replay Incident",
                    "Export to XLSX", "Rows per page", "Search logs"):
        assert control in body, control


def test_the_six_approved_dependencies_are_shown(client):
    names = {d["name"] for d in _s(client)["dependencies"]}
    # Screen 12 names the broker dependency "Broker API"; the rest match exactly
    assert {"Chartink", "Database", "Tailscale", "Internet", "VM"} <= names
    assert any("Broker" in n for n in names)


def test_header_and_body_share_one_column_list_so_they_cannot_desync():
    assert _tpl().count('x-for="c in cols"') == 2


def test_column_drag_reuses_the_established_convention():
    tpl = _tpl()
    for hook in ("onDragStart", "onDrop", "onDragEnd", "initCols", "saveCols",
                 "resetCols", "isDefaultOrder", "localStorage"):
        assert hook in tpl, hook
    assert "screen15.systemlogs.colOrder.v1" in tpl
    assert "next.length === this.DEFAULT_COLS.length" in tpl


def test_a_stored_order_can_permute_but_never_change_the_column_set():
    body = re.search(r"initCols\(\)\s*\{(.*?)\n    \},", _tpl(), re.S).group(1)
    assert "byKey[k]" in body
    assert "next.indexOf(byKey[k]) === -1" in body
    assert "next.length === this.DEFAULT_COLS.length" in body


def test_the_rail_holds_only_the_three_panels_the_png_places_there():
    """🔴 THE LAYOUT DEFECT, PINNED. Stacking SEARCH and EVENT TYPES into the
    rail made it overshoot the left column by ~780px (measured off the render),
    so everything below waited for the grid to close and a dead band opened
    under the table. The binding PNG's rail is Services / Severity / Trading
    Impact — three panels, ending level with the table."""
    tpl = _tpl()
    rail = tpl[tpl.index('<aside class="slg-rail"'):tpl.index("</aside>")]
    panels = re.findall(r'class="panel slg-([a-z-]+)"', rail)
    assert panels == ["svcs", "sev-p", "impact"], panels
    # ⛔ THE STANDALONE ROW IS GONE, ⛔ not merely hidden. Moving SEARCH and
    # EVENT TYPES out of the rail fixed the overshoot but cost a whole row's
    # height, and the binding PNG places NEITHER panel. They now sit inside the
    # FILTERS panel — the one control block the artwork does draw.
    for dead in ('class="slg-rowx"', "slg-searchp", "slg-etypes"):
        assert dead not in tpl, dead
    # ⭐ and BOTH survive there in full — ⛔ integration, never deletion
    i = tpl.index('<section class="panel slg-filters">')
    filters = tpl[i:tpl.index("</section>", i)]
    assert "SEARCH" in filters and "EVENT TYPES" in filters
    for field in ("Service Name", "Module", "Error Code", "Message",
                  "Reference ID"):
        assert field in filters, field


def test_no_x_if_branch_has_two_root_elements():
    """🔴 THE BLANK-CELL DEFECT, PINNED — and it is a CORRECTNESS rule, not
    style. Alpine 3 builds an x-if branch with
    `content.cloneNode(true).firstElementChild` and DISCARDS every later
    sibling. Event Type and Status each held TWO spans: the value span, and the
    NOT INSTRUMENTED span. Only the first was ever created, and it carried
    x-show="r.status" — display:none exactly when the value was missing. So the
    gap state rendered as an EMPTY CELL, the one thing this screen must never
    do, and it did it silently: no error, no warning, no failing test."""
    tpl = _tpl()

    def root_count(inner: str) -> int:
        depth = n = 0
        void = ("br", "img", "input", "hr", "meta", "link")
        for m in re.finditer(r"<(/?)([a-zA-Z][\w-]*)([^>]*?)(/?)>", inner):
            closing, tag, _attrs, selfclose = m.groups()
            if closing:
                depth -= 1
            elif selfclose or tag.lower() in void:
                if depth == 0:
                    n += 1
            else:
                if depth == 0:
                    n += 1
                depth += 1
        return n

    offenders = [
        m.group(0)[:60]
        for m in re.finditer(r"<template\s+x-if=[^>]*>(.*?)</template>", tpl, re.S)
        if root_count(m.group(1)) > 1
    ]
    assert not offenders, offenders

    # ⭐ and the two cells state the gap on ONE element, so it cannot recur
    for key in ("status", "event_type"):
        blk = re.search(r"x-if=\"c\.key === '" + key + r"'\">(.*?)</template>",
                        tpl, re.S)
        assert blk, key
        assert "NOT INSTRUMENTED" in blk.group(1), key
        assert "x-show" not in blk.group(1), key


def test_the_component_timeline_is_the_pngs_horizontal_diagram():
    """⛔ The binding TXT forbids replacing a pictorial element with plain text
    where the original design shows a visual component, and the PNG draws the
    four stages as circular nodes on a horizontal rail. A vertical <ol> was
    doing exactly what the TXT forbids."""
    tpl, css = _tpl(), _css()
    assert "slg-tlh-track" in tpl
    assert 'class="slg-tl"' not in tpl, "the vertical list is still there"
    # four pictorial nodes, one per approved stage, injected as real SVG
    assert tpl.count("<svg viewBox=\"0 0 16 16\"", tpl.index("TL_ICONS:")) >= 4
    assert 'x-html="stageIcon(i)"' in tpl
    # laid out as four columns joined by a rail drawn off the nodes themselves
    assert "grid-template-columns: repeat(4, 1fr)" in css
    assert ".slg-tlh-stage::before" in css
    # ⛔ and an unmeasured stage still says so rather than showing a time
    seg = tpl[tpl.index("slg-tlh-track"):tpl.index("slg-tlh-foot")]
    assert 'x-show="!s.measured"' in seg and "NOT INSTRUMENTED" in seg
    # ⛔ no fabricated duration: the PNG's "2 sec" has no source here
    assert "Duration" in tpl and "tlResolved()" in tpl


def test_no_artificial_height_device_closes_a_gap():
    """⛔ A min-height / fixed height / stretch would HIDE a layout imbalance
    rather than fix it. The columns are balanced by what is IN them."""
    css = _css()
    start = css.index("SCREEN 15 — SYSTEM LOGS")
    nxt = re.search(r"SCREEN \d+ [—-]", css[start + 20:])
    block = css[start: start + 20 + nxt.start()] if nxt else css[start:]
    for rule in re.findall(r"^\.slg-page[^{]*\{[^}]*\}", block, re.M | re.S):
        if ".slg-row" in rule.split("{")[0] or ".panel" in rule.split("{")[0]:
            assert "min-height" not in rule, rule
            assert "height:" not in rule, rule
    assert "align-items: start" in block      # ⛔ never `stretch` on these rows


def test_alignment_is_by_role_and_not_blanket_centred():
    css = _css()
    start = css.index("SCREEN 15 — SYSTEM LOGS")
    nxt = re.search(r"SCREEN \d+ [—-]", css[start + 20:])
    block = css[start: start + 20 + nxt.start()] if nxt else css[start:]
    assert ".slg-tbl th, .slg-page .slg-tbl td { text-align: left; }" in block
    cols = re.search(r"DEFAULT_COLS:\s*\[(.*?)\n    \],", _tpl(), re.S).group(1)
    assert cols.count("ctr: true") == 2      # Severity + Status only


def test_every_table_mixin_member_the_page_calls_actually_exists():
    """⚠️ Carried forward: Alpine swallows an undefined-method binding, so this
    class of bug is invisible to every contract test. It bit Screen 13 twice."""
    tpl = _tpl()
    mixin = _read("frontend", "static", "components.js")
    defined = set(re.findall(r"^\s{4}(t[A-Za-z]+)\s*:", mixin, re.M))
    used = set(re.findall(r"\b(t[A-Z][A-Za-z]*)\b", tpl))
    unknown = sorted(u for u in used if u not in defined)
    assert not unknown, "template calls table-mixin members that do not exist: %s" % unknown


def test_every_page_method_the_markup_calls_is_defined_on_the_component():
    tpl = _tpl()
    markup = tpl[:tpl.index("<script>")]
    defined = set(re.findall(r"^\s{4}(?:async\s+)?([a-zA-Z_][a-zA-Z0-9_]*)\(", tpl, re.M))
    defined |= set(re.findall(r"^\s{4}([a-zA-Z_][a-zA-Z0-9_]*):", tpl, re.M))
    mixin = _read("frontend", "static", "components.js")
    defined |= set(re.findall(r"^\s{4}([a-zA-Z_][a-zA-Z0-9_]*)\s*[:(]", mixin, re.M))
    defined |= {"load", "Math", "Date", "JSON", "Number", "String", "Object",
                "encodeURIComponent", "fetch", "includes"}
    called = set(re.findall(
        r'(?:x-text|x-show|x-html|@click|:class|:href|:disabled|x-model|:title)='
        r'"[^"]*?(?<![.\w])([a-zA-Z_][a-zA-Z0-9_]*)\(', markup))
    unknown = sorted(c for c in called if c not in defined)
    assert not unknown, "markup calls undefined page methods: %s" % unknown


def test_real_time_uses_the_existing_refresh_path(client):
    body = _shipped(client)
    assert "ops-refresh.window" in body
    assert 'pageBase("/api/system-logs/screen")' in body


def test_the_m13_raw_log_endpoint_is_untouched(client):
    """⭐ ADDITIVE: this screen reads THROUGH the hardened log_reader; the raw
    tail endpoint keeps its own contract."""
    assert client.get("/api/logs").status_code == 200


# ── IST / readability / scoping ────────────────────────────────────────────
def test_timestamps_are_ist_and_never_converted(client):
    src = _read("backend", "services", "system_logs.py")
    assert "astimezone" not in src and "utcnow" not in src
    for r in _s(client)["records"]:
        assert re.match(r"\d{4}-\d{2}-\d{2}$", r["date"])
        assert re.match(r"\d{2}:\d{2}:\d{2}$", r["time"])


def test_text_meets_the_thirteen_pixel_floor():
    css = _css()
    start = css.index("SCREEN 15 — SYSTEM LOGS")
    nxt = re.search(r"SCREEN \d+ [—-]", css[start + 20:])
    block = css[start: start + 20 + nxt.start()] if nxt else css[start:]
    small = [float(m) for m in re.findall(r"font-size:\s*(\d+(?:\.\d+)?)px", block)
             if float(m) < 13]
    assert not small, small
    for rule in (".slg-page th { font-size: 13px; }",
                 ".slg-page .mono { font-size: 13px; }",
                 ".slg-page .rt-btn { font-size: 13px; }",
                 ".slg-page .panel-sub { font-size: 13px; }"):
        assert rule in block, rule


def test_the_screen_is_scoped_and_cannot_repaint_another():
    css = _css()
    start = css.index("SCREEN 15 — SYSTEM LOGS")
    nxt = re.search(r"SCREEN \d+ [—-]", css[start + 20:])
    block = css[start: start + 20 + nxt.start()] if nxt else css[start:]
    for line in block.splitlines():
        line = line.strip()
        if line.startswith(".") and "{" in line:
            assert line.startswith(".slg-page"), line


def test_no_secret_reaches_the_screen(client):
    blob = json.dumps(_s(client)).lower()
    for secret in ("password", "totp", "api_key", "access_token", "secret"):
        assert secret not in blob
