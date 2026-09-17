#!/usr/bin/env python
"""deploy/hooks/secret_scan.py -- pre-commit secret scanner (C-1 recurrence guard).

Blocks a commit that stages a real credential. Layers:
  1. A real `.env` file must never be committed.
  2. PATTERN scan (any staged text file): telegram bot-token shape, and a
     credential-style assignment (KEY named like *API_KEY/API_SECRET/TOKEN/SECRET/
     PASSWORD/TOTP*) whose VALUE looks like a real secret (>=16 chars of key-ish
     charset, or 64-hex, or base32) and is NOT a recognised placeholder.
  3. `.env.example` must not carry a real credential -- a credential-named key (or a
     secret-looking value) must be a placeholder. This is the exact failure C-1
     exposed (real values committed in the template + git history). Non-secret
     template defaults (LOG_LEVEL=INFO) are intentionally allowed.
  4. `.xlsx` workbooks are parsed cell-by-cell (`scan_xlsx`) and scanned for the same
     secret shapes -- a binary spreadsheet (e.g. `credentials.xlsx`) would otherwise
     bypass the text scanner. FAIL CLOSED: an un-inspectable .xlsx is BLOCKED.

`scan_content(path, text) -> list[str]` is a PURE function (no git) so it is unit
-testable. `scan_staged()` scans `git diff --cached` content and is what the
pre-commit hook runs. Findings NEVER contain the secret value -- only file:line +
the rule name, so the scanner's own output can't leak a secret.

Deliberately conservative to avoid false positives: only credential-NAMED
assignments (or the unambiguous telegram/hex shapes) are flagged, and any value that
matches a placeholder / env-var reference / template token is allowed.
"""
from __future__ import annotations

import os
import re
import subprocess
import sys

# ── Recognised placeholders / non-literal references (ALLOWED values) ──────────
_PLACEHOLDER_RE = re.compile(
    r'^(?:["\']?)'
    r'(?:fill[_-]?\w*|your[_-]?\w*|<[^>]*>|change[_-]?me\w*|replace[_\w]*|example\w*|'
    r'dummy\w*|placeholder\w*|sample\w*|redacted\w*|x{3,}|\*{3,}|todo|tbd|none|null|'
    r'true|false|\d+(?:\.\d+)?|'                       # bare numbers/bools are not secrets
    r'\$\{[^}]*\}|\{\{[^}]*\}\}|%\([^)]*\)s?|'         # ${VAR} {{var}} %(var)s
    r'os\.environ\S*|getenv\S*|process\.env\S*)'
    r'(?:["\']?)$',
    re.IGNORECASE,
)

# ── Credential-ish key words (shared by the assignment + .env.example rules) ───
_CRED_KEY_WORDS = (
    r'(?:API[_-]?KEY|API[_-]?SECRET|ACCESS[_-]?TOKEN|AUTH[_-]?TOKEN|BOT[_-]?TOKEN|'
    r'WEBHOOK[_-]?SECRET|CLIENT[_-]?SECRET|PRIVATE[_-]?KEY|SECRET|TOKEN|PASSWORD|'
    r'PASSWD|TOTP)'
)
_CRED_KEY_RE = re.compile(_CRED_KEY_WORDS, re.IGNORECASE)

# ── High-confidence standalone shapes ─────────────────────────────────────────
_BOT_TOKEN_RE = re.compile(r'\b\d{8,10}:[A-Za-z0-9_-]{35}\b')          # Telegram bot token

# ── Credential-style assignment: KEY (credential-ish) = VALUE ─────────────────
_CRED_ASSIGN_RE = re.compile(
    r'(?P<key>[A-Za-z0-9_]*' + _CRED_KEY_WORDS + r'[A-Za-z0-9_]*)'
    r'\s*[=:]\s*'
    r'(?P<q>["\']?)(?P<val>[^"\'\s#]+)(?P=q)',
    re.IGNORECASE,
)

_SECRETY_VAL_RE = re.compile(r'^[A-Za-z0-9+/=_\-]{16,}$')   # key-ish charset, >=16
_HEX64_RE = re.compile(r'^[0-9a-fA-F]{64}$')                # e.g. token_hex(32)
_BASE32_RE = re.compile(r'^[A-Z2-7]{16,}$')                 # e.g. a TOTP seed
_ENV_LINE_RE = re.compile(r'^\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*)$')


def _is_placeholder(val: str) -> bool:
    return bool(_PLACEHOLDER_RE.match(val.strip()))


def _looks_secret(val: str) -> bool:
    """True only for values that plausibly ARE a real credential."""
    v = val.strip().strip('"\'')
    if not v or _is_placeholder(v):
        return False
    if _HEX64_RE.match(v) or _BASE32_RE.match(v):
        return True
    if not _SECRETY_VAL_RE.match(v):
        return False
    # require mixed alnum (a path/word/number alone is not treated as a secret)
    return any(c.isalpha() for c in v) and any(c.isdigit() for c in v)


def scan_content(path: str, text: str) -> list[str]:
    """Return a list of finding descriptions (file:line + rule; NO secret value)."""
    findings: list[str] = []
    base = os.path.basename(path)

    # Rule 0: a real .env (not a template) must never be committed.
    if base == ".env":
        findings.append(f"{path}: a real .env file must never be committed (rule=env_file)")

    is_env_example = base == ".env.example" or base.endswith(".env.example")

    for i, line in enumerate(text.splitlines(), 1):
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue

        # Rule 1: telegram bot-token shape anywhere.
        if _BOT_TOKEN_RE.search(line):
            findings.append(f"{path}:{i}: telegram bot-token shape (rule=bot_token)")

        # Rule 2: credential-named assignment with a real-looking value.
        m = _CRED_ASSIGN_RE.search(line)
        if m and _looks_secret(m.group("val")):
            findings.append(
                f"{path}:{i}: credential assignment '{m.group('key')}' has a "
                f"real-looking value (rule=cred_assign)"
            )

        # Rule 3: .env.example must not carry a real credential (a credential-named
        # key, or a secret-looking value, must be a placeholder).
        if is_env_example:
            em = _ENV_LINE_RE.match(line)
            if em:
                key, val = em.group(1), em.group(2).strip()
                if val and not _is_placeholder(val) and (
                    _CRED_KEY_RE.search(key) or _looks_secret(val)
                ):
                    findings.append(
                        f"{path}:{i}: non-placeholder credential value for '{key}' in a "
                        f"template (.env.example) (rule=env_example_nonplaceholder)"
                    )
    return findings


def scan_xlsx(path: str, data: bytes) -> list[str]:
    """Scan an .xlsx workbook's cell values for secret shapes.

    C-1 gap: a binary .xlsx (e.g. `credentials.xlsx`) bypasses the text scanner
    entirely (scan_staged skips null-byte blobs). This inspects every cell string.
    FAIL CLOSED: if the workbook cannot be parsed (openpyxl missing / corrupt), BLOCK
    — an un-inspectable spreadsheet could hide a credential. Findings never contain
    the value (only file[sheet] + rule)."""
    try:
        import io
        import openpyxl
    except Exception:  # noqa: BLE001
        return [f"{path}: cannot scan .xlsx — openpyxl unavailable; refusing to commit an "
                f"un-inspectable spreadsheet (rule=xlsx_unscannable)"]
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        return [f"{path}: cannot open .xlsx ({type(exc).__name__}); refusing to commit an "
                f"un-inspectable spreadsheet (rule=xlsx_unparseable)"]
    findings: list[str] = []
    seen: set = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is None:
                    continue
                v = str(cell).strip()
                if not v:
                    continue
                rule = None
                if _BOT_TOKEN_RE.search(v):
                    rule = "xlsx_bot_token"
                elif _looks_secret(v):
                    rule = "xlsx_secret_value"
                if rule and (ws.title, rule) not in seen:
                    seen.add((ws.title, rule))
                    findings.append(f"{path} [sheet={ws.title}]: real-looking secret in a cell (rule={rule})")
    return findings


def _staged_files() -> list[str]:
    out = subprocess.run(
        ["git", "diff", "--cached", "--name-only", "--diff-filter=ACM", "-z"],
        capture_output=True, text=True, encoding="utf-8", errors="replace", check=True,
    )
    return [f for f in (out.stdout or "").split("\0") if f]


def scan_staged() -> int:
    """Scan staged blob content; return 1 (block) if any finding, else 0."""
    try:
        files = _staged_files()
    except Exception as exc:  # noqa: BLE001 -- git infra hiccup: don't brick all commits
        print(f"secret_scan: could not list staged files ({exc}); skipping", file=sys.stderr)
        return 0
    findings: list[str] = []
    for f in files:
        # .xlsx is binary (a zip) → read the staged blob as BYTES and parse cells
        # (the text path below would skip it as binary — the C-1 credentials.xlsx gap).
        if f.lower().endswith(".xlsx"):
            raw = subprocess.run(["git", "show", f":{f}"], capture_output=True)  # bytes
            if raw.returncode == 0 and raw.stdout:
                findings.extend(scan_xlsx(f, raw.stdout))
            continue
        # Decode git output as UTF-8 explicitly: on Windows text=True would use the
        # locale codec (cp1252) and crash on UTF-8 bytes (arrows/em-dashes in files).
        blob = subprocess.run(
            ["git", "show", f":{f}"],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
        )
        if blob.returncode != 0 or not blob.stdout:
            continue
        if "\x00" in blob.stdout[:2048]:   # binary-ish; skip
            continue
        findings.extend(scan_content(f, blob.stdout))
    if findings:
        bar = "=" * 64
        print(bar, file=sys.stderr)
        print("SECRET SCAN BLOCKED THE COMMIT -- remove the secret(s) below:", file=sys.stderr)
        for fnd in findings:
            print("  BLOCKED " + fnd, file=sys.stderr)
        print("(values redacted. Use a placeholder, or unstage the file.)", file=sys.stderr)
        print(bar, file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(scan_staged())
